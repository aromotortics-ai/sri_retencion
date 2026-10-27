# -*- coding: utf-8 -*-
"""
Conciliación de FACTURAS RECIBIDAS (compras): SRI vs Odoo.

Entradas (carpeta facturas_recibidas/):
  - fact_reci_MMAAAA.txt: TXT del portal SRI "Comprobantes electrónicos recibidos"
    (ISO-8859-1, tab). Trae base, IVA e importe total de cada factura.
  - fact_reci_odoo.xlsx: facturas de proveedor exportadas de Odoo (16 columnas:
    Número, Número de factura, Socio, Sucursal, Fecha, ..., Importe sin impuestos
    con signo, Total con signo, Estado, Número de autorización).

Uso:
    python conciliar_recibidas.py facturas_recibidas salida.xlsx

Cruce en tres niveles (se registra en Metodo_Cruce):
  1. Clave de acceso (el Excel pierde el cero inicial: se completa a 49 dígitos).
  2. N° de factura + importe total (clave de Odoo mal digitada o sin clave).
  3. N° de factura único en el SRI, importe distinto (se reporta como Diferencia de valor).
Documentos propios de Odoo que nunca estarán en "recibidas" se separan:
  LEAR/LECO/LETC = liquidaciones de compra emitidas por la empresa;
  CFAR/CFCO/CFTE/VBICO/NTAR/CEXAR sin clave electrónica = documentos no electrónicos.
"""
import glob
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # raíz del repo
from app.reportes.estilos import FONT, _enc, _fuente  # noqa: E402

TOL = 0.02
LIQ_PROPIAS = ("LEAR", "LECO", "LETC")
NO_ELECTRONICOS = ("CFAR", "CFCO", "CFTE", "VBICO", "NTAR", "CEXAR")
CAT_LIQ = "Liquidación propia (no aplica)"
CAT_NOELEC = "No electrónico (informativo)"
CAT_SOLO_ODOO = "Solo en Odoo"
CAT_DUP = "Duplicado en Odoo"
COLORES = {"Resumen_Mensual": "1F4E79", "Conciliacion": "1F4E79", "Solo_en_SRI": "ED7D31",
           "Solo_en_Odoo": "C00000", "Diferencia_Valor": "FFC000", "Duplicados_Odoo": "7030A0",
           "SRI_Listado": "808080", "Odoo_Lineas": "808080"}


def leer_sri(carpeta):
    archivos = sorted(glob.glob(os.path.join(carpeta, "fact_reci_*.txt")))
    if not archivos:
        raise ValueError("No se encontraron archivos fact_reci_*.txt del SRI.")
    s = pd.concat([pd.read_csv(f, sep="\t", encoding="iso-8859-1", dtype=str) for f in archivos])
    s = s.drop_duplicates("CLAVE_ACCESO").copy()
    s["fe"] = pd.to_datetime(s["FECHA_EMISION"], format="%d/%m/%Y")
    for col, n in (("VALOR_SIN_IMPUESTOS", "base"), ("IVA", "iva"), ("IMPORTE_TOTAL", "total")):
        s[n] = pd.to_numeric(s[col], errors="coerce").fillna(0.0)
    s["k2"] = s["SERIE_COMPROBANTE"] + "|" + s["total"].round(2).astype(str)
    return s.sort_values(["fe", "SERIE_COMPROBANTE"]).reset_index(drop=True)


def leer_odoo(path):
    o = pd.read_excel(path, dtype={15: str})
    if len(o.columns) < 16:
        raise ValueError(f"El Excel de Odoo tiene {len(o.columns)} columnas; se esperan 16.")
    o = o.iloc[:, :16].copy()
    o.columns = ["num", "nfac", "socio", "suc", "fecha", "venc", "ref", "act", "iva", "iva0",
                 "exe", "noobj", "sinimp", "total", "estado", "aut"]
    o["prefijo"] = o["num"].astype(str).str.split("/").str[0]
    aut = o["aut"].astype("string").str.strip()
    o["clave"] = aut.where(aut.str.len() >= 45).str.zfill(49)
    o["fecha"] = pd.to_datetime(o["fecha"])
    o["tot"] = o["total"].abs().round(2)
    o["sub"] = o["sinimp"].abs().round(2)
    o["k2"] = o["nfac"].astype(str) + "|" + o["tot"].astype(str)
    o.insert(0, "id", range(1, len(o) + 1))
    return o


def conciliar(o, s):
    """DataFrame de pares (idx_sri, id_odoo, metodo, cat); una fila por cada lado sin pareja."""
    clave_sri = dict(zip(s["CLAVE_ACCESO"], s.index))
    k2_sri = {}
    for i, k in zip(s.index, s["k2"]):
        k2_sri.setdefault(k, i)
    serie_cnt = s["SERIE_COMPROBANTE"].value_counts()
    serie_idx = dict(zip(s["SERIE_COMPROBANTE"], s.index))
    pares, usados = [], set()
    for r in o.itertuples():
        idx, met = None, ""
        if r.prefijo in LIQ_PROPIAS:
            pares.append((None, r.id, "", CAT_LIQ))
            continue
        if r.clave in clave_sri:
            idx, met = clave_sri[r.clave], "Clave de acceso"
        elif r.k2 in k2_sri:
            idx, met = k2_sri[r.k2], "N° factura + importe"
        elif serie_cnt.get(r.nfac, 0) == 1:
            idx, met = serie_idx[r.nfac], "N° factura (importe distinto)"
        if idx is not None:
            cat = CAT_DUP if idx in usados else ""
            usados.add(idx)
            pares.append((idx, r.id, met, cat))
        else:
            cat = CAT_NOELEC if (r.prefijo in NO_ELECTRONICOS and pd.isna(r.clave)) else CAT_SOLO_ODOO
            pares.append((None, r.id, "", cat))
    p = pd.DataFrame(pares, columns=["idx_sri", "id", "metodo", "cat"])
    solo_sri = [i for i in s.index if i not in usados]
    extra = pd.DataFrame({"idx_sri": solo_sri, "id": None, "metodo": "", "cat": ""})
    return pd.concat([p, extra], ignore_index=True)


# ----------------------------------------------------------------- Excel ---
def _hoja(wb, titulo, enc, filas, anchos, fmt=None):
    ws = wb.create_sheet(titulo)
    ws.append(enc)
    _enc(ws, 1, len(enc), COLORES[titulo])
    for f in filas:
        ws.append(f)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(enc))}{max(ws.max_row, 2)}"
    ws.sheet_properties.tabColor = COLORES[titulo]
    for col, nf in (fmt or {}).items():
        for r in range(2, ws.max_row + 1):
            ws[f"{col}{r}"].number_format = nf
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _fuente(ws)
    return ws


def _nan(v):
    return "" if pd.isna(v) else v


def construir_excel(o, s, p, salida):
    wb = Workbook()
    wb.remove(wb.active)
    ns, no = len(s), len(o)
    od = o.set_index("id")

    _hoja(wb, "SRI_Listado",
          ["Clave_Acceso", "Serie", "RUC_Emisor", "Proveedor", "Fecha_Emision", "Mes", "Base", "IVA", "Total"],
          [[r.CLAVE_ACCESO, r.SERIE_COMPROBANTE, r.RUC_EMISOR, r.RAZON_SOCIAL_EMISOR, r.fe.to_pydatetime(),
            r.fe.strftime("%Y-%m"), r.base, r.iva, r.total] for r in s.itertuples()],
          [52, 18, 15, 45, 13, 9, 12, 12, 12], {"E": "dd/mm/yyyy", "G": "#,##0.00", "H": "#,##0.00", "I": "#,##0.00"})
    _hoja(wb, "Odoo_Lineas",
          ["ID_Odoo", "Numero_Odoo", "Factura", "Proveedor", "Sucursal", "Fecha", "Mes", "Tipo_Doc", "Referencia",
           "Subtotal", "Total", "Clave_Odoo"],
          [[r.id, r.num, _nan(r.nfac), r.socio, r.suc, r.fecha.to_pydatetime(), r.fecha.strftime("%Y-%m"),
            r.prefijo, _nan(r.ref), r.sub, r.tot, _nan(r.clave)] for r in o.itertuples()],
          [8, 20, 19, 45, 13, 12, 9, 9, 35, 12, 12, 52], {"F": "dd/mm/yyyy", "J": "#,##0.00", "K": "#,##0.00"})

    sa, so = f"SRI_Listado!$A$2:$A${ns + 1}", f"Odoo_Lineas!$A$2:$A${no + 1}"

    def idx(col, hoja, n, pos):
        return f"INDEX({hoja}!${col}$2:${col}${n + 1},{pos})"

    p = p.copy()
    p["_ord"] = [(s.at[int(i), "fe"] if pd.notna(i) else od.at[j, "fecha"]) for i, j in zip(p["idx_sri"], p["id"])]
    p["mes"] = [d.strftime("%Y-%m") for d in p["_ord"]]
    p = p.sort_values("_ord", kind="stable").reset_index(drop=True)

    enc = ["Clave_SRI", "ID_Odoo", "Mes", "Fecha_SRI", "Fecha_Odoo", "Proveedor", "Factura", "Numero_Odoo",
           "Sucursal", "Metodo_Cruce", "Total_SRI", "Total_Odoo", "Diferencia", "Resultado", "Categoria_Odoo",
           "_fila_SRI", "_fila_Odoo"]
    ws = wb.create_sheet("Conciliacion", 0)
    ws.append(enc)
    _enc(ws, 1, len(enc), COLORES["Conciliacion"])
    for i, r in enumerate(p.itertuples(), start=2):
        clave = s.at[int(r.idx_sri), "CLAVE_ACCESO"] if pd.notna(r.idx_sri) else ""
        idod = int(r.id) if pd.notna(r.id) else ""
        ws.append([
            clave, idod, r.mes,
            f'=IF(P{i}="","",{idx("E", "SRI_Listado", ns, f"P{i}")})',
            f'=IF(Q{i}="","",{idx("F", "Odoo_Lineas", no, f"Q{i}")})',
            f'=IF(P{i}<>"",{idx("D", "SRI_Listado", ns, f"P{i}")},{idx("D", "Odoo_Lineas", no, f"Q{i}")})',
            f'=IF(Q{i}<>"",{idx("C", "Odoo_Lineas", no, f"Q{i}")},{idx("B", "SRI_Listado", ns, f"P{i}")})',
            f'=IF(Q{i}="","",{idx("B", "Odoo_Lineas", no, f"Q{i}")})',
            f'=IF(Q{i}="","",{idx("E", "Odoo_Lineas", no, f"Q{i}")})',
            r.metodo,
            f'=IF(P{i}="","",{idx("I", "SRI_Listado", ns, f"P{i}")})',
            f'=IF(Q{i}="","",{idx("K", "Odoo_Lineas", no, f"Q{i}")})',
            f'=IF(AND(P{i}<>"",Q{i}<>""),ROUND(K{i}-L{i},2),"")',
            f'=IF(O{i}="{CAT_DUP}",O{i},IF(AND(P{i}<>"",Q{i}<>""),IF(ABS(M{i})>{TOL},"Diferencia de valor","Coincide"),'
            f'IF(P{i}<>"","Solo en SRI",O{i})))',
            r.cat,
            f'=IF(A{i}="","",MATCH(A{i},{sa},0))',
            f'=IF(B{i}="","",MATCH(B{i},{so},0))',
        ])
    n = len(p)
    for r in range(2, n + 2):
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = "dd/mm/yyyy"
        for col in "KLM":
            ws[f"{col}{r}"].number_format = "#,##0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:Q{n + 1}"
    ws.sheet_properties.tabColor = COLORES["Conciliacion"]
    for i, w in enumerate([52, 9, 9, 12, 12, 42, 19, 20, 13, 26, 12, 12, 11, 22, 28, 9, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["P"].hidden = ws.column_dimensions["Q"].hidden = True
    _fuente(ws)

    res = p.copy()
    res["tot_sri"] = [s.at[int(i), "total"] if pd.notna(i) else None for i in res["idx_sri"]]
    res["tot_od"] = [od.at[j, "tot"] if pd.notna(j) else None for j in res["id"]]
    res["dif"] = (res["tot_sri"].astype(float) - res["tot_od"].astype(float)).round(2)
    es_dif = res["tot_sri"].notna() & res["tot_od"].notna() & (res["dif"].abs() > TOL) & (res["cat"] != CAT_DUP)

    sol_sri = res[res["idx_sri"].notna() & res["id"].isna()]
    _hoja(wb, "Solo_en_SRI", ["Fecha_Emision", "Serie", "RUC_Emisor", "Proveedor", "Base", "IVA", "Total", "Clave_Acceso"],
          [[s.at[int(i), "fe"].to_pydatetime(), s.at[int(i), "SERIE_COMPROBANTE"], s.at[int(i), "RUC_EMISOR"],
            s.at[int(i), "RAZON_SOCIAL_EMISOR"], s.at[int(i), "base"], s.at[int(i), "iva"], s.at[int(i), "total"],
            s.at[int(i), "CLAVE_ACCESO"]] for i in sol_sri["idx_sri"]],
          [13, 18, 15, 45, 12, 12, 12, 52], {"A": "dd/mm/yyyy", "E": "#,##0.00", "F": "#,##0.00", "G": "#,##0.00"})
    sol_od = res[res["idx_sri"].isna() & res["id"].notna()]
    _hoja(wb, "Solo_en_Odoo", ["Categoria", "Numero_Odoo", "Factura", "Proveedor", "Sucursal", "Fecha", "Tipo_Doc",
                                "Referencia", "Total", "Clave_Odoo"],
          [[r.cat, od.at[r.id, "num"], _nan(od.at[r.id, "nfac"]), od.at[r.id, "socio"], od.at[r.id, "suc"],
            od.at[r.id, "fecha"].to_pydatetime(), od.at[r.id, "prefijo"], _nan(od.at[r.id, "ref"]),
            od.at[r.id, "tot"], _nan(od.at[r.id, "clave"])] for r in sol_od.sort_values(["cat", "_ord"]).itertuples()],
          [30, 20, 19, 45, 13, 12, 9, 35, 12, 52], {"F": "dd/mm/yyyy", "I": "#,##0.00"})
    dif = res[es_dif]
    _hoja(wb, "Diferencia_Valor", ["Fecha_SRI", "Proveedor", "Factura", "Numero_Odoo", "Metodo_Cruce", "Total_SRI",
                                    "Total_Odoo", "Diferencia"],
          [[s.at[int(r.idx_sri), "fe"].to_pydatetime(), s.at[int(r.idx_sri), "RAZON_SOCIAL_EMISOR"],
            s.at[int(r.idx_sri), "SERIE_COMPROBANTE"], od.at[r.id, "num"], r.metodo, r.tot_sri, r.tot_od, r.dif]
           for r in dif.itertuples()],
          [12, 45, 19, 20, 26, 12, 12, 11], {"A": "dd/mm/yyyy", "F": "#,##0.00", "G": "#,##0.00", "H": "#,##0.00"})
    claves_dup = set(res.loc[res["cat"] == CAT_DUP, "idx_sri"])
    dd = res[res["idx_sri"].isin(claves_dup) & res["id"].notna()].sort_values(["idx_sri", "_ord"])
    _hoja(wb, "Duplicados_Odoo", ["Factura_SRI", "Proveedor", "Numero_Odoo", "Factura_Odoo", "Sucursal", "Fecha_Odoo",
                                   "Total_Odoo"],
          [[s.at[int(r.idx_sri), "SERIE_COMPROBANTE"], s.at[int(r.idx_sri), "RAZON_SOCIAL_EMISOR"],
            od.at[r.id, "num"], _nan(od.at[r.id, "nfac"]), od.at[r.id, "suc"], od.at[r.id, "fecha"].to_pydatetime(),
            od.at[r.id, "tot"]] for r in dd.itertuples()],
          [19, 45, 20, 19, 13, 12, 12], {"F": "dd/mm/yyyy", "G": "#,##0.00"})

    rs = wb.create_sheet("Resumen_Mensual", 0)
    rs["A1"] = "Conciliación de facturas RECIBIDAS: SRI vs Odoo — resumen mensual"
    rs["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E79")
    rs["A2"] = ("RUC 2390000831001  |  Cruce: clave de acceso, luego N° factura + importe  |  "
                "Mes = mes de emisión (SRI) o fecha de la factura (Odoo si no está en SRI)")
    rs["A2"].font = Font(name=FONT, italic=True, size=9)
    e = ["Mes", "Facturas SRI", "Facturas Odoo (aplicables)", "Coinciden", "Diferencia de valor", "Solo en SRI",
         "Solo en Odoo", "Duplicado en Odoo", "No electrónico (info)", "Liquidación propia (no aplica)",
         "% Coincidencia", "Total SRI", "Total Odoo", "Dif. total"]
    for j, h in enumerate(e, 1):
        rs.cell(row=4, column=j, value=h)
    _enc(rs, 4, len(e), COLORES["Resumen_Mensual"])
    meses = sorted(p["mes"].unique())
    cm, cr = f"Conciliacion!$C$2:$C${n + 1}", f"Conciliacion!$N$2:$N${n + 1}"
    for k, m in enumerate(meses):
        r = 5 + k
        rs[f"A{r}"] = m
        rs[f"B{r}"] = f"=COUNTIFS(SRI_Listado!$F$2:$F${ns + 1},$A{r})"
        rs[f"C{r}"] = f"=D{r}+E{r}+G{r}+H{r}"
        rs[f"D{r}"] = f'=COUNTIFS({cm},$A{r},{cr},"Coincide")'
        rs[f"E{r}"] = f'=COUNTIFS({cm},$A{r},{cr},"Diferencia de valor")'
        rs[f"F{r}"] = f'=COUNTIFS({cm},$A{r},{cr},"Solo en SRI")'
        rs[f"G{r}"] = f'=COUNTIFS({cm},$A{r},{cr},"{CAT_SOLO_ODOO}")'
        rs[f"H{r}"] = f'=COUNTIFS({cm},$A{r},{cr},"{CAT_DUP}")'
        rs[f"I{r}"] = f'=COUNTIFS({cm},$A{r},{cr},"{CAT_NOELEC}")'
        rs[f"J{r}"] = f'=COUNTIFS({cm},$A{r},{cr},"{CAT_LIQ}")'
        rs[f"K{r}"] = f"=IF(B{r}=0,0,(D{r}+E{r})/B{r})"
        rs[f"L{r}"] = f"=SUMIFS(SRI_Listado!$I$2:$I${ns + 1},SRI_Listado!$F$2:$F${ns + 1},$A{r})"
        rs[f"M{r}"] = "=" + "+".join(
            f'SUMIFS(Conciliacion!$L$2:$L${n + 1},{cm},$A{r},{cr},"{x}")'
            for x in ("Coincide", "Diferencia de valor", CAT_SOLO_ODOO, CAT_DUP))
        rs[f"N{r}"] = f"=L{r}-M{r}"
    ult = 4 + len(meses)
    tot = ult + 1
    rs[f"A{tot}"] = "Total"
    for col in "BCDEFGHIJLMN":
        rs[f"{col}{tot}"] = f"=SUM({col}5:{col}{ult})"
    rs[f"K{tot}"] = f"=IF(B{tot}=0,0,(D{tot}+E{tot})/B{tot})"
    for r in range(5, tot + 1):
        rs[f"K{r}"].number_format = "0.0%"
        for col in "LMN":
            rs[f"{col}{r}"].number_format = "#,##0.00"
    _fuente(rs, 4)
    for col in "ABCDEFGHIJKLMN":
        rs[f"{col}{tot}"].font = Font(name=FONT, bold=True)
        rs[f"{col}{tot}"].fill = PatternFill("solid", fgColor="DDEBF7")
    obs = tot + 2
    notas = [
        "Observaciones",
        "Coinciden = factura en SRI y Odoo con importe total igual (tolerancia $0.02). 'Facturas Odoo (aplicables)' excluye "
        "documentos no electrónicos y liquidaciones propias.",
        "Solo en SRI = factura recibida autorizada en el SRI sin registro en Odoo. Solo en Odoo = registrada en Odoo con "
        "clave/N° que no figura en el SRI (clave mal digitada, otro período o factura no recibida).",
        "Duplicado en Odoo = la misma factura del SRI registrada más de una vez en Odoo (ver hoja Duplicados_Odoo).",
        "No electrónico: CFAR, CFCO, CFTE, VBICO, NTAR, CEXAR sin clave de acceso (RISE, recibos, asientos, aduana, exterior).",
        "Liquidación propia: LEAR/LECO/LETC son liquidaciones de compra emitidas por la empresa; no aparecen en recibidas.",
        "Total Odoo (col. M) suma Coinciden + Diferencia + Solo en Odoo + Duplicados; compárese con Total SRI (col. L).",
    ]
    for k, t in enumerate(notas):
        c = rs[f"A{obs + k}"]
        c.value = t
        c.font = Font(name=FONT, bold=(k == 0), size=9 if k else 10)
    rs.freeze_panes = "B5"
    rs.sheet_properties.tabColor = COLORES["Resumen_Mensual"]
    for i, w in enumerate([10, 12, 16, 12, 13, 12, 12, 13, 14, 16, 13, 15, 15, 13], 1):
        rs.column_dimensions[get_column_letter(i)].width = w
    rs.row_dimensions[4].height = 42
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(salida)
    return res


def main(carpeta, salida):
    s = leer_sri(carpeta)
    o = leer_odoo(os.path.join(carpeta, "fact_reci_odoo.xlsx"))
    p = conciliar(o, s)
    res = construir_excel(o, s, p, salida)
    print(res["cat"].replace("", "(emparejadas / solo SRI)").value_counts().to_string())
    return res


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    main(*sys.argv[1:3])
