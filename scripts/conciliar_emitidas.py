# -*- coding: utf-8 -*-
"""
Conciliación de retenciones EMITIDAS por la empresa: SRI vs Odoo.

Entradas:
  - Excel exportado de Odoo (retenciones de compras): Factura, Contacto,
    Sucursal, Posición fiscal, No. de retención, Fecha de retención, Asiento
    contable, Asiento contable/Fecha, y por línea: Base imponible, Porcentaje,
    Valor retenido, % de retención (código), Categoría (IR / IVA).
  - Carpeta con los TXT diarios del portal SRI ("Comprobantes electrónicos
    emitidos" -> Comprobante de Retención -> Descargar reporte):
    emitidos_AAAA-MM-DD.txt (ISO-8859-1, tab).

Uso:
    python conciliar_emitidas.py retenciones_emitidas/retenciones_emitidas_2025_2026.xlsx \
        retenciones_emitidas salida.xlsx [AAAA-MM-DD desde] [AAAA-MM-DD hasta]

Clave de cruce: N° de retención (NNN-NNN-NNNNNNNNN). El TXT del SRI no trae
montos para retenciones y el web service público solo conserva los XML de los
últimos ~30 días, por lo que la conciliación es a nivel de comprobante
(presencia en SRI vs Odoo); los valores retenidos provienen de Odoo.
"""
import glob
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

FONT = "Arial"
COLORES = {  # convención del proyecto
    "Resumen_Mensual": "1F4E79", "Conciliacion": "1F4E79", "Solo_en_SRI": "ED7D31",
    "Solo_en_Odoo": "C00000", "Fecha_distinta": "FFC000",
    "SRI_Listado": "808080", "Odoo_Lineas": "808080",
}
COLS_ODOO = ["factura", "contacto", "sucursal", "pos_fiscal", "n_ret", "fecha_ret",
             "asiento", "fecha_asiento", "base", "pct", "valor", "codigo", "categoria"]
RE_NRET = re.compile(r"^\d{3}-\d{3}-\d{9}$")


def leer_odoo(path, desde=None, hasta=None):
    o = pd.read_excel(path)
    if len(o.columns) < len(COLS_ODOO):
        raise ValueError(f"El Excel de Odoo tiene {len(o.columns)} columnas; se esperan {len(COLS_ODOO)}.")
    o = o.iloc[:, :len(COLS_ODOO)].copy()
    o.columns = COLS_ODOO
    o = o[o["n_ret"].notna()].copy()
    o["n_ret"] = o["n_ret"].astype(str).str.strip()
    malos = o[~o["n_ret"].str.match(RE_NRET)]
    if len(malos):
        raise ValueError(f"{len(malos)} líneas de Odoo con N° de retención inválido, p.ej. {malos['n_ret'].iloc[0]!r}")
    o["fecha_ret"] = pd.to_datetime(o["fecha_ret"])
    o["tipo"] = o["categoria"].astype(str).str.split(" ").str[0]
    if desde:
        o = o[o["fecha_ret"] >= pd.Timestamp(desde)]
    if hasta:
        o = o[o["fecha_ret"] <= pd.Timestamp(hasta)]
    return o.reset_index(drop=True)


def leer_sri(carpeta, desde=None, hasta=None):
    archivos = sorted(glob.glob(os.path.join(carpeta, "emitidos_*.txt")))
    if not archivos:
        raise ValueError("No se encontraron archivos emitidos_*.txt del SRI.")
    filas = []
    for f in archivos:
        with open(f, encoding="iso-8859-1") as fh:
            lineas = fh.read().splitlines()
        cab = lineas[0].split("\t")
        for col in ("SERIE_COMPROBANTE", "CLAVE_ACCESO", "FECHA_EMISION", "FECHA_AUTORIZACION"):
            if col not in cab:
                raise ValueError(f"{os.path.basename(f)}: falta la columna {col}")
        i = {c: cab.index(c) for c in cab}
        for ln in lineas[1:]:
            p = ln.split("\t")
            if len(p) <= max(i["SERIE_COMPROBANTE"], i["CLAVE_ACCESO"], i["FECHA_EMISION"]) or not p[i["CLAVE_ACCESO"]].strip():
                continue
            filas.append({
                "n_ret": p[i["SERIE_COMPROBANTE"]].strip(),
                "fecha_emision": pd.to_datetime(p[i["FECHA_EMISION"]][:10], format="%d/%m/%Y"),
                "fecha_aut": p[i["FECHA_AUTORIZACION"]].strip(),
                "clave": p[i["CLAVE_ACCESO"]].strip(),
            })
    s = pd.DataFrame(filas).drop_duplicates("clave")
    if desde:
        s = s[s["fecha_emision"] >= pd.Timestamp(desde)]
    if hasta:
        s = s[s["fecha_emision"] <= pd.Timestamp(hasta)]
    return s.sort_values(["fecha_emision", "n_ret"]).reset_index(drop=True)


def conciliar_emitidas(odoo, sri):
    """Un registro por N° de retención presente en SRI y/o Odoo."""
    por_ret = odoo.groupby("n_ret").agg(
        fecha_odoo=("fecha_ret", "first"), proveedor=("contacto", "first"),
        sucursal=("sucursal", "first"), asiento=("asiento", "first"),
        facturas=("factura", lambda x: ", ".join(dict.fromkeys(x.astype(str)))),
        valor_ir=("valor", lambda x: x[odoo.loc[x.index, "tipo"] == "IR"].sum()),
        valor_iva=("valor", lambda x: x[odoo.loc[x.index, "tipo"] == "IVA"].sum()),
    ).reset_index()
    sri1 = sri.drop_duplicates("n_ret")[["n_ret", "fecha_emision", "clave"]]
    c = sri1.merge(por_ret, on="n_ret", how="outer")
    c["en_sri"] = c["clave"].notna()
    c["en_odoo"] = c["fecha_odoo"].notna()
    c["fecha_ref"] = c["fecha_emision"].fillna(c["fecha_odoo"])
    c["mes"] = c["fecha_ref"].dt.strftime("%Y-%m")
    c["resultado"] = "Coincide"
    c.loc[c["en_sri"] & ~c["en_odoo"], "resultado"] = "Solo en SRI"
    c.loc[~c["en_sri"] & c["en_odoo"], "resultado"] = "Solo en Odoo"
    return c.sort_values(["fecha_ref", "n_ret"]).reset_index(drop=True)


# ----------------------------------------------------------------- Excel ---
def _estilo_encabezado(ws, fila, ncols, color="1F4E79"):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=fila, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fuente_hoja(ws, desde_fila=1):
    for row in ws.iter_rows(min_row=desde_fila):
        for cell in row:
            if not cell.font or cell.font.name != FONT or not cell.font.bold:
                cell.font = Font(name=FONT, bold=cell.font.bold if cell.font else False,
                                 color=cell.font.color if cell.font and cell.font.color else None,
                                 size=cell.font.size if cell.font and cell.font.size else 10)


def _anchos(ws, anchos):
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hoja_datos(wb, titulo, encabezados, filas, anchos, formatos=None):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    _estilo_encabezado(ws, 1, len(encabezados), COLORES.get(titulo, "808080"))
    for f in filas:
        ws.append(f)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{max(ws.max_row, 2)}"
    ws.sheet_properties.tabColor = COLORES.get(titulo, "808080")
    for col, fmt in (formatos or {}).items():
        for r in range(2, ws.max_row + 1):
            ws[f"{col}{r}"].number_format = fmt
    _anchos(ws, anchos)
    _fuente_hoja(ws)
    return ws


def construir_excel(odoo, sri, c, salida, desde, hasta):
    wb = Workbook()
    wb.remove(wb.active)
    n_sri, n_odoo, n_c = len(sri), len(odoo), len(c)

    # --- datos base ------------------------------------------------------
    ws_sri = _hoja_datos(
        wb, "SRI_Listado", ["N_Retencion", "Fecha_Emision", "Fecha_Autorizacion", "Clave_Acceso"],
        [[r.n_ret, r.fecha_emision.to_pydatetime(), r.fecha_aut, r.clave] for r in sri.itertuples()],
        [20, 14, 22, 52], {"B": "dd/mm/yyyy"})
    ws_od = _hoja_datos(
        wb, "Odoo_Lineas",
        ["N_Retencion", "Fecha_Retencion", "Factura_Odoo", "Proveedor", "Sucursal", "Asiento",
         "Tipo", "Porcentaje", "Base", "Valor_Retenido", "Retencion_Codigo"],
        [[r.n_ret, r.fecha_ret.to_pydatetime(), r.factura, r.contacto, r.sucursal, r.asiento,
          r.tipo, r.pct, r.base, r.valor, r.codigo] for r in odoo.itertuples()],
        [20, 14, 20, 40, 14, 20, 7, 11, 13, 14, 60], {"B": "dd/mm/yyyy", "I": "#,##0.00", "J": "#,##0.00"})

    sri_a = f"SRI_Listado!$A$2:$A${n_sri + 1}"
    sri_b = f"SRI_Listado!$B$2:$B${n_sri + 1}"
    od_a = f"Odoo_Lineas!$A$2:$A${n_odoo + 1}"

    # --- Conciliacion (fórmulas) ------------------------------------------
    enc = ["N_Retencion", "Mes", "Fecha_SRI", "Fecha_Odoo", "Proveedor", "Facturas_Odoo", "Sucursal",
           "En_SRI", "En_Odoo", "Resultado", "Dif_Fecha_Dias", "Valor_IR_Odoo", "Valor_IVA_Odoo",
           "Valor_Total_Odoo", "Clave_Acceso"]
    ws = wb.create_sheet("Conciliacion", 0)
    ws.append(enc)
    _estilo_encabezado(ws, 1, len(enc), COLORES["Conciliacion"])
    for i, r in enumerate(c.itertuples(), start=2):
        ws.append([
            r.n_ret, r.mes,
            f'=IFERROR(INDEX({sri_b},MATCH($A{i},{sri_a},0)),"")',
            f'=IFERROR(INDEX(Odoo_Lineas!$B$2:$B${n_odoo + 1},MATCH($A{i},{od_a},0)),"")',
            "" if pd.isna(r.proveedor) else r.proveedor,
            "" if pd.isna(r.facturas) else r.facturas,
            "" if pd.isna(r.sucursal) else r.sucursal,
            f'=IF(COUNTIF({sri_a},$A{i})>0,"Sí","No")',
            f'=IF(COUNTIF({od_a},$A{i})>0,"Sí","No")',
            f'=IF(AND(H{i}="Sí",I{i}="Sí"),"Coincide",IF(H{i}="Sí","Solo en SRI","Solo en Odoo"))',
            f'=IF(AND(H{i}="Sí",I{i}="Sí"),C{i}-D{i},"")',
            f'=SUMIFS(Odoo_Lineas!$J$2:$J${n_odoo + 1},{od_a},$A{i},Odoo_Lineas!$G$2:$G${n_odoo + 1},"IR")',
            f'=SUMIFS(Odoo_Lineas!$J$2:$J${n_odoo + 1},{od_a},$A{i},Odoo_Lineas!$G$2:$G${n_odoo + 1},"IVA")',
            f"=L{i}+M{i}",
            "" if pd.isna(r.clave) else r.clave,
        ])
    for r in range(2, n_c + 2):
        ws[f"C{r}"].number_format = ws[f"D{r}"].number_format = "dd/mm/yyyy"
        for col in "LMN":
            ws[f"{col}{r}"].number_format = "#,##0.00"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:O{n_c + 1}"
    ws.sheet_properties.tabColor = COLORES["Conciliacion"]
    _anchos(ws, [20, 9, 12, 12, 40, 24, 13, 8, 9, 14, 10, 13, 13, 14, 52])
    _fuente_hoja(ws)

    # --- Solo en SRI / Solo en Odoo / Fecha distinta (listas) -------------
    solo_sri = c[c["resultado"] == "Solo en SRI"]
    _hoja_datos(wb, "Solo_en_SRI", ["N_Retencion", "Fecha_Emision", "Clave_Acceso"],
                [[r.n_ret, r.fecha_emision.to_pydatetime(), r.clave] for r in solo_sri.itertuples()],
                [20, 14, 52], {"B": "dd/mm/yyyy"})
    solo_od = odoo[odoo["n_ret"].isin(c.loc[c["resultado"] == "Solo en Odoo", "n_ret"])]
    _hoja_datos(wb, "Solo_en_Odoo",
                ["N_Retencion", "Fecha_Retencion", "Factura_Odoo", "Proveedor", "Sucursal", "Asiento",
                 "Tipo", "Porcentaje", "Base", "Valor_Retenido"],
                [[r.n_ret, r.fecha_ret.to_pydatetime(), r.factura, r.contacto, r.sucursal, r.asiento,
                  r.tipo, r.pct, r.base, r.valor] for r in solo_od.itertuples()],
                [20, 14, 20, 40, 14, 20, 7, 11, 13, 14], {"B": "dd/mm/yyyy", "I": "#,##0.00", "J": "#,##0.00"})
    dist = c[c["en_sri"] & c["en_odoo"] & (c["fecha_emision"] != c["fecha_odoo"])]
    _hoja_datos(wb, "Fecha_distinta",
                ["N_Retencion", "Fecha_SRI", "Fecha_Odoo", "Dif_Dias", "Proveedor", "Facturas_Odoo"],
                [[r.n_ret, r.fecha_emision.to_pydatetime(), r.fecha_odoo.to_pydatetime(),
                  (r.fecha_emision - r.fecha_odoo).days, r.proveedor, r.facturas] for r in dist.itertuples()],
                [20, 12, 12, 9, 40, 28], {"B": "dd/mm/yyyy", "C": "dd/mm/yyyy"})

    # --- Resumen_Mensual (fórmulas) ----------------------------------------
    rs = wb.create_sheet("Resumen_Mensual", 0)
    rs["A1"] = "Conciliación de retenciones EMITIDAS: SRI vs Odoo — resumen mensual"
    rs["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E79")
    rs["A2"] = (f"RUC 2390000831001  |  Período {desde} a {hasta}  |  Clave de cruce: N° de retención  |  "
                "Mes = mes de emisión según SRI (si no está en SRI, mes de Odoo)")
    rs["A2"].font = Font(name=FONT, italic=True, size=9)
    enc_r = ["Mes", "Retenciones SRI", "Retenciones Odoo", "Coinciden", "Solo en SRI", "Solo en Odoo",
             "% Coincidencia", "Valor IR Odoo", "Valor IVA Odoo", "Valor total Odoo"]
    for j, h in enumerate(enc_r, 1):
        rs.cell(row=4, column=j, value=h)
    _estilo_encabezado(rs, 4, len(enc_r), COLORES["Resumen_Mensual"])
    meses = sorted(c["mes"].unique())
    cm = f"Conciliacion!$B$2:$B${n_c + 1}"
    cres = f"Conciliacion!$J$2:$J${n_c + 1}"
    cen_sri = f"Conciliacion!$H$2:$H${n_c + 1}"
    cen_od = f"Conciliacion!$I$2:$I${n_c + 1}"
    for k, m in enumerate(meses):
        r = 5 + k
        rs[f"A{r}"] = m
        rs[f"B{r}"] = f'=COUNTIFS({cm},$A{r},{cen_sri},"Sí")'
        rs[f"C{r}"] = f'=COUNTIFS({cm},$A{r},{cen_od},"Sí")'
        rs[f"D{r}"] = f'=COUNTIFS({cm},$A{r},{cres},"Coincide")'
        rs[f"E{r}"] = f'=COUNTIFS({cm},$A{r},{cres},"Solo en SRI")'
        rs[f"F{r}"] = f'=COUNTIFS({cm},$A{r},{cres},"Solo en Odoo")'
        rs[f"G{r}"] = f"=IF(B{r}=0,0,D{r}/B{r})"
        rs[f"H{r}"] = f"=SUMIFS(Conciliacion!$L$2:$L${n_c + 1},{cm},$A{r})"
        rs[f"I{r}"] = f"=SUMIFS(Conciliacion!$M$2:$M${n_c + 1},{cm},$A{r})"
        rs[f"J{r}"] = f"=H{r}+I{r}"
    ult = 4 + len(meses)
    tot = ult + 1
    rs[f"A{tot}"] = "Total"
    for col in "BCDEFHIJ":
        rs[f"{col}{tot}"] = f"=SUM({col}5:{col}{ult})"
    rs[f"G{tot}"] = f"=IF(B{tot}=0,0,D{tot}/B{tot})"
    for r in range(5, tot + 1):
        rs[f"G{r}"].number_format = "0.0%"
        for col in "HIJ":
            rs[f"{col}{r}"].number_format = "#,##0.00"
        if r == tot:
            for col in "ABCDEFGHIJ":
                rs[f"{col}{r}"].font = Font(name=FONT, bold=True)
                rs[f"{col}{r}"].fill = PatternFill("solid", fgColor="DDEBF7")
    obs = tot + 2
    rs[f"A{obs}"] = "Observaciones"
    rs[f"A{obs}"].font = Font(name=FONT, bold=True)
    rs[f"A{obs + 1}"] = (f'="Coinciden con fecha distinta entre SRI y Odoo: "&(COUNT(Conciliacion!$K$2:$K${n_c + 1})'
                         f"-COUNTIF(Conciliacion!$K$2:$K${n_c + 1},0))&\" (ver hoja Fecha_distinta)\"")
    rs[f"A{obs + 2}"] = ("El TXT del SRI no trae montos de retenciones y el web service público solo conserva XML "
                         "recientes; por eso se concilia a nivel de comprobante y los valores son los de Odoo.")
    rs[f"A{obs + 3}"] = ("Solo en SRI = autorizada en el SRI sin registro en Odoo. Solo en Odoo = registrada en Odoo "
                         "pero no figura entre los AUTORIZADOS del SRI (no autorizada / por procesar / N° erróneo).")
    for r in (obs + 1, obs + 2, obs + 3):
        rs[f"A{r}"].font = Font(name=FONT, size=9)
    rs.freeze_panes = "B5"
    rs.sheet_properties.tabColor = COLORES["Resumen_Mensual"]
    _anchos(rs, [12, 15, 15, 12, 12, 12, 14, 15, 15, 16])
    for row in rs.iter_rows(min_row=4, max_row=tot):
        for cell in row:
            if cell.font.name != FONT:
                cell.font = Font(name=FONT, bold=cell.font.bold, color=cell.font.color, size=10)
    rs.row_dimensions[4].height = 30

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(salida)


def main(odoo_xlsx, carpeta_sri, salida, desde="2025-04-01", hasta="2026-08-31"):
    odoo = leer_odoo(odoo_xlsx, desde, hasta)
    sri = leer_sri(carpeta_sri, desde, hasta)
    c = conciliar_emitidas(odoo, sri)
    construir_excel(odoo, sri, c, salida, desde, hasta)
    print(c["resultado"].value_counts().to_string())
    return c


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(1)
    main(*sys.argv[1:6])
