# -*- coding: utf-8 -*-
"""Hojas de retenciones recibidas (SRI vs Odoo) para el libro Excel de Compras."""
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import conciliar_retenciones as cr
from conciliar_recibidas import _enc, _fuente
from excel_compras import COLORES, TOL, _hoja, _proteger

HOJA_CONC = "Ret_Conciliacion"
_ENC_CONC = ["Estado", "N_Retencion", "Fecha", "RUC", "Proveedor", "Factura_SRI", "Factura_Odoo", "Pct_SRI",
             "Pct_Odoo", "Base_SRI", "Base_Odoo", "Valor_SRI", "Valor_Odoo", "Diferencia", "ID_Odoo", "Clave"]
_ANCH_CONC = [30, 18, 12, 15, 42, 19, 19, 8, 8, 12, 12, 12, 12, 11, 9, 52]


def _v(x):
    """Valor de celda: NA/NaN/NaT -> vacio; numpy/pandas -> tipos nativos."""
    if x is None or (not isinstance(x, str) and pd.isna(x)):
        return ""
    if isinstance(x, pd.Timestamp):
        return x.to_pydatetime()
    if hasattr(x, "item"):
        return x.item()
    return x


def _tiene(x):
    return x is not None and not pd.isna(x)


def _es_formula(r):
    """El estado se calcula en la hoja solo si es Coincide/Diferencia con ambos valores y mismo porcentaje.
    Las parejas por porcentaje distinto se quedan como texto: deben seguir siendo diferencia."""
    if r["estado"] not in (cr.EST_OK, cr.EST_DIF) or not (_tiene(r["valor_sri"]) and _tiene(r["valor_odoo"])):
        return False
    ps, po = r["pct_sri"], r["pct_odoo"]
    return (not _tiene(ps) and not _tiene(po)) or (_tiene(ps) and _tiene(po) and abs(ps - po) < 1e-9)


def _hoja_conciliacion(wb, ret):
    ws = wb.create_sheet(HOJA_CONC)
    ws.append(_ENC_CONC)
    _enc(ws, 1, len(_ENC_CONC), COLORES[HOJA_CONC])
    for i, r in enumerate((x for _, x in ret.iterrows()), start=2):
        ambos = _tiene(r["valor_sri"]) and _tiene(r["valor_odoo"])
        estado = (f'=IF(ABS(N{i})>{TOL},"{cr.EST_DIF}","{cr.EST_OK}")' if _es_formula(r) else r["estado"])
        dif = f'=ROUND(L{i}-M{i},2)' if ambos else ""
        ws.append([estado, _v(r["n_ret"]), _v(r["fecha"]), _v(r["ruc"]), _v(r["proveedor"]),
                   _v(r["factura_sri"]), _v(r["factura_odoo"]), _v(r["pct_sri"]), _v(r["pct_odoo"]),
                   _v(r["base_sri"]), _v(r["base_odoo"]), _v(r["valor_sri"]), _v(r["valor_odoo"]), dif,
                   _v(r["id_ret"]), _v(r["clave"])])
    n = len(ret)
    for rr in range(2, n + 2):
        ws[f"C{rr}"].number_format = "dd/mm/yyyy"
        for col in "JKLMN":
            ws[f"{col}{rr}"].number_format = "#,##0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_ENC_CONC))}{max(n, 1) + 1}"
    ws.sheet_properties.tabColor = COLORES[HOJA_CONC]
    for k, w in enumerate(_ANCH_CONC, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    _fuente(ws)
    _proteger(ws, exentas="AN")
    return ws


def _filas(df, campos):
    return [[_v(r[c]) for c in campos] for _, r in df.iterrows()]


def agregar_hojas_retenciones(wb, ret):
    """Crea las hojas Ret_* en `wb` y devuelve los contadores del resumen."""
    ret = ret.reset_index(drop=True)
    _hoja_conciliacion(wb, ret)
    est = ret["estado"]

    _hoja(wb, "Ret_Solo_SRI",
          ["N_Retencion", "Fecha", "RUC", "Proveedor", "Factura", "Pct", "Base", "Valor", "Clave"],
          _filas(ret[est == cr.EST_SOLO_SRI],
                 ["n_ret", "fecha", "ruc", "proveedor", "factura_sri", "pct_sri", "base_sri", "valor_sri", "clave"]),
          [18, 12, 15, 42, 19, 8, 12, 12, 52], {"B": "dd/mm/yyyy", "G": "#,##0.00", "H": "#,##0.00"})
    _hoja(wb, "Ret_Solo_Odoo",
          ["N_Retencion", "Fecha", "RUC", "Proveedor", "Factura", "Pct", "Base", "Valor", "ID_Odoo"],
          _filas(ret[est == cr.EST_SOLO_ODOO],
                 ["n_ret", "fecha", "ruc", "proveedor", "factura_odoo", "pct_odoo", "base_odoo", "valor_odoo",
                  "id_ret"]),
          [18, 12, 15, 42, 19, 8, 12, 12, 9], {"B": "dd/mm/yyyy", "G": "#,##0.00", "H": "#,##0.00"})
    ws_dif = _hoja(wb, "Ret_Diferencia_Valor",
                   ["N_Retencion", "Fecha", "Proveedor", "Factura", "Pct_SRI", "Pct_Odoo", "Base_SRI", "Base_Odoo",
                    "Valor_SRI", "Valor_Odoo", "Diferencia"],
                   _filas(ret[est == cr.EST_DIF],
                          ["n_ret", "fecha", "proveedor", "factura_sri", "pct_sri", "pct_odoo", "base_sri",
                           "base_odoo", "valor_sri", "valor_odoo"]),
                   [18, 12, 42, 19, 8, 8, 12, 12, 12, 12, 11],
                   {"B": "dd/mm/yyyy", **{c: "#,##0.00" for c in "GHIJK"}})
    # Diferencia = Valor_SRI - Valor_Odoo como formula (se escribe tras _proteger, que la volveria texto).
    for rr in range(2, ws_dif.max_row + 1):
        ws_dif[f"K{rr}"] = f"=ROUND(I{rr}-J{rr},2)"
    _hoja(wb, "Ret_Factura_Incorrecta",
          ["N_Retencion", "Fecha", "Proveedor", "Factura_SRI", "Factura_Odoo", "Pct", "Base", "Valor_SRI",
           "Valor_Odoo"],
          _filas(ret[est == cr.EST_FAC],
                 ["n_ret", "fecha", "proveedor", "factura_sri", "factura_odoo", "pct_sri", "base_sri", "valor_sri",
                  "valor_odoo"]),
          [18, 12, 42, 19, 19, 8, 12, 12, 12], {"B": "dd/mm/yyyy", **{c: "#,##0.00" for c in "GHI"}})
    _hoja(wb, "Ret_Informativos",
          ["Estado", "N_Retencion", "Fecha", "RUC", "Proveedor", "Factura_Odoo", "Pct", "Base", "Valor",
           "ID_Odoo", "Clave"],
          _filas(ret[est.isin([cr.EST_XML, cr.EST_INFO])],
                 ["estado", "n_ret", "fecha", "ruc", "proveedor", "factura_odoo", "pct_odoo", "base_odoo",
                  "valor_odoo", "id_ret", "clave"]),
          [34, 18, 12, 15, 42, 19, 8, 12, 12, 9, 52], {"C": "dd/mm/yyyy", "H": "#,##0.00", "I": "#,##0.00"})

    sin_info = ret[est != cr.EST_INFO]
    return {
        "coinciden": int((est == cr.EST_OK).sum()),
        "diferencia": int((est == cr.EST_DIF).sum()),
        "solo_sri": int((est == cr.EST_SOLO_SRI).sum()),
        "solo_odoo": int((est == cr.EST_SOLO_ODOO).sum()),
        "factura_incorrecta": int((est == cr.EST_FAC).sum()),
        "informativos": int(est.isin([cr.EST_XML, cr.EST_INFO]).sum()),
        "total_sri": round(float(ret["valor_sri"].sum()), 2),
        "total_odoo": round(float(sin_info["valor_odoo"].sum()), 2),
    }


def _texto_periodo(desde, hasta):
    """'MM/AAAA' si es un mes calendario completo; si no, 'dd/mm/aaaa a dd/mm/aaaa'."""
    d, h = pd.Timestamp(desde), pd.Timestamp(hasta)
    if d.day == 1 and h == d + pd.offsets.MonthEnd(0):
        return f"{d:%m/%Y}"
    return f"{d:%d/%m/%Y} a {h:%d/%m/%Y}"


def fila_resumen_retenciones(rs, r, ret_n, periodo=None, periodo_compras=None):
    """Escribe en la hoja Resumen (fila `r`) la fila 'Retenciones' con COUNTIFS sobre Ret_Conciliacion.
    Mismas columnas que la tabla por tipo. Devuelve la fila siguiente libre.

    periodo: (desde, hasta) de las retenciones; va en la columna Mes y, si difiere de
    periodo_compras, se rotula en la nota bajo la fila."""
    lim = max(ret_n, 1) + 1
    est, vs, vo = (f"{HOJA_CONC}!${c}$2:${c}${lim}" for c in ("A", "L", "M"))
    rs[f"A{r}"] = "Retenciones"
    nota_periodo = ""
    if periodo is not None:
        d, h = (pd.Timestamp(x) for x in periodo)
        rs[f"B{r}"] = (f"{d:%Y-%m}" if _texto_periodo(d, h) == f"{d:%m/%Y}"
                       else _texto_periodo(d, h))
        if periodo_compras is None or tuple(pd.Timestamp(x) for x in periodo_compras) != (d, h):
            nota_periodo = f'&"  |  Período de retenciones: {_texto_periodo(d, h)}"'
    rs[f"C{r}"] = f"=COUNT({vs})"
    for col, txt in zip("DEFG", (cr.EST_OK, cr.EST_DIF, cr.EST_SOLO_SRI, cr.EST_SOLO_ODOO)):
        rs[f"{col}{r}"] = f'=COUNTIFS({est},"{txt}")'
    rs[f"H{r}"] = 0
    rs[f"I{r}"] = f'=COUNTIFS({est},"{cr.EST_INFO}")+COUNTIFS({est},"{cr.EST_XML}")'
    rs[f"J{r}"] = f"=IF(C{r}=0,0,(D{r}+E{r})/C{r})"
    rs[f"K{r}"] = f"=SUM({vs})"
    rs[f"L{r}"] = f'=SUMIFS({vo},{est},"<>{cr.EST_INFO}")'
    rs[f"M{r}"] = f"=K{r}-L{r}"
    rs[f"J{r}"].number_format = "0.0%"
    for col in "KLM":
        rs[f"{col}{r}"].number_format = "#,##0.00"
    rs[f"A{r + 1}"] = (f'="Retenciones - factura incorrecta en Odoo: "&COUNTIFS({est},"{cr.EST_FAC}")'
                       f'&"  |  XML no disponible: "&COUNTIFS({est},"{cr.EST_XML}")' + nota_periodo)
    rs[f"A{r + 1}"].font = Font(name="Arial", italic=True, size=9)
    return r + 2
