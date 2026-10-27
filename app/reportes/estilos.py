# -*- coding: utf-8 -*-
"""Estilos compartidos de los libros Excel (fuente Arial, encabezados)."""
from openpyxl.styles import Alignment, Font, PatternFill

FONT = "Arial"


def _enc(ws, fila, n, color):
    for c in range(1, n + 1):
        cell = ws.cell(row=fila, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fuente(ws, desde=2):
    for row in ws.iter_rows(min_row=desde):
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=cell.font.bold)
