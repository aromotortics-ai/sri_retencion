# -*- coding: utf-8 -*-
"""Excel de observaciones: solo lo que hay que corregir en Odoo, una hoja por estado."""
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.compras.excel_compras import _proteger
from app.reportes.estilos import _enc, _fuente

QUE_CORREGIR = {
    "Solo en SRI": "Registrar el comprobante en Odoo",
    "Solo en Odoo": "Verificar el número/clave del documento en Odoo o anularlo si no existe en el SRI",
    "Diferencia de valor": "Corregir el valor en Odoo (ver detalle)",
    "Diferencia de datos": "Corregir clave de acceso, fecha o marca electrónica en Odoo (ver detalle)",
    "Factura incorrecta en Odoo": "Corregir el número de la factura sustento en la retención",
    "Odoo anulado, SRI autorizado": "Reactivar el documento en Odoo o anular el comprobante en el SRI",
    "Duplicado en Odoo": "Anular el registro duplicado en Odoo",
    "XML no disponible": "Reintentar: no se pudo leer el XML del SRI",
}
ACCIONABLES = frozenset(QUE_CORREGIR)
COLUMNAS_OBS = ["tipo", "estado", "documento", "fecha", "contraparte", "ruc", "valor_sri", "valor_odoo",
                "diferencia", "detalle", "que_corregir", "accionable"]

# Mismos títulos que las columnas de la tabla en pantalla (tablero.js, COLUMNAS)
TITULOS_OBS = ["Tipo", "Estado", "Documento", "Fecha", "Contraparte", "RUC", "Valor SRI", "Valor Odoo",
               "Diferencia", "Detalle", "Qué corregir en Odoo"]
_COLOR = "1F4E79"
_ANCHOS = [8, 28, 22, 12, 34, 16, 13, 13, 13, 50, 62]
_INVALIDOS = re.compile(r"[\[\]:*?/\\]")


def _nombre_hoja(estado):
    return _INVALIDOS.sub("_", estado)[:31]


def _fecha(x):
    """Fecha ISO (texto del tablero) -> date; vacío o no interpretable queda como estaba."""
    if isinstance(x, str) and x:
        try:
            return datetime.fromisoformat(x[:10]).date()
        except ValueError:
            return x
    return x if isinstance(x, (date, datetime)) else None


def escribir_excel_observaciones(tablero: dict, ruta: str) -> int:
    """Escribe en `ruta` las excepciones accionables del tablero; devuelve cuántas son."""
    obs = [e for e in tablero.get("excepciones", []) if e.get("accionable")]
    wb = Workbook()
    wb.remove(wb.active)
    if not obs:
        ws = wb.create_sheet("Sin_observaciones")
        ws["A1"] = "No hay observaciones: todo lo conciliado coincide."
        _fuente(ws, desde=1)
        ws.column_dimensions["A"].width = 60
        wb.save(ruta)
        return 0
    columnas = COLUMNAS_OBS[:-1]
    estados = [e for e in QUE_CORREGIR if any(o["estado"] == e for o in obs)]
    estados += sorted({o["estado"] for o in obs} - set(estados))  # defensivo: estado accionable no listado
    for estado in estados:
        ws = wb.create_sheet(_nombre_hoja(estado))
        ws.append(TITULOS_OBS)
        _enc(ws, 1, len(columnas), _COLOR)
        for o in (x for x in obs if x["estado"] == estado):
            ws.append([_fecha(o.get(c)) if c == "fecha" else o.get(c) for c in columnas])
        _fuente(ws)
        _proteger(ws)  # texto de SRI/Odoo que empieza con = + - @ no se guarda como fórmula
        for fila in range(2, ws.max_row + 1):
            ws.cell(row=fila, column=4).number_format = "dd/mm/yyyy"
            for col in (7, 8, 9):
                ws.cell(row=fila, column=col).number_format = "#,##0.00"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{ws.max_row}"
        for i, w in enumerate(_ANCHOS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(ruta)
    return len(obs)
