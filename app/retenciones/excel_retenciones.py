# -*- coding: utf-8 -*-
"""Hojas de retenciones recibidas (SRI vs Odoo) para el libro Excel de Compras.

Vista por DOCUMENTO (un comprobante de retencion = una fila, hoja Ret_Documentos) con el detalle por
LINEA (Renta, IVA...) en Ret_Conciliacion. Los totales por documento son formulas SUMIFS/COUNTIFS sobre
Ret_Conciliacion, enlazadas por la columna Documento.
"""
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import app.retenciones.conciliar_retenciones as cr
from app.reportes.estilos import _enc, _fuente
from app.compras.excel_compras import COLORES, TOL, _hoja, _proteger

HOJA_CONC = "Ret_Conciliacion"
HOJA_DOC = "Ret_Documentos"
_ENC_CONC = ["Estado", "N_Retencion", "Fecha", "RUC", "Proveedor", "Factura_SRI", "Factura_Odoo", "Pct_SRI",
             "Pct_Odoo", "Base_SRI", "Base_Odoo", "Valor_SRI", "Valor_Odoo", "Diferencia", "ID_Odoo", "Clave",
             "Documento", "Estado_Documento"]
_ANCH_CONC = [30, 18, 12, 15, 42, 19, 19, 8, 8, 12, 12, 12, 12, 11, 9, 52, 22, 30]
_ENC_DOC = ["Estado", "N_Retencion", "Fecha", "RUC", "Proveedor", "Factura_SRI", "Factura_Odoo", "Lineas_SRI",
            "Lineas_Odoo", "Valor_SRI", "Valor_Odoo", "Diferencia", "Detalle", "Documento"]
_ANCH_DOC = [30, 18, 12, 15, 42, 19, 19, 10, 10, 12, 12, 11, 44, 22]


def _v(x):
    """Valor de celda: NA/NaN/NaT -> vacio; numpy/pandas -> tipos nativos."""
    if x is None or (not isinstance(x, str) and pd.isna(x)):
        return ""
    if isinstance(x, pd.Timestamp):
        return x.to_pydatetime()
    if hasattr(x, "item"):
        return x.item()
    return x


def _tiene(x):
    return x is not None and not pd.isna(x)


def _es_formula(r):
    """El estado se calcula en la hoja solo si es Coincide/Diferencia con ambos valores y mismo porcentaje.
    Las parejas por porcentaje distinto se quedan como texto: deben seguir siendo diferencia."""
    if r["estado"] not in (cr.EST_OK, cr.EST_DIF) or not (_tiene(r["valor_sri"]) and _tiene(r["valor_odoo"])):
        return False
    ps, po = r["pct_sri"], r["pct_odoo"]
    return (not _tiene(ps) and not _tiene(po)) or (_tiene(ps) and _tiene(po) and abs(ps - po) < 1e-9)


class _Formulas:
    """Formulas por documento sobre Ret_Conciliacion (clave en la columna Documento = Q).

    El criterio se escribe "="&celda para que una clave que empiece con = < > se compare literal.
    Lado Odoo: las lineas informativas (cancelada/borrador) no cuentan, salvo en un documento que sea
    solo informativo (asi muestra su valor)."""

    def __init__(self, n_lineas):
        lim = max(n_lineas, 1) + 1
        self.rango = {c: f"{HOJA_CONC}!${c}$2:${c}${lim}" for c in ("A", "L", "M", "Q")}

    def _extra(self, lado, estado_doc):
        if lado == "sri" or estado_doc == cr.EST_INFO:
            return ""
        return f',{self.rango["A"]},"<>{cr.EST_INFO}"'

    def lineas(self, ref, lado, estado_doc):
        col, ex = self.rango["L" if lado == "sri" else "M"], self._extra(lado, estado_doc)
        clave = f'{self.rango["Q"]},"="&{ref}'
        return f'=COUNTIFS({clave},{col},">=0"{ex})+COUNTIFS({clave},{col},"<0"{ex})'

    def valor(self, ref, lado, estado_doc):
        col, ex = self.rango["L" if lado == "sri" else "M"], self._extra(lado, estado_doc)
        return f'=SUMIFS({col},{self.rango["Q"]},"="&{ref}{ex})'


def _reactivar(ws, cols):
    """Vuelve formula las celdas de `cols` (las nuestras), que _hoja/_proteger dejaron como texto."""
    for rr in range(2, ws.max_row + 1):
        for col in cols:
            c = ws[f"{col}{rr}"]
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = c.value


def _hoja_conciliacion(wb, ret, n_docs):
    ws = wb.create_sheet(HOJA_CONC)
    ws.append(_ENC_CONC)
    _enc(ws, 1, len(_ENC_CONC), COLORES[HOJA_CONC])
    lim_d = max(n_docs, 1) + 1
    for i, r in enumerate((x for _, x in ret.iterrows()), start=2):
        ambos = _tiene(r["valor_sri"]) and _tiene(r["valor_odoo"])
        estado = (f'=IF(ABS(N{i})>{TOL},"{cr.EST_DIF}","{cr.EST_OK}")' if _es_formula(r) else r["estado"])
        dif = f'=ROUND(L{i}-M{i},2)' if ambos else ""
        est_doc = (f"=INDEX({HOJA_DOC}!$A$2:$A${lim_d},MATCH(Q{i},{HOJA_DOC}!$N$2:$N${lim_d},0))")
        ws.append([estado, _v(r["n_ret"]), _v(r["fecha"]), _v(r["ruc"]), _v(r["proveedor"]),
                   _v(r["factura_sri"]), _v(r["factura_odoo"]), _v(r["pct_sri"]), _v(r["pct_odoo"]),
                   _v(r["base_sri"]), _v(r["base_odoo"]), _v(r["valor_sri"]), _v(r["valor_odoo"]), dif,
                   _v(r["id_ret"]), _v(r["clave"]), r["_doc"], est_doc])
    n = len(ret)
    for rr in range(2, n + 2):
        ws[f"C{rr}"].number_format = "dd/mm/yyyy"
        for col in "JKLMN":
            ws[f"{col}{rr}"].number_format = "#,##0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_ENC_CONC))}{max(n, 1) + 1}"
    ws.sheet_properties.tabColor = COLORES[HOJA_CONC]
    for k, w in enumerate(_ANCH_CONC, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    _fuente(ws)
    _proteger(ws, exentas="ANR")  # Documento (Q) es texto de datos: se protege
    return ws


def _hoja_documentos(wb, docs, fx):
    ws = wb.create_sheet(HOJA_DOC)
    ws.append(_ENC_DOC)
    _enc(ws, 1, len(_ENC_DOC), COLORES[HOJA_DOC])
    for i, d in enumerate((x for _, x in docs.iterrows()), start=2):
        e, ref = d["estado"], f"N{i}"
        ws.append([e, _v(d["n_ret"]), _v(d["fecha"]), _v(d["ruc"]), _v(d["proveedor"]),
                   _v(d["factura_sri"]), _v(d["factura_odoo"]),
                   fx.lineas(ref, "sri", e), fx.lineas(ref, "odoo", e),
                   fx.valor(ref, "sri", e), fx.valor(ref, "odoo", e), f"=ROUND(J{i}-K{i},2)",
                   _v(d["detalle"]), d["doc"]])
    n = len(docs)
    for rr in range(2, n + 2):
        ws[f"C{rr}"].number_format = "dd/mm/yyyy"
        for col in "JKL":
            ws[f"{col}{rr}"].number_format = "#,##0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_ENC_DOC))}{max(n, 1) + 1}"
    ws.sheet_properties.tabColor = COLORES[HOJA_DOC]
    for k, w in enumerate(_ANCH_DOC, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    _fuente(ws)
    _proteger(ws, exentas="HIJKL")
    return ws


def _hoja_docs_estado(wb, titulo, enc, docs, armar, anchos, fmt, cols_formula):
    """Hoja de resultado con UNA fila por documento; `armar(d, i)` devuelve la fila i."""
    filas = [armar(d, i) for i, d in enumerate((x for _, x in docs.iterrows()), start=2)]
    ws = _hoja(wb, titulo, enc, filas, anchos, fmt)
    _reactivar(ws, cols_formula)
    return ws


def agregar_hojas_retenciones(wb, ret):
    """Crea las hojas Ret_* en `wb` y devuelve los contadores del resumen (por DOCUMENTO; `lineas` es
    el total de filas de detalle y total_sri/total_odoo suman las lineas, como antes)."""
    ret = ret.reset_index(drop=True)
    docs = cr.resumir_por_documento(ret)
    rango = {d: k for k, d in enumerate(docs["doc"])}
    claves = cr.documento_por_linea(ret)
    ret = ret.assign(_doc=claves, _rk=claves.map(rango))
    ret = ret.sort_values("_rk", kind="mergesort").reset_index(drop=True)  # lineas contiguas por documento
    fx = _Formulas(len(ret))

    _hoja_documentos(wb, docs, fx)
    _hoja_conciliacion(wb, ret, len(docs))
    est = docs["estado"]
    f_fecha = {"B": "dd/mm/yyyy"}

    _hoja_docs_estado(
        wb, "Ret_Solo_SRI",
        ["N_Retencion", "Fecha", "RUC", "Proveedor", "Factura", "Lineas", "Valor", "Clave", "Documento"],
        docs[est == cr.EST_SOLO_SRI],
        lambda d, i: [_v(d["n_ret"]), _v(d["fecha"]), _v(d["ruc"]), _v(d["proveedor"]), _v(d["factura_sri"]),
                      fx.lineas(f"I{i}", "sri", d["estado"]), fx.valor(f"I{i}", "sri", d["estado"]),
                      _v(d["clave"]), d["doc"]],
        [18, 12, 15, 42, 19, 8, 12, 52, 22], {**f_fecha, "G": "#,##0.00"}, "FG")
    _hoja_docs_estado(
        wb, "Ret_Solo_Odoo",
        ["N_Retencion", "Fecha", "RUC", "Proveedor", "Factura", "Lineas", "Valor", "ID_Odoo", "Documento"],
        docs[est == cr.EST_SOLO_ODOO],
        lambda d, i: [_v(d["n_ret"]), _v(d["fecha"]), _v(d["ruc"]), _v(d["proveedor"]), _v(d["factura_odoo"]),
                      fx.lineas(f"I{i}", "odoo", d["estado"]), fx.valor(f"I{i}", "odoo", d["estado"]),
                      _v(d["id_ret"]), d["doc"]],
        [18, 12, 15, 42, 19, 8, 12, 12, 22], {**f_fecha, "G": "#,##0.00"}, "FG")

    # Diferencia de valor: TODAS las lineas (no informativas) de los documentos con diferencia, agrupadas
    # por documento, con el estado propio de cada linea.
    docs_dif = set(docs.loc[est == cr.EST_DIF, "doc"])
    lin = ret[ret["_doc"].isin(docs_dif) & (ret["estado"] != cr.EST_INFO)]
    filas_dif = []
    for i, (_, r) in enumerate(lin.iterrows(), start=2):
        ambos = _tiene(r["valor_sri"]) and _tiene(r["valor_odoo"])
        fac = r["factura_sri"] if _tiene(r["factura_sri"]) else r["factura_odoo"]
        filas_dif.append([r["estado"], _v(r["n_ret"]), _v(r["fecha"]), _v(r["proveedor"]), _v(fac),
                          _v(r["pct_sri"]), _v(r["pct_odoo"]), _v(r["base_sri"]), _v(r["base_odoo"]),
                          _v(r["valor_sri"]), _v(r["valor_odoo"]),
                          f"=ROUND(J{i}-K{i},2)" if ambos else "", r["_doc"]])
    ws_dif = _hoja(wb, "Ret_Diferencia_Valor",
                   ["Estado_Linea", "N_Retencion", "Fecha", "Proveedor", "Factura", "Pct_SRI", "Pct_Odoo",
                    "Base_SRI", "Base_Odoo", "Valor_SRI", "Valor_Odoo", "Diferencia", "Documento"],
                   filas_dif, [24, 18, 12, 42, 19, 8, 8, 12, 12, 12, 12, 11, 22],
                   {"C": "dd/mm/yyyy", **{c: "#,##0.00" for c in "HIJKL"}})
    _reactivar(ws_dif, "L")

    _hoja_docs_estado(
        wb, "Ret_Factura_Incorrecta",
        ["N_Retencion", "Fecha", "RUC", "Proveedor", "Factura_SRI", "Factura_Odoo", "Lineas_SRI", "Lineas_Odoo",
         "Valor_SRI", "Valor_Odoo", "Documento"],
        docs[est == cr.EST_FAC],
        lambda d, i: [_v(d["n_ret"]), _v(d["fecha"]), _v(d["ruc"]), _v(d["proveedor"]), _v(d["factura_sri"]),
                      _v(d["factura_odoo"]), fx.lineas(f"K{i}", "sri", d["estado"]),
                      fx.lineas(f"K{i}", "odoo", d["estado"]), fx.valor(f"K{i}", "sri", d["estado"]),
                      fx.valor(f"K{i}", "odoo", d["estado"]), d["doc"]],
        [18, 12, 15, 42, 19, 19, 10, 10, 12, 12, 22], {**f_fecha, "I": "#,##0.00", "J": "#,##0.00"}, "GHIJ")
    _hoja_docs_estado(
        wb, "Ret_Informativos",
        ["Estado", "N_Retencion", "Fecha", "RUC", "Proveedor", "Factura_Odoo", "Lineas", "Valor", "ID_Odoo",
         "Clave", "Documento"],
        docs[est.isin([cr.EST_XML, cr.EST_INFO])],
        lambda d, i: [d["estado"], _v(d["n_ret"]), _v(d["fecha"]), _v(d["ruc"]), _v(d["proveedor"]),
                      _v(d["factura_odoo"]), fx.lineas(f"K{i}", "odoo", d["estado"]),
                      fx.valor(f"K{i}", "odoo", d["estado"]), _v(d["id_ret"]), _v(d["clave"]), d["doc"]],
        [34, 18, 12, 15, 42, 19, 8, 12, 12, 52, 22], {"C": "dd/mm/yyyy", "H": "#,##0.00"}, "GH")

    sin_info = ret[ret["estado"] != cr.EST_INFO]
    return {
        "coinciden": int((est == cr.EST_OK).sum()),
        "diferencia": int((est == cr.EST_DIF).sum()),
        "solo_sri": int((est == cr.EST_SOLO_SRI).sum()),
        "solo_odoo": int((est == cr.EST_SOLO_ODOO).sum()),
        "factura_incorrecta": int((est == cr.EST_FAC).sum()),
        "informativos": int(est.isin([cr.EST_XML, cr.EST_INFO]).sum()),
        "total_sri": round(float(ret["valor_sri"].sum()), 2),
        "total_odoo": round(float(sin_info["valor_odoo"].sum()), 2),
        "documentos": int(len(docs)),
        "lineas": int(len(ret)),
    }


def _texto_periodo(desde, hasta):
    """'MM/AAAA' si es un mes calendario completo; si no, 'dd/mm/aaaa a dd/mm/aaaa'."""
    d, h = pd.Timestamp(desde), pd.Timestamp(hasta)
    if d.day == 1 and h == d + pd.offsets.MonthEnd(0):
        return f"{d:%m/%Y}"
    return f"{d:%d/%m/%Y} a {h:%d/%m/%Y}"


def fila_resumen_retenciones(rs, r, ret_n, periodo=None, periodo_compras=None, doc_n=None):
    """Escribe en la hoja Resumen (fila `r`) la fila 'Retenciones' con COUNTIFS sobre Ret_Documentos
    (cuenta DOCUMENTOS, no lineas). Mismas columnas que la tabla por tipo. Devuelve la fila siguiente libre.

    ret_n: lineas de Ret_Conciliacion; doc_n: filas de Ret_Documentos (por defecto ret_n).
    periodo: (desde, hasta) de las retenciones; va en la columna Mes y, si difiere de
    periodo_compras, se rotula en la nota bajo la fila."""
    doc_n = ret_n if doc_n is None else doc_n
    lim, lim_l = max(doc_n, 1) + 1, max(ret_n, 1) + 1
    est, lsri, vs, vo = (f"{HOJA_DOC}!${c}$2:${c}${lim}" for c in ("A", "H", "J", "K"))
    rs[f"A{r}"] = "Retenciones"
    nota_periodo = ""
    if periodo is not None:
        d, h = (pd.Timestamp(x) for x in periodo)
        rs[f"B{r}"] = (f"{d:%Y-%m}" if _texto_periodo(d, h) == f"{d:%m/%Y}"
                       else _texto_periodo(d, h))
        if periodo_compras is None or tuple(pd.Timestamp(x) for x in periodo_compras) != (d, h):
            nota_periodo = f'&"  |  Período de retenciones: {_texto_periodo(d, h)}"'
    rs[f"C{r}"] = f'=COUNTIFS({lsri},">0")'
    for col, txt in zip("DEFG", (cr.EST_OK, cr.EST_DIF, cr.EST_SOLO_SRI, cr.EST_SOLO_ODOO)):
        rs[f"{col}{r}"] = f'=COUNTIFS({est},"{txt}")'
    rs[f"H{r}"] = 0
    rs[f"I{r}"] = f'=COUNTIFS({est},"{cr.EST_INFO}")+COUNTIFS({est},"{cr.EST_XML}")'
    rs[f"J{r}"] = f"=IF(C{r}=0,0,(D{r}+E{r})/C{r})"
    rs[f"K{r}"] = f"=SUM({vs})"
    rs[f"L{r}"] = f'=SUMIFS({vo},{est},"<>{cr.EST_INFO}")'
    rs[f"M{r}"] = f"=K{r}-L{r}"
    rs[f"J{r}"].number_format = "0.0%"
    for col in "KLM":
        rs[f"{col}{r}"].number_format = "#,##0.00"
    rs[f"A{r + 1}"] = (f'="Retenciones (documentos) - factura incorrecta en Odoo: "&COUNTIFS({est},"{cr.EST_FAC}")'
                       f'&"  |  XML no disponible: "&COUNTIFS({est},"{cr.EST_XML}")'
                       f'&"  |  Detalle por línea: "&COUNTA({HOJA_CONC}!$Q$2:$Q${lim_l})&" líneas"'
                       + nota_periodo)
    rs[f"A{r + 1}"].font = Font(name="Arial", italic=True, size=9)
    return r + 2
