# -*- coding: utf-8 -*-
"""Libro Excel de la conciliacion de comprobantes recibidos (FAC/NC/ND) SRI vs Odoo."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from conciliar_compras import CAT_DUP, CAT_INFO, CAT_SOLO_ODOO, TIPOS
from conciliar_recibidas import FONT, _enc, _fuente

TOL = 0.02
COLORES = {"Resumen": "1F4E79", "Conciliacion": "1F4E79", "Solo_en_SRI": "ED7D31",
           "Solo_en_Odoo": "C00000", "Diferencia_Valor": "FFC000", "Duplicados_Odoo": "7030A0",
           "Informativos": "808080",
           "Ret_Conciliacion": "1F4E79", "Ret_Solo_SRI": "ED7D31", "Ret_Solo_Odoo": "C00000",
           "Ret_Diferencia_Valor": "FFC000", "Ret_Factura_Incorrecta": "7030A0", "Ret_Informativos": "808080",
           "SRI_Listado": "808080", "Odoo_Lineas": "808080"}
_PELIGROSOS = ("=", "+", "-", "@")
FUERA = "Fecha fuera del período"


def _nan(v):
    return "" if pd.isna(v) else v


def _fecha(v):
    return pd.Timestamp(v).to_pydatetime()


def _proteger(ws, exentas=""):
    """Texto de Odoo/SRI que empieza con = + - @ nunca se guarda como fórmula (commit 96f5249)."""
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column_letter not in exentas and isinstance(c.value, str) and c.value.startswith(_PELIGROSOS):
                c.data_type = "s"


def _hoja(wb, titulo, enc, filas, anchos, fmt=None):
    ws = wb.create_sheet(titulo)
    ws.append(enc)
    _enc(ws, 1, len(enc), COLORES[titulo])
    for f in filas:
        ws.append(f)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(enc))}{max(ws.max_row, 2)}"
    ws.sheet_properties.tabColor = COLORES[titulo]
    for col, nf in (fmt or {}).items():
        for r in range(2, ws.max_row + 1):
            ws[f"{col}{r}"].number_format = nf
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _fuente(ws)
    _proteger(ws)
    return ws


def construir_excel(sri, odoo, pares, desde, hasta, salida, ret=None):
    wb = Workbook()
    wb.remove(wb.active)
    ns, no = len(sri), len(odoo)
    od = odoo.set_index("id")
    sri_n, odo_n = max(ns, 1) + 1, max(no, 1) + 1

    # ---- pares enriquecidos (orden: fecha, tipo) ----
    p = pares.copy().reset_index(drop=True)
    p["_ord"] = [sri.at[int(i), "fecha"] if pd.notna(i) else od.at[int(j), "fecha"]
                 for i, j in zip(p["idx_sri"], p["id_odoo"])]
    p["mes"] = [pd.Timestamp(d).strftime("%Y-%m") for d in p["_ord"]]
    p["tot_sri"] = [sri.at[int(i), "total"] if pd.notna(i) else None for i in p["idx_sri"]]
    p["tot_od"] = [od.at[int(j), "total"] if pd.notna(j) else None for j in p["id_odoo"]]
    ambos = (p["idx_sri"].notna() & p["id_odoo"].notna()).tolist()
    p["dif"] = [round(float(a) - float(b), 2) if m else None for a, b, m in zip(p["tot_sri"], p["tot_od"], ambos)]
    no_dup = (p["cat"] != CAT_DUP).tolist()
    p["es_dif"] = [m and nd and abs(d) > TOL for m, nd, d in zip(ambos, no_dup, p["dif"].fillna(0))]
    p["es_ok"] = [m and nd and not df for m, nd, df in zip(ambos, no_dup, p["es_dif"])]
    p = p.sort_values(["_ord", "tipo"], kind="stable").reset_index(drop=True)

    # ---- hojas de datos ----
    _hoja(wb, "SRI_Listado",
          ["Clave", "Tipo", "Serie/Número", "RUC", "Proveedor", "Fecha", "Mes", "Base", "IVA", "Total"],
          [[r.clave, r.tipo, r.serie, r.ruc, r.proveedor, _fecha(r.fecha), pd.Timestamp(r.fecha).strftime("%Y-%m"),
            r.base, r.iva, r.total] for r in sri.itertuples()],
          [52, 7, 19, 15, 45, 12, 9, 12, 12, 12],
          {"F": "dd/mm/yyyy", "H": "#,##0.00", "I": "#,##0.00", "J": "#,##0.00"})
    _hoja(wb, "Odoo_Lineas",
          ["ID_Odoo", "Numero_Odoo", "Tipo", "Clave", "Serie/Número", "RUC", "Proveedor", "Fecha", "Base", "IVA", "Total"],
          [[r.id, r.num, r.tipo, _nan(r.clave), _nan(r.nfac), _nan(r.ruc), r.socio, _fecha(r.fecha),
            r.base, r.iva, r.total] for r in odoo.itertuples()],
          [9, 20, 7, 52, 19, 15, 45, 12, 12, 12, 12],
          {"H": "dd/mm/yyyy", "I": "#,##0.00", "J": "#,##0.00", "K": "#,##0.00"})

    sa, so = f"SRI_Listado!$A$2:$A${sri_n}", f"Odoo_Lineas!$A$2:$A${odo_n}"

    def idx(col, hoja, n, pos):
        return f"INDEX({hoja}!${col}$2:${col}${n},{pos})"

    # ---- Conciliacion ----
    enc = ["Tipo", "Mes", "Clave_SRI", "ID_Odoo", "Fecha_SRI", "Fecha_Odoo", "RUC", "Proveedor", "Serie/Número",
           "Numero_Odoo", "Metodo_Cruce", "Total_SRI", "Total_Odoo", "Diferencia", "Resultado",
           "Categoria_Odoo", "Observacion", "_fila_SRI", "_fila_Odoo"]
    ws = wb.create_sheet("Conciliacion", 0)
    ws.append(enc)
    _enc(ws, 1, len(enc), COLORES["Conciliacion"])
    d0, d1 = desde, hasta
    for i, r in enumerate(p.itertuples(), start=2):
        clave = sri.at[int(r.idx_sri), "clave"] if pd.notna(r.idx_sri) else ""
        idod = int(r.id_odoo) if pd.notna(r.id_odoo) else ""
        ws.append([
            r.tipo, r.mes, clave, idod,
            f'=IF(R{i}="","",{idx("F", "SRI_Listado", sri_n, f"R{i}")})',
            f'=IF(S{i}="","",{idx("H", "Odoo_Lineas", odo_n, f"S{i}")})',
            f'=IF(R{i}<>"",{idx("D", "SRI_Listado", sri_n, f"R{i}")},{idx("F", "Odoo_Lineas", odo_n, f"S{i}")})&""',
            f'=IF(R{i}<>"",{idx("E", "SRI_Listado", sri_n, f"R{i}")},{idx("G", "Odoo_Lineas", odo_n, f"S{i}")})&""',
            f'=IF(R{i}<>"",{idx("C", "SRI_Listado", sri_n, f"R{i}")},{idx("E", "Odoo_Lineas", odo_n, f"S{i}")})&""',
            f'=IF(S{i}="","",{idx("B", "Odoo_Lineas", odo_n, f"S{i}")}&"")',
            r.metodo,
            f'=IF(R{i}="","",{idx("J", "SRI_Listado", sri_n, f"R{i}")})',
            f'=IF(S{i}="","",{idx("K", "Odoo_Lineas", odo_n, f"S{i}")})',
            f'=IF(AND(R{i}<>"",S{i}<>""),ROUND(L{i}-M{i},2),"")',
            f'=IF(P{i}="{CAT_DUP}",P{i},IF(AND(R{i}<>"",S{i}<>""),IF(ABS(N{i})>{TOL},"Diferencia de valor","Coincide"),'
            f'IF(R{i}<>"","Solo en SRI",P{i})))',
            r.cat,
            f'=IF(AND(R{i}<>"",S{i}<>"",OR(F{i}<DATE({d0.year},{d0.month},{d0.day}),'
            f'F{i}>DATE({d1.year},{d1.month},{d1.day}))),"{FUERA}","")',
            f'=IF(C{i}="","",MATCH(C{i},{sa},0))',
            f'=IF(D{i}="","",MATCH(D{i},{so},0))',
        ])
    n = len(p)
    for r in range(2, n + 2):
        ws[f"E{r}"].number_format = ws[f"F{r}"].number_format = "dd/mm/yyyy"
        for col in "LMN":
            ws[f"{col}{r}"].number_format = "#,##0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:S{max(n, 1) + 1}"
    ws.sheet_properties.tabColor = COLORES["Conciliacion"]
    for i, w in enumerate([7, 9, 52, 9, 12, 12, 15, 42, 19, 20, 28, 12, 12, 11, 22, 26, 22, 9, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["R"].hidden = ws.column_dimensions["S"].hidden = True
    _fuente(ws)
    _proteger(ws, exentas="EFGHIJLMNOQRS")

    # ---- hojas de resultado ----
    enc_d = ["Tipo", "Clave", "Serie/Número", "RUC", "Proveedor", "Fecha", "Base", "IVA", "Total"]
    fmt_d = {"F": "dd/mm/yyyy", "G": "#,##0.00", "H": "#,##0.00", "I": "#,##0.00"}
    anch_d = [7, 52, 19, 15, 45, 12, 12, 12, 12]

    sol_sri = p[p["idx_sri"].notna() & p["id_odoo"].isna()]
    _hoja(wb, "Solo_en_SRI", enc_d,
          [[sri.at[k, "tipo"], sri.at[k, "clave"], sri.at[k, "serie"], sri.at[k, "ruc"], sri.at[k, "proveedor"],
            _fecha(sri.at[k, "fecha"]), sri.at[k, "base"], sri.at[k, "iva"], sri.at[k, "total"]]
           for k in (int(i) for i in sol_sri["idx_sri"])],
          anch_d, fmt_d)
    enc_o = ["Categoria", "Numero_Odoo"] + enc_d
    fmt_o = {"H": "dd/mm/yyyy", "I": "#,##0.00", "J": "#,##0.00", "K": "#,##0.00"}
    anch_o = [28, 20] + anch_d

    def filas_cat(cat):
        sel = p[(p["cat"] == cat) & p["id_odoo"].notna()]
        out = []
        for j in (int(x) for x in sel["id_odoo"]):
            o = od.loc[j]
            out.append([cat, o["num"], o["tipo"], _nan(o["clave"]), _nan(o["nfac"]), _nan(o["ruc"]), o["socio"],
                        _fecha(o["fecha"]), o["base"], o["iva"], o["total"]])
        return out

    _hoja(wb, "Solo_en_Odoo", enc_o, filas_cat(CAT_SOLO_ODOO), anch_o, fmt_o)
    dif = p[p["es_dif"]]
    _hoja(wb, "Diferencia_Valor",
          ["Tipo", "Fecha_SRI", "Proveedor", "Serie/Número", "Numero_Odoo", "Metodo_Cruce", "Total_SRI",
           "Total_Odoo", "Diferencia"],
          [[r.tipo, _fecha(sri.at[int(r.idx_sri), "fecha"]), sri.at[int(r.idx_sri), "proveedor"],
            sri.at[int(r.idx_sri), "serie"], od.at[int(r.id_odoo), "num"], r.metodo, r.tot_sri, r.tot_od, r.dif]
           for r in dif.itertuples()],
          [7, 12, 45, 19, 20, 28, 12, 12, 11],
          {"B": "dd/mm/yyyy", "G": "#,##0.00", "H": "#,##0.00", "I": "#,##0.00"})
    dup_idx = set(p.loc[p["cat"] == CAT_DUP, "idx_sri"])
    dd = p[p["idx_sri"].isin(dup_idx) & p["id_odoo"].notna()].sort_values(["idx_sri", "_ord"], kind="stable")
    _hoja(wb, "Duplicados_Odoo",
          ["Tipo", "Serie/Número_SRI", "Proveedor", "Numero_Odoo", "Serie/Número_Odoo", "Fecha_Odoo", "Total_Odoo"],
          [[r.tipo, sri.at[int(r.idx_sri), "serie"], sri.at[int(r.idx_sri), "proveedor"], od.at[int(r.id_odoo), "num"],
            _nan(od.at[int(r.id_odoo), "nfac"]), _fecha(od.at[int(r.id_odoo), "fecha"]), od.at[int(r.id_odoo), "total"]]
           for r in dd.itertuples()],
          [7, 19, 45, 20, 19, 12, 12], {"F": "dd/mm/yyyy", "G": "#,##0.00"})
    _hoja(wb, "Informativos", enc_o, filas_cat(CAT_INFO), anch_o, fmt_o)

    # ---- Resumen ----
    rs = wb.create_sheet("Resumen", 0)
    rs["A1"] = "Conciliación de comprobantes RECIBIDOS (FAC / NC / ND): SRI vs Odoo"
    rs["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E79")
    rs["A2"] = (f"Período {desde:%d/%m/%Y} a {hasta:%d/%m/%Y}  |  Mes = mes de emisión (SRI) o fecha del "
                "documento (Odoo si no está en SRI)")
    rs["A2"].font = Font(name=FONT, italic=True, size=9)
    e = ["Tipo", "Mes", "Comprobantes SRI", "Coinciden", "Diferencia de valor", "Solo en SRI", "Solo en Odoo",
         "Duplicado en Odoo", "Informativos", "% Coincidencia", "Total SRI", "Total Odoo", "Dif. total"]
    for j, h in enumerate(e, 1):
        rs.cell(row=4, column=j, value=h)
    _enc(rs, 4, len(e), COLORES["Resumen"])
    lim = max(n, 1) + 1
    ct, cm, cr, cl = (f"Conciliacion!${c}$2:${c}${lim}" for c in ("A", "B", "O", "M"))
    st, sm, sj = (f"SRI_Listado!${c}$2:${c}${sri_n}" for c in ("B", "G", "J"))
    r = 5
    for t in TIPOS:
        meses = sorted(set(p.loc[p["tipo"] == t, "mes"]))
        for m in meses:
            rs[f"A{r}"], rs[f"B{r}"] = t, m
            rs[f"C{r}"] = f"=COUNTIFS({st},$A{r},{sm},$B{r})"
            for col, txt in zip("DEFGHI", ("Coincide", "Diferencia de valor", "Solo en SRI", CAT_SOLO_ODOO,
                                           CAT_DUP, CAT_INFO)):
                rs[f"{col}{r}"] = f'=COUNTIFS({ct},$A{r},{cm},$B{r},{cr},"{txt}")'
            rs[f"J{r}"] = f"=IF(C{r}=0,0,(D{r}+E{r})/C{r})"
            rs[f"K{r}"] = f"=SUMIFS({sj},{st},$A{r},{sm},$B{r})"
            rs[f"L{r}"] = "=" + "+".join(
                f'SUMIFS({cl},{ct},$A{r},{cm},$B{r},{cr},"{x}")'
                for x in ("Coincide", "Diferencia de valor", CAT_SOLO_ODOO, CAT_DUP))
            rs[f"M{r}"] = f"=K{r}-L{r}"
            r += 1
    ult, tot = r - 1, r
    rs[f"A{tot}"] = "Total"
    for col in "CDEFGHIKLM":
        rs[f"{col}{tot}"] = f"=SUM({col}5:{col}{ult})" if ult >= 5 else 0
    rs[f"J{tot}"] = f"=IF(C{tot}=0,0,(D{tot}+E{tot})/C{tot})"
    for rr in range(5, tot + 1):
        rs[f"J{rr}"].number_format = "0.0%"
        for col in "KLM":
            rs[f"{col}{rr}"].number_format = "#,##0.00"
    contadores_ret, sig = None, tot + 2
    if ret is not None:
        from excel_retenciones import agregar_hojas_retenciones, fila_resumen_retenciones
        contadores_ret = agregar_hojas_retenciones(wb, ret)
        sig = fila_resumen_retenciones(rs, tot + 1, len(ret)) + 1
    _fuente(rs, 4)
    for col in "ABCDEFGHIJKLM":
        rs[f"{col}{tot}"].font = Font(name=FONT, bold=True)
        rs[f"{col}{tot}"].fill = PatternFill("solid", fgColor="DDEBF7")
    notas = [
        "Observaciones",
        "Coinciden = comprobante en SRI y Odoo con total igual (tolerancia $0.02). Duplicado en Odoo = el mismo "
        "comprobante del SRI registrado más de una vez (hoja Duplicados_Odoo).",
        "Solo en SRI = autorizado en el SRI sin registro en Odoo. Solo en Odoo = registrado en Odoo y ausente en el SRI.",
        "Informativos = documentos de Odoo que no aplican a la conciliación electrónica (hoja Informativos).",
        f"{FUERA} (columna Observacion de Conciliacion) marca documentos de Odoo emparejados con fecha fuera del período.",
    ]
    for k, tx in enumerate(notas):
        c = rs[f"A{sig + k}"]
        c.value = tx
        c.font = Font(name=FONT, bold=(k == 0), size=9 if k else 10)
    rs.freeze_panes = "C5"
    rs.sheet_properties.tabColor = COLORES["Resumen"]
    for i, w in enumerate([7, 10, 14, 12, 13, 12, 12, 13, 13, 13, 15, 15, 13], 1):
        rs.column_dimensions[get_column_letter(i)].width = w
    rs.row_dimensions[4].height = 42

    orden = list(COLORES)
    wb._sheets.sort(key=lambda w: orden.index(w.title))
    wb.active = 0
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(salida)

    resumen = {}
    for t in TIPOS:
        q = p[p["tipo"] == t]
        resumen[t] = {
            "coinciden": int(q["es_ok"].sum()),
            "solo_sri": int((q["idx_sri"].notna() & q["id_odoo"].isna()).sum()),
            "solo_odoo": int((q["cat"] == CAT_SOLO_ODOO).sum()),
            "diferencias": int(q["es_dif"].sum()),
            "informativos": int((q["cat"] == CAT_INFO).sum()),
        }
    if contadores_ret is not None:
        resumen["retenciones"] = contadores_ret
    return resumen
