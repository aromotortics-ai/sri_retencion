# -*- coding: utf-8 -*-
"""Libro Excel de la conciliación de comprobantes EMITIDOS (FAC/NC/ND/LIQ y retenciones) SRI vs Odoo.

Misma estructura y estilo que `excel_compras.construir_excel`: el Estado de cada documento es el valor
calculado (y testeado) por `conciliar_emitidos.clasificar_documentos`; conteos, sumas y diferencias son
fórmulas. Las retenciones reutilizan las hojas Ret_* de `excel_retenciones`.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

import conciliar_retenciones as cr
from conciliar_compras import _recalcular
from conciliar_emitidos import (EST_ANULADO, EST_DIF_DATOS, EST_DIF_VALOR, EST_DUP, EST_INFO, EST_OK,
                                EST_SOLO_ODOO, EST_SOLO_SRI)
from conciliar_recibidas import FONT, _enc, _fuente
from emitidos_sri import TIPOS_DOC
from excel_compras import COLORES, _fecha, _hoja, _proteger
from excel_retenciones import _reactivar, _texto_periodo, agregar_hojas_retenciones, fila_resumen_retenciones

ESTADOS_DOC = (EST_OK, EST_DIF_VALOR, EST_DIF_DATOS, EST_SOLO_SRI, EST_SOLO_ODOO, EST_ANULADO, EST_DUP,
               EST_INFO)
ESTADOS_RET = (cr.EST_OK, cr.EST_DIF, cr.EST_SOLO_SRI, cr.EST_SOLO_ODOO, cr.EST_FAC, cr.EST_ANULADO,
               cr.EST_XML, cr.EST_INFO)
NOMBRES = {"FAC": "Facturas", "NC": "Notas de crédito", "ND": "Notas de débito", "LIQ": "Liquidaciones"}
ENC = ["Tipo", "Estado", "Número", "Fecha_SRI", "Fecha_Odoo", "Cliente", "Total_SRI", "Total_Odoo",
       "Diferencia", "Detalle", "Numero_Odoo", "Metodo_Cruce", "Clave_SRI", "ID_Odoo"]
ANCHOS = [7, 26, 19, 12, 12, 42, 12, 12, 11, 48, 20, 24, 52, 9]
FMT = {"D": "dd/mm/yyyy", "E": "dd/mm/yyyy", "G": "#,##0.00", "H": "#,##0.00", "I": "#,##0.00"}
ORDEN = ["Resumen", "Conciliacion", "Solo_en_SRI", "Solo_en_Odoo", "Diferencias", "Anulados_Odoo",
         "Ret_Documentos", "Ret_Conciliacion", "Ret_Solo_SRI", "Ret_Solo_Odoo", "Ret_Diferencia_Valor",
         "Ret_Factura_Incorrecta", "Ret_Informativos"]


def _o(v):
    """Celda: NA/NaN -> None (celda vacía, no cuenta en COUNTIFS "<>"); numpy -> nativo."""
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    return v.item() if hasattr(v, "item") else v


def _total(v):
    """Totales en valor absoluto (NC)."""
    v = _o(v)
    return None if v is None else round(abs(float(v)), 2)


def _filas_documentos(sri, odoo, docs):
    """Una fila (sin la fórmula de Diferencia) por documento, ordenadas por tipo y fecha."""
    od = odoo.set_index("id")
    filas = []
    for d in docs.itertuples(index=False):
        s = sri.loc[int(d.idx_sri)] if pd.notna(d.idx_sri) else None
        o = od.loc[int(d.id_odoo)] if pd.notna(d.id_odoo) else None
        numero = _o(s["serie"]) if s is not None else None
        if numero is None and o is not None:
            numero = _o(o["nfac"]) or _o(o["num"])
        cliente = _o(o["socio"]) if o is not None else None
        if cliente is None and s is not None:
            cliente = _o(s["proveedor"])
        f_sri = _fecha(s["fecha"]) if s is not None and pd.notna(s["fecha"]) else None
        f_odoo = _fecha(o["fecha"]) if o is not None and pd.notna(o["fecha"]) else None
        filas.append({
            "_orden": (TIPOS_DOC.index(d.tipo) if d.tipo in TIPOS_DOC else len(TIPOS_DOC),
                       f_sri or f_odoo or pd.Timestamp.max.to_pydatetime()),
            "fila": [d.tipo, d.estado, numero, f_sri, f_odoo, cliente,
                     _total(s["total"]) if s is not None else None,
                     _total(o["total"]) if o is not None else None,
                     None, _o(d.detalle), _o(o["num"]) if o is not None else None, _o(d.metodo),
                     _o(s["clave"]) if s is not None else None,
                     int(d.id_odoo) if pd.notna(d.id_odoo) else None],
        })
    filas.sort(key=lambda x: x["_orden"])
    return [x["fila"] for x in filas]


def _con_diferencia(filas):
    """Copia las filas poniendo Diferencia (columna I) = fórmula cuando hay ambos totales."""
    out = []
    for i, f in enumerate(filas, start=2):
        f = list(f)
        f[8] = f"=ROUND(G{i}-H{i},2)" if f[6] is not None and f[7] is not None else None
        out.append(f)
    return out


def _hoja_filtrada(wb, titulo, filas, estados):
    ws = _hoja(wb, titulo, ENC, _con_diferencia([f for f in filas if f[1] in estados]), ANCHOS, FMT)
    _reactivar(ws, "I")  # _hoja/_proteger dejan como texto todo lo que empieza con "="
    return ws


def _hoja_conciliacion(wb, filas):
    ws = _hoja(wb, "Conciliacion", ENC, _con_diferencia(filas), ANCHOS, FMT)
    _reactivar(ws, "I")
    ws.freeze_panes = "D2"
    ws.column_dimensions["N"].hidden = True
    return ws


def _resumen(wb, n, desde, hasta, ret, advertencias):
    rs = wb.create_sheet("Resumen", 0)
    rs["A1"] = "Conciliación de comprobantes EMITIDOS (FAC / NC / ND / LIQ / RET): SRI vs Odoo"
    rs["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E79")
    rs["A2"] = (f"Período {desde:%d/%m/%Y} a {hasta:%d/%m/%Y}  |  Una fila por documento en Conciliacion; "
                "las retenciones se cuentan por documento (Ret_Documentos)")
    rs["A2"].font = Font(name=FONT, italic=True, size=9)
    e = ["Tipo", "Período", "Comprobantes SRI", "Coinciden", "Diferencia de valor", "Solo en SRI",
         "Solo en Odoo", "Duplicado en Odoo", "Informativos", "% Coincidencia", "Total SRI", "Total Odoo",
         "Dif. total", "Diferencia de datos", "Anulado en Odoo"]
    for j, h in enumerate(e, 1):
        rs.cell(row=4, column=j, value=h)
    _enc(rs, 4, len(e), COLORES["Resumen"])
    lim = max(n, 1) + 1
    ct, ce, cs, co = (f"Conciliacion!${c}$2:${c}${lim}" for c in ("A", "B", "G", "H"))
    # Comprobantes del SRI: sin las filas que repiten un documento ya contado (duplicado en Odoo o
    # duplicado anulado, que queda Informativo con lado SRI).
    no_rep = f'{ce},"<>{EST_DUP}",{ce},"<>{EST_INFO}"'
    periodo = _texto_periodo(desde, hasta)  # mismo rótulo que la fila de retenciones
    if periodo == f"{desde:%m/%Y}":
        periodo = f"{desde:%Y-%m}"
    r = 5
    for t in TIPOS_DOC:
        rs[f"A{r}"], rs[f"B{r}"] = t, periodo
        rs[f"C{r}"] = f'=COUNTIFS({ct},$A{r},{cs},"<>",{no_rep})'
        for col, est in zip("DEFGHINO", (EST_OK, EST_DIF_VALOR, EST_SOLO_SRI, EST_SOLO_ODOO, EST_DUP, EST_INFO,
                                         EST_DIF_DATOS, EST_ANULADO)):
            rs[f"{col}{r}"] = f'=COUNTIFS({ct},$A{r},{ce},"{est}")'
        rs[f"J{r}"] = f"=IF(C{r}=0,0,(D{r}+E{r}+N{r})/C{r})"
        rs[f"K{r}"] = f"=SUMIFS({cs},{ct},$A{r},{no_rep})"
        rs[f"L{r}"] = "=" + "+".join(f'SUMIFS({co},{ct},$A{r},{ce},"{x}")'
                                     for x in (EST_OK, EST_DIF_VALOR, EST_DIF_DATOS, EST_SOLO_ODOO))
        rs[f"M{r}"] = f"=K{r}-L{r}"
        r += 1
    ult, tot = r - 1, r
    rs[f"A{tot}"] = "Total"
    for col in "CDEFGHIKLMNO":
        rs[f"{col}{tot}"] = f"=SUM({col}5:{col}{ult})"
    rs[f"J{tot}"] = f"=IF(C{tot}=0,0,(D{tot}+E{tot}+N{tot})/C{tot})"
    for rr in range(5, tot + 1):
        rs[f"J{rr}"].number_format = "0.0%"
        for col in "KLM":
            rs[f"{col}{rr}"].number_format = "#,##0.00"
    sig = tot + 2
    if ret is not None:
        contadores = agregar_hojas_retenciones(wb, ret)
        rr = tot + 1
        sig = fila_resumen_retenciones(rs, rr, len(ret), periodo=(desde, hasta), periodo_compras=(desde, hasta),
                                       doc_n=contadores["documentos"]) + 1
        lim_d = max(contadores["documentos"], 1) + 1
        rs[f"N{rr}"] = 0  # la diferencia de datos no aplica a retenciones
        rs[f"O{rr}"] = f'=COUNTIFS(Ret_Documentos!$A$2:$A${lim_d},"{cr.EST_ANULADO}")'
    _fuente(rs, 4)
    for col in "ABCDEFGHIJKLMNO":
        rs[f"{col}{tot}"].font = Font(name=FONT, bold=True)
        rs[f"{col}{tot}"].fill = PatternFill("solid", fgColor="DDEBF7")
    notas = [
        "Observaciones",
        "Coinciden = documento en SRI y Odoo con mismo total, base, clave y fecha (tolerancia $0.02). "
        "Diferencia de valor = total o base distintos. Diferencia de datos = mismos valores pero clave de acceso, "
        "fecha o marca electrónica distintas.",
        "Solo en SRI = autorizado en el SRI sin registro en Odoo. Solo en Odoo = publicado en Odoo y ausente en el SRI. "
        "Anulado en Odoo = cancelado en Odoo pero autorizado en el SRI (hoja Anulados_Odoo).",
        "Informativos = registros de Odoo en borrador, cancelados sin autorización o no electrónicos. "
        "Las retenciones usan sus propios estados (hojas Ret_*).",
    ]
    if advertencias:
        notas += ["Advertencias"] + list(advertencias)
    for k, tx in enumerate(notas):
        c = rs[f"A{sig + k}"]
        c.value = tx
        c.font = Font(name=FONT, bold=tx in ("Observaciones", "Advertencias"), size=10 if k == 0 else 9)
    _proteger(rs)  # notas/advertencias son texto (pueden traer datos del TXT)
    for row in rs.iter_rows(min_row=5, max_row=sig - 1):
        for c in row:  # las fórmulas del resumen sí son fórmulas
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = c.value
    rs.freeze_panes = "C5"
    rs.sheet_properties.tabColor = COLORES["Resumen"]
    for i, w in enumerate([12, 12, 14, 12, 13, 12, 12, 13, 13, 13, 15, 15, 13, 13, 13], 1):
        rs.column_dimensions[get_column_letter(i)].width = w
    rs.row_dimensions[4].height = 42
    return rs


def _conteos(estados, posibles):
    vc = pd.Series(list(estados), dtype=object).value_counts()
    out = {"total": int(len(estados))}
    out.update({e: int(vc.get(e, 0)) for e in posibles})
    return out


def construir_excel_emitidos(sri, odoo, docs, ret, desde, hasta, salida, advertencias=None):
    """Escribe el libro de emitidos en `salida` y devuelve {"documentos": {tipo: {"total", estado: n}},
    "recalculado": bool}. `docs` = `clasificar_documentos(sri, odoo, desde, hasta)`; `ret` = líneas de
    `conciliar_retenciones_emitidas` o None (sin TXT de retenciones: sin hojas Ret_* ni fila RET)."""
    wb = Workbook()
    wb.remove(wb.active)
    filas = _filas_documentos(sri, odoo, docs)
    _hoja_conciliacion(wb, filas)
    _hoja_filtrada(wb, "Solo_en_SRI", filas, (EST_SOLO_SRI,))
    _hoja_filtrada(wb, "Solo_en_Odoo", filas, (EST_SOLO_ODOO,))
    _hoja_filtrada(wb, "Diferencias", filas, (EST_DIF_VALOR, EST_DIF_DATOS))
    _hoja_filtrada(wb, "Anulados_Odoo", filas, (EST_ANULADO,))
    if ret is not None:
        ret = ret.reset_index(drop=True)
    _resumen(wb, len(filas), desde, hasta, ret, advertencias)

    wb._sheets.sort(key=lambda w: ORDEN.index(w.title))
    wb.active = 0
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(salida)

    documentos = {t: _conteos(docs.loc[docs["tipo"] == t, "estado"], ESTADOS_DOC) for t in TIPOS_DOC}
    if ret is not None:
        documentos["RET"] = _conteos(cr.resumir_por_documento(ret)["estado"] if len(ret) else [], ESTADOS_RET)
    return {"documentos": documentos, "recalculado": _recalcular(salida)}
