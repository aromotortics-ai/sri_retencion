# -*- coding: utf-8 -*-
"""
Conciliación de facturas de VENTA: Odoo vs SRI (comprobantes emitidos autorizados).

Entradas:
  - Excel exportado de Odoo (facturas de cliente): columnas Número, Número de
    autorización, Nombre del socio, Sucursal, Fecha de Factura/Recibo, Fecha de
    autorización, Importe sin impuestos con signo, Total con signo, Estado,
    Estado electrónico.
  - Carpeta con los TXT diarios del portal SRI ("Comprobantes electrónicos
    emitidos" -> Descargar reporte): emitidos_AAAA-MM-DD.txt (ISO-8859-1, tab).

Uso:
    python conciliar_ventas_sri.py venta/odoo_venta01.xlsx venta/sri_emitidos salida.xlsx

Clave de cruce: número de factura (NNN-NNN-NNNNNNNNN). Se compara además la
clave de acceso, la fecha de emisión y base/total (tolerancia $0.02).
"""
import glob
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TOL = 0.02
FONT = "Arial"
COLORES = {  # convención del proyecto
    "Resumen": "1F4E79", "Solo_en_SRI": "ED7D31", "Solo_en_Odoo": "C00000",
    "Diferencias": "FFC000", "Odoo_no_autorizado": "7030A0", "Detalle": "808080",
}

COLS_ODOO = ["num", "clave", "cliente", "sucursal", "fecha", "f_aut",
             "base", "total", "estado", "edi"]


def leer_odoo(path):
    o = pd.read_excel(path, dtype=str)
    if len(o.columns) < len(COLS_ODOO):
        raise ValueError(
            f"El Excel de Odoo tiene {len(o.columns)} columnas; se esperan "
            f"{len(COLS_ODOO)} (Número, Número de autorización, Nombre del socio, "
            "Sucursal, Fecha de Factura, Fecha de autorización, Importe sin impuestos, "
            "Total, Estado, Estado electrónico).")
    o = o.iloc[:, :len(COLS_ODOO)].copy()
    o.columns = COLS_ODOO[:len(o.columns)]
    o["base"] = o["base"].astype(float)
    o["total"] = o["total"].astype(float)
    o["fecha"] = pd.to_datetime(o["fecha"].str[:10])
    return o


def leer_sri(carpeta):
    partes = []
    archivos = sorted(glob.glob(os.path.join(carpeta, "emitidos_*.txt")))
    if not archivos:
        raise ValueError("No se encontraron archivos emitidos_*.txt del SRI.")
    requeridas = {"SERIE_COMPROBANTE", "CLAVE_ACCESO", "FECHA_EMISION",
                  "VALOR_SIN_IMPUESTOS", "IVA", "IMPORTE_TOTAL"}
    for f in archivos:
        d = pd.read_csv(f, sep="\t", encoding="iso-8859-1", dtype=str)
        faltan = requeridas - set(d.columns)
        if faltan:
            raise ValueError(
                f"{os.path.basename(f)}: no parece un TXT de emitidos del SRI "
                f"(faltan columnas {sorted(faltan)}).")
        d["archivo"] = os.path.basename(f)
        partes.append(d)
    s = pd.concat(partes, ignore_index=True)
    s = s.rename(columns={"SERIE_COMPROBANTE": "num", "CLAVE_ACCESO": "clave"})
    s["base"] = s["VALOR_SIN_IMPUESTOS"].astype(float)
    s["iva"] = s["IVA"].astype(float)
    s["total"] = s["IMPORTE_TOTAL"].astype(float)
    s["fecha"] = pd.to_datetime(s["FECHA_EMISION"].str[:10], format="%d/%m/%Y")
    return s[["num", "clave", "fecha", "base", "iva", "total", "archivo"]]


def conciliar(odoo, sri):
    dup_sri = sri[sri.duplicated("num", keep=False)]
    sri = sri.drop_duplicates("num")
    # Odoo: lo que debería estar autorizado vs lo que no
    o_ok = odoo[(odoo["estado"] == "Publicado") & (odoo["edi"] == "authorized")]
    o_no = odoo.drop(o_ok.index)

    m = sri.merge(o_ok, on="num", how="outer", suffixes=("_sri", "_odoo"),
                  indicator=True)
    solo_sri = m[m["_merge"] == "left_only"].copy()
    solo_odoo = m[m["_merge"] == "right_only"].copy()
    ambos = m[m["_merge"] == "both"].copy()

    ambos["dif_base"] = (ambos["base_sri"] - ambos["base_odoo"]).round(2)
    ambos["dif_total"] = (ambos["total_sri"] - ambos["total_odoo"]).round(2)
    mal_valor = (ambos["dif_base"].abs() > TOL) | (ambos["dif_total"].abs() > TOL)
    mal_clave = ambos["clave_sri"] != ambos["clave_odoo"]
    mal_fecha = ambos["fecha_sri"] != ambos["fecha_odoo"]
    dif = ambos[mal_valor | mal_clave | mal_fecha].copy()
    dif["motivo"] = ""
    dif.loc[mal_valor, "motivo"] += "Valor; "
    dif.loc[mal_clave, "motivo"] += "Clave de acceso; "
    dif.loc[mal_fecha, "motivo"] += "Fecha; "

    # Odoo no autorizado: ¿existe en el SRI de todos modos?
    o_no = o_no.copy()
    o_no["en_sri"] = o_no["num"].isin(sri["num"])

    estado = pd.concat([
        ambos.loc[~(mal_valor | mal_clave | mal_fecha)].assign(resultado="Coincide"),
        dif.assign(resultado="Diferencia"),
        solo_sri.assign(resultado="Solo en SRI"),
        solo_odoo.assign(resultado="Solo en Odoo"),
    ], ignore_index=True)
    return estado, solo_sri, solo_odoo, dif, o_no, dup_sri


def _hoja(wb, nombre, df, color, anchos=None):
    ws = wb.create_sheet(nombre)
    ws.sheet_properties.tabColor = COLORES[color]
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else
                   (v.to_pydatetime().date() if isinstance(v, pd.Timestamp) else v)
                   for v in r])
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COLORES[color])
        c.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            if hasattr(c.value, "year"):
                c.number_format = "dd/mm/yyyy"
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = (anchos or {}).get(col, 18)
    ws.freeze_panes = "A2"
    return ws


def escribir(estado, solo_sri, solo_odoo, dif, o_no, dup_sri, odoo, sri, salida):
    wb = Workbook()
    wb.remove(wb.active)

    det = estado[["num", "resultado", "fecha_sri", "fecha_odoo", "base_sri", "base_odoo",
                  "total_sri", "total_odoo", "clave_sri", "clave_odoo", "sucursal",
                  "cliente"]].copy()
    det["sucursal"] = det["sucursal"].fillna("")
    det["serie"] = det["num"].str[:7]
    det["mes"] = det["fecha_sri"].fillna(det["fecha_odoo"]).dt.strftime("%Y-%m")
    det = det.sort_values(["resultado", "num"])
    ws_det = _hoja(wb, "Detalle", det, "Detalle", {"clave_sri": 52, "clave_odoo": 52, "cliente": 36})
    n = len(det) + 1
    # columnas del Detalle: B=resultado, E/F=base, G/H=total, M=serie, N=mes
    col = {c: get_column_letter(i + 1) for i, c in enumerate(det.columns)}

    ws = wb.create_sheet("Resumen", 0)
    ws.sheet_properties.tabColor = COLORES["Resumen"]
    ws["A1"] = "Conciliación facturas de venta: Odoo vs SRI"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = f"Período: {odoo['fecha'].min():%d/%m/%Y} al {odoo['fecha'].max():%d/%m/%Y}  |  Tolerancia valor: ${TOL}"
    ws["A2"].font = Font(name=FONT, italic=True)
    hdr = ["Resultado", "Facturas", "Base SRI", "Base Odoo", "Total SRI", "Total Odoo"]
    ws.append([])
    ws.append(hdr)
    for c in ws[4]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COLORES["Resumen"])
    D = "Detalle"
    for r, nombre in enumerate(["Coincide", "Diferencia", "Solo en SRI", "Solo en Odoo"], 5):
        ws.cell(r, 1, nombre)
        ws.cell(r, 2, f'=COUNTIF({D}!${col["resultado"]}$2:${col["resultado"]}${n},A{r})')
        for j, cn in enumerate(["base_sri", "base_odoo", "total_sri", "total_odoo"], 3):
            L = col[cn]
            ws.cell(r, j, f'=SUMIFS({D}!${L}$2:${L}${n},{D}!${col["resultado"]}$2:${col["resultado"]}${n},$A{r})')
    ws.cell(9, 1, "Total")
    for j in range(2, 7):
        L = get_column_letter(j)
        ws.cell(9, j, f"=SUM({L}5:{L}8)")
    ws.cell(11, 1, "Odoo no autorizado (Cancelado / issued / sin estado)")
    ws.cell(11, 2, f"=COUNTA(Odoo_no_autorizado!A2:A{len(o_no) + 1})")
    ws.cell(12, 1, "  de los cuales existen en el SRI")
    ws.cell(12, 2, f'=COUNTIF(Odoo_no_autorizado!I2:I{len(o_no) + 1},TRUE)')
    ws.cell(13, 1, "Números duplicados en TXT del SRI")
    ws.cell(13, 2, len(dup_sri))
    for row in ws.iter_rows(min_row=5, max_row=13):
        for c in row:
            c.font = Font(name=FONT, bold=(c.row == 9))
            if c.column >= 3:
                c.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 52
    for L in "BCDEF":
        ws.column_dimensions[L].width = 16

    # por serie y mes (fórmulas COUNTIFS sobre Detalle)
    ws.cell(15, 1, "Por serie / mes").font = Font(name=FONT, bold=True)
    series = sorted(det["serie"].unique())
    meses = sorted(det["mes"].unique())
    ws.cell(16, 1, "Serie")
    ws.cell(16, 2, "Mes")
    for j, nombre in enumerate(["Coincide", "Diferencia", "Solo en SRI", "Solo en Odoo"], 3):
        ws.cell(16, j, nombre)
    for c in ws[16]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COLORES["Resumen"])
    r = 17
    for s in series:
        for m_ in meses:
            ws.cell(r, 1, s)
            ws.cell(r, 2, m_)
            for j in range(3, 7):
                ws.cell(r, j, f'=COUNTIFS({D}!${col["serie"]}$2:${col["serie"]}${n},$A{r},'
                              f'{D}!${col["mes"]}$2:${col["mes"]}${n},$B{r},'
                              f'{D}!${col["resultado"]}$2:${col["resultado"]}${n},{get_column_letter(j)}$16)')
            for c in ws[r]:
                c.font = Font(name=FONT)
            r += 1

    _hoja(wb, "Solo_en_SRI", solo_sri[["num", "fecha_sri", "base_sri", "iva", "total_sri", "clave_sri"]]
          .rename(columns=lambda x: x), "Solo_en_SRI", {"clave_sri": 52})
    _hoja(wb, "Solo_en_Odoo", solo_odoo[["num", "fecha_odoo", "base_odoo", "total_odoo", "clave_odoo",
                                          "cliente", "sucursal"]], "Solo_en_Odoo", {"clave_odoo": 52, "cliente": 36})
    _hoja(wb, "Diferencias", dif[["num", "motivo", "fecha_sri", "fecha_odoo", "base_sri", "base_odoo",
                                   "dif_base", "total_sri", "total_odoo", "dif_total", "clave_sri", "clave_odoo"]],
          "Diferencias", {"clave_sri": 52, "clave_odoo": 52})
    nn = o_no[["num", "fecha", "base", "total", "estado", "edi", "cliente", "sucursal", "en_sri"]]
    _hoja(wb, "Odoo_no_autorizado", nn, "Odoo_no_autorizado", {"cliente": 36})
    # En Resumen, COUNTIF sobre TRUE: la columna I es en_sri (booleano)
    wb.save(salida)


def conciliar_ventas(odoo_path, emitidos_dir, salida):
    """Concilia Odoo vs SRI, escribe el Excel en `salida` y devuelve el resumen."""
    odoo = leer_odoo(odoo_path)
    sri = leer_sri(emitidos_dir)
    # la carpeta puede tener otros períodos: limitar al rango del Excel de Odoo
    sri = sri[(sri["fecha"] >= odoo["fecha"].min()) & (sri["fecha"] <= odoo["fecha"].max())]
    estado, solo_sri, solo_odoo, dif, o_no, dup = conciliar(odoo, sri)
    escribir(estado, solo_sri, solo_odoo, dif, o_no, dup, odoo, sri, salida)
    cuenta = estado["resultado"].value_counts()
    return {
        "coinciden": int(cuenta.get("Coincide", 0)),
        "diferencias": int(cuenta.get("Diferencia", 0)),
        "solo_sri": int(cuenta.get("Solo en SRI", 0)),
        "solo_odoo": int(cuenta.get("Solo en Odoo", 0)),
        "odoo_no_autorizado": len(o_no),
        "no_autorizado_en_sri": int(o_no["en_sri"].sum()),
        "duplicados_sri": len(dup),
        "n_sri": int(sri["num"].nunique()),
        "n_odoo": len(odoo),
    }


def main():
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    r = conciliar_ventas(*sys.argv[1:4])
    for k, v in r.items():
        print(f"{k}: {v}")
    print("Escrito:", sys.argv[3])


if __name__ == "__main__":
    main()
