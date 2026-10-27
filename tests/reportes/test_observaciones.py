# -*- coding: utf-8 -*-
import openpyxl

from app.reportes.observaciones import ACCIONABLES, COLUMNAS_OBS, QUE_CORREGIR, escribir_excel_observaciones


def _fila(estado, n, **kw):
    f = {"tipo": "FAC", "estado": estado, "documento": f"001-001-{n:09d}", "fecha": "2026-09-05",
         "contraparte": "X", "ruc": None, "valor_sri": 10.0, "valor_odoo": 9.0, "diferencia": 1.0,
         "detalle": "", "que_corregir": QUE_CORREGIR.get(estado, ""), "accionable": estado in ACCIONABLES}
    f.update(kw)
    return f


def test_constantes():
    assert ACCIONABLES == frozenset(QUE_CORREGIR)
    assert "Coincide" not in ACCIONABLES
    assert COLUMNAS_OBS[-1] == "accionable" and len(COLUMNAS_OBS) == 12


def test_excel_observaciones_solo_accionables(tmp_path):
    exc = [_fila("Diferencia de valor", 3), _fila("Solo en SRI", 1), _fila("Solo en SRI", 2),
           _fila("Informativo", 4), _fila("Odoo anulado, SRI autorizado", 5)]
    ruta = str(tmp_path / "o.xlsx")
    n = escribir_excel_observaciones({"excepciones": exc}, ruta)
    assert n == sum(f["accionable"] for f in exc) == 4
    wb = openpyxl.load_workbook(ruta)
    # orden de QUE_CORREGIR, no de aparición
    assert wb.sheetnames == ["Solo en SRI", "Diferencia de valor", "Odoo anulado, SRI autorizado"]
    ws = wb["Solo en SRI"]
    assert [c.value for c in ws[1]] == COLUMNAS_OBS[:-1]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    textos = [c.value for w in wb for r in w.iter_rows() for c in r]
    assert "Coincide" not in textos and "Informativo" not in textos
    assert sum(w.max_row - 1 for w in wb) == 4
    assert all(c.font.name == "Arial" for w in wb for r in w.iter_rows() for c in r if c.value is not None)


def test_excel_sin_observaciones_es_valido(tmp_path):
    ruta = str(tmp_path / "o.xlsx")
    n = escribir_excel_observaciones({"excepciones": []}, ruta)
    assert n == 0
    wb = openpyxl.load_workbook(ruta)
    assert wb.sheetnames == ["Sin_observaciones"]
    assert "No hay observaciones" in wb["Sin_observaciones"]["A1"].value
