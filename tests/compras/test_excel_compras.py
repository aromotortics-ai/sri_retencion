# -*- coding: utf-8 -*-
from datetime import date

import openpyxl
import pandas as pd

from app.compras.conciliar_compras import CAT_DUP, CAT_INFO, CAT_SOLO_ODOO
from app.compras.excel_compras import construir_excel

DESDE, HASTA = date(2026, 9, 1), date(2026, 9, 30)


def _sri(filas):
    return pd.DataFrame(filas, columns=["clave", "serie", "ruc", "proveedor", "fecha", "tipo",
                                        "base", "iva", "total", "afectada"])


def _odoo(filas):
    return pd.DataFrame(filas, columns=["id", "num", "nfac", "socio", "ruc", "fecha", "tipo",
                                        "cod_comp", "clave", "base", "iva", "total", "afectada",
                                        "elegible"])


def _datos(proveedor="PROVEEDOR SA"):
    f = pd.Timestamp("2026-09-05")
    sri = _sri([
        ["1" * 49, "001-001-000000001", "1790000000001", proveedor, f, "FAC", 100.0, 15.0, 115.0, pd.NA],
        ["2" * 49, "001-001-000000002", "1790000000001", "OTRO SA", f, "FAC", 200.0, 30.0, 230.0, pd.NA],
        ["3" * 49, "001-001-000000003", "1790000000001", "OTRO SA", f, "NC", 10.0, 1.5, 11.5, pd.NA],
    ])
    odoo = _odoo([
        [10, "BILL/1", "001-001-000000001", proveedor, "1790000000001", f, "FAC", "01", "1" * 49,
         100.0, 15.0, 115.0, pd.NA, True],
        [11, "BILL/2", "001-001-000000002", "OTRO SA", "1790000000001", f, "FAC", "01", "2" * 49,
         200.0, 30.0, 230.0, pd.NA, True],
    ])
    pares = pd.DataFrame({
        "idx_sri": pd.array([0, 1, 2], dtype="Int64"),
        "id_odoo": pd.array([10, 11, None], dtype="Int64"),
        "tipo": ["FAC", "FAC", "NC"], "metodo": ["Clave de acceso", "Clave de acceso", ""],
        "cat": ["", "", ""]})
    return sri, odoo, pares


def _construir(tmp_path, **kw):
    sri, odoo, pares = _datos(**kw)
    salida = str(tmp_path / "out.xlsx")
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, salida)
    return res, openpyxl.load_workbook(salida)


def test_hojas_esperadas_y_colores(tmp_path):
    _, wb = _construir(tmp_path)
    assert wb.sheetnames == ["Resumen", "Conciliacion", "Solo_en_SRI", "Solo_en_Odoo", "Diferencia_Valor",
                             "Duplicados_Odoo", "Informativos", "SRI_Listado", "Odoo_Lineas"]
    colores = {"Resumen": "1F4E79", "Conciliacion": "1F4E79", "Solo_en_SRI": "ED7D31",
               "Solo_en_Odoo": "C00000", "Diferencia_Valor": "FFC000", "Duplicados_Odoo": "7030A0",
               "Informativos": "808080", "SRI_Listado": "808080", "Odoo_Lineas": "808080"}
    for nombre, color in colores.items():
        assert wb[nombre].sheet_properties.tabColor.rgb.endswith(color), nombre
    assert wb.calculation.fullCalcOnLoad is True


def test_resultado_es_formula_no_valor(tmp_path):
    _, wb = _construir(tmp_path)
    ws = wb["Conciliacion"]
    col = [c.value for c in ws[1]].index("Resultado") + 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        assert isinstance(v, str) and v.startswith("=")
    assert "0.02" in ws.cell(row=2, column=col).value


def test_resumen_cuenta_por_tipo(tmp_path):
    res, wb = _construir(tmp_path)
    assert res["FAC"] == {"coinciden": 2, "solo_sri": 0, "solo_odoo": 0, "diferencias": 0, "informativos": 0}
    assert res["NC"] == {"coinciden": 0, "solo_sri": 1, "solo_odoo": 0, "diferencias": 0, "informativos": 0}
    assert res["ND"] == {"coinciden": 0, "solo_sri": 0, "solo_odoo": 0, "diferencias": 0, "informativos": 0}
    formulas = [c.value for row in wb["Resumen"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=COUNTIFS")]
    assert any("Conciliacion!" in f for f in formulas)


def test_diferencia_solo_odoo_e_informativo(tmp_path):
    sri, odoo, pares = _datos()
    sri.at[1, "total"] = 240.0
    f = pd.Timestamp("2026-09-07")
    odoo = pd.concat([odoo, _odoo([
        [12, "BILL/3", "001-001-000000009", "X", "1790000000001", f, "FAC", "01", "9" * 49, 5.0, 0.0, 5.0, pd.NA, True],
        [13, "BILL/4", "", "Y", "", f, "FAC", "02", pd.NA, 5.0, 0.0, 5.0, pd.NA, False],
    ])], ignore_index=True)
    pares = pd.concat([pares, pd.DataFrame({
        "idx_sri": pd.array([None, None], dtype="Int64"), "id_odoo": pd.array([12, 13], dtype="Int64"),
        "tipo": ["FAC", "FAC"], "metodo": ["", ""], "cat": [CAT_SOLO_ODOO, CAT_INFO]})], ignore_index=True)
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, str(tmp_path / "o.xlsx"))
    assert res["FAC"] == {"coinciden": 1, "solo_sri": 0, "solo_odoo": 1, "diferencias": 1, "informativos": 1}


def test_texto_que_empieza_con_igual_no_se_guarda_como_formula(tmp_path):
    _, wb = _construir(tmp_path, proveedor='=HYPERLINK("http://x","click")')
    celdas = [c for ws in wb for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert celdas and all(c.data_type == "s" for c in celdas)


def test_fuente_arial(tmp_path):
    _, wb = _construir(tmp_path)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    assert c.font.name == "Arial", (ws.title, c.coordinate)


def _col(ws, nombre):
    return [c.value for c in ws[1]].index(nombre) + 1


def _filas(ws):
    return [[c.value for c in row] for row in ws.iter_rows(min_row=2)]


def test_fecha_fuera_del_periodo_marca_y_sigue_coincidiendo(tmp_path):
    sri, odoo, pares = _datos()
    odoo.at[0, "fecha"] = pd.Timestamp("2026-10-03")
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, str(tmp_path / "f.xlsx"))
    assert res["FAC"]["coinciden"] == 2
    ws = openpyxl.load_workbook(str(tmp_path / "f.xlsx"))["Conciliacion"]
    f = ws.cell(row=2, column=_col(ws, "Observacion")).value
    assert f.startswith("=") and "Fecha fuera del período" in f and "DATE(2026,9,1)" in f and "DATE(2026,9,30)" in f
    fila_odoo10 = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=_col(ws, "ID_Odoo")).value == 10][0]
    assert "Fecha fuera del período" in ws.cell(row=fila_odoo10, column=_col(ws, "Observacion")).value


def test_duplicados_en_hoja_resultado_y_resumen(tmp_path):
    sri, odoo, pares = _datos()
    f = pd.Timestamp("2026-09-06")
    odoo = pd.concat([odoo, _odoo([[12, "BILL/9", "001-001-000000001", "PROVEEDOR SA", "1790000000001", f, "FAC",
                                    "01", "1" * 49, 100.0, 15.0, 115.0, pd.NA, True]])], ignore_index=True)
    pares = pd.concat([pares, pd.DataFrame({
        "idx_sri": pd.array([0], dtype="Int64"), "id_odoo": pd.array([12], dtype="Int64"),
        "tipo": ["FAC"], "metodo": ["Clave de acceso"], "cat": [CAT_DUP]})], ignore_index=True)
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, str(tmp_path / "d.xlsx"))
    assert res["FAC"]["coinciden"] == 2 and res["FAC"]["solo_odoo"] == 0
    wb = openpyxl.load_workbook(str(tmp_path / "d.xlsx"))
    dup = [r for r in _filas(wb["Duplicados_Odoo"])]
    assert any("BILL/9" in r for r in dup)
    ws = wb["Conciliacion"]
    rc = _col(ws, "Resultado")
    assert f'IF(P2="{CAT_DUP}",P2' in ws.cell(row=2, column=rc).value
    rs = [c.value for row in wb["Resumen"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(v.startswith("=COUNTIFS") and f'"{CAT_DUP}"' in v for v in rs)


def test_filas_de_las_hojas_de_resultado(tmp_path):
    sri, odoo, pares = _datos()
    sri.at[1, "total"] = 240.0
    f = pd.Timestamp("2026-09-07")
    odoo = pd.concat([odoo, _odoo([
        [12, "BILL/3", "001-001-000000009", "X", "1790000000001", f, "FAC", "01", "9" * 49, 5.0, 0.0, 5.0, pd.NA, True],
        [13, "BILL/4", "001-001-000000010", "Y", "1790000000001", f, "FAC", "02", pd.NA, 6.0, 0.0, 6.0, pd.NA, False],
    ])], ignore_index=True)
    pares = pd.concat([pares, pd.DataFrame({
        "idx_sri": pd.array([None, None], dtype="Int64"), "id_odoo": pd.array([12, 13], dtype="Int64"),
        "tipo": ["FAC", "FAC"], "metodo": ["", ""], "cat": [CAT_SOLO_ODOO, CAT_INFO]})], ignore_index=True)
    construir_excel(sri, odoo, pares, DESDE, HASTA, str(tmp_path / "r.xlsx"))
    wb = openpyxl.load_workbook(str(tmp_path / "r.xlsx"))

    d = _filas(wb["Diferencia_Valor"])
    assert len(d) == 1 and "001-001-000000002" in d[0] and 240.0 in d[0] and 230.0 in d[0] and 10.0 in d[0]
    so = _filas(wb["Solo_en_Odoo"])
    assert len(so) == 1 and "001-001-000000009" in so[0] and 5.0 in so[0] and CAT_SOLO_ODOO in so[0]
    inf = _filas(wb["Informativos"])
    assert len(inf) == 1 and "001-001-000000010" in inf[0] and 6.0 in inf[0] and CAT_INFO in inf[0]
    ss = _filas(wb["Solo_en_SRI"])
    assert len(ss) == 1 and "001-001-000000003" in ss[0] and "NC" in ss[0]

    rs = [c.value for row in wb["Resumen"].iter_rows() for c in row
          if isinstance(c.value, str) and c.value.startswith("=COUNTIFS")]
    for texto in ("Coincide", "Diferencia de valor", "Solo en SRI", CAT_SOLO_ODOO, "Duplicado en Odoo", CAT_INFO):
        assert any(f'"{texto}")' in v for v in rs), texto


def test_pares_vacios_sin_formulas_circulares(tmp_path):
    sri, odoo, _ = _datos()
    sri, odoo = sri.iloc[0:0], odoo.iloc[0:0]
    pares = pd.DataFrame({"idx_sri": pd.array([], dtype="Int64"), "id_odoo": pd.array([], dtype="Int64"),
                          "tipo": [], "metodo": [], "cat": []})
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, str(tmp_path / "v.xlsx"))
    assert all(v == 0 for t in res.values() for v in t.values())
    ws = openpyxl.load_workbook(str(tmp_path / "v.xlsx"))["Resumen"]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=SUM("):
                assert c.coordinate not in c.value.replace("$", "")
                a, b = c.value[5:-1].split(":")
                assert not (int("".join(ch for ch in a if ch.isdigit())) <= c.row
                            <= int("".join(ch for ch in b if ch.isdigit())))
