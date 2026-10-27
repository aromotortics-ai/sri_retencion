import json
from datetime import date

import openpyxl
import pandas as pd

import app.odoo.client as odoo_client
from app.job_manager import create_job, get_job_dir, read_status
from app.odoo.client import OdooError
from _helpers import _autorizacion


def _job(tmp_path, monkeypatch):
    monkeypatch.setattr("app.job_manager.JOBS_DIR", tmp_path / "jobs")
    import app.compras.pipeline_compras as pipeline_compras
    job_id = create_job()
    (get_job_dir(job_id) / "sri").mkdir()
    monkeypatch.setattr(pipeline_compras, "schedule_cleanup", lambda *a, **k: None)
    # Estos tests simulan solo el lado de compras: sin retenciones en los TXT.
    monkeypatch.setattr(pipeline_compras, "leer_sri_retenciones_txt",
                        lambda rutas: pd.DataFrame(columns=["clave", "serie", "ruc", "proveedor", "fecha"]))
    monkeypatch.setattr(pipeline_compras, "CACHE_XML_DIR", tmp_path / "cache" / "xml")
    return pipeline_compras, job_id


class _Lector:
    llamado = None

    def fetch_compras(self, desde, hasta, margen_dias=5):
        _Lector.llamado = (desde, hasta)
        return pd.DataFrame()


def test_pipeline_compras_ok(tmp_path, monkeypatch):
    pc, job_id = _job(tmp_path, monkeypatch)
    sri = pd.DataFrame({"fecha": pd.to_datetime(["2026-09-05", "2026-09-20"])})
    monkeypatch.setattr(pc, "leer_sri_txts", lambda rutas: sri)
    monkeypatch.setattr(pc.OdooReader, "from_env", classmethod(lambda cls: _Lector()))
    visto = {}

    def falso(rutas, odoo, desde, hasta, salida, ret=None, periodo_ret=None):
        visto.update(rutas=rutas, desde=desde, hasta=hasta, salida=salida, ret=ret,
                     periodo_ret=periodo_ret)
        return {"por_tipo": {"FAC": {"coinciden": 1}}, "advertencias": []}
    monkeypatch.setattr(pc, "conciliar_compras", falso)
    (get_job_dir(job_id) / "sri" / "a.txt").write_text("x")
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done"
    assert s["resumen"]["por_tipo"]["FAC"]["coinciden"] == 1
    assert _Lector.llamado == (date(2026, 9, 1), date(2026, 9, 30))
    assert (visto["desde"], visto["hasta"]) == (date(2026, 9, 1), date(2026, 9, 30))
    assert visto["salida"].endswith("resultado.xlsx")
    assert visto["ret"] is None and visto["periodo_ret"] is None


def _sri(*fechas):
    return pd.DataFrame({"fecha": pd.to_datetime(list(fechas))})


def test_periodo_mes_dominante_con_atipico_y_ventana_amplia(tmp_path, monkeypatch):
    pc, _ = _job(tmp_path, monkeypatch)
    sri = _sri("2026-07-20", "2026-09-05", "2026-09-06", "2026-09-30")
    assert pc.periodo_sri(sri) == (date(2026, 9, 1), date(2026, 9, 30),
                                   date(2026, 7, 20), date(2026, 9, 30))


def test_periodo_empate_gana_el_mes_mas_reciente(tmp_path, monkeypatch):
    pc, _ = _job(tmp_path, monkeypatch)
    r = pc.periodo_sri(_sri("2026-08-10", "2026-09-05"))
    assert r[:2] == (date(2026, 9, 1), date(2026, 9, 30))
    assert r[2:] == (date(2026, 8, 10), date(2026, 9, 30))


def test_periodo_varios_meses_domina_uno(tmp_path, monkeypatch):
    pc, _ = _job(tmp_path, monkeypatch)
    sri = _sri("2026-06-01", "2026-08-02", "2026-08-03", "2026-08-04", "2026-08-31", "2026-09-01")
    assert pc.periodo_sri(sri)[:2] == (date(2026, 8, 1), date(2026, 8, 31))
    assert pc.periodo_sri(sri)[2:] == (date(2026, 6, 1), date(2026, 9, 1))


def test_pipeline_compras_sin_credenciales_error_claro(tmp_path, monkeypatch):
    pc, job_id = _job(tmp_path, monkeypatch)
    sri = pd.DataFrame({"fecha": pd.to_datetime(["2026-09-05"])})
    monkeypatch.setattr(pc, "leer_sri_txts", lambda rutas: sri)

    def sin_env(cls):
        raise OdooError("Faltan variables de entorno de Odoo: ODOO_URL")
    monkeypatch.setattr(pc.OdooReader, "from_env", classmethod(sin_env))
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert "ODOO_URL" in s["mensaje"]


def test_pipeline_compras_sin_filas_sri_es_error(tmp_path, monkeypatch):
    pc, job_id = _job(tmp_path, monkeypatch)
    monkeypatch.setattr(pc, "leer_sri_txts", lambda rutas: pd.DataFrame({"fecha": []}))
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert "factura" in s["mensaje"].lower()


def test_pipeline_ventana_amplia_para_odoo_y_mes_dominante_para_conciliar(tmp_path, monkeypatch):
    pc, job_id = _job(tmp_path, monkeypatch)
    monkeypatch.setattr(pc, "leer_sri_txts", lambda rutas: _sri("2026-07-20", "2026-09-05", "2026-09-06"))
    monkeypatch.setattr(pc.OdooReader, "from_env", classmethod(lambda cls: _Lector()))
    visto = {}
    monkeypatch.setattr(pc, "conciliar_compras",
                        lambda r, o, d, h, s, ret=None, periodo_ret=None: visto.update(d=d, h=h) or {"por_tipo": {}, "advertencias": []})
    pc.run_pipeline_compras(job_id)
    assert read_status(job_id)["estado"] == "done"
    assert _Lector.llamado == (date(2026, 7, 20), date(2026, 9, 30))
    assert (visto["d"], visto["h"]) == (date(2026, 9, 1), date(2026, 9, 30))


def test_periodo_diciembre_con_atipico_de_enero(tmp_path, monkeypatch):
    pc, _ = _job(tmp_path, monkeypatch)
    sri = _sri("2026-12-05", "2026-12-06", "2026-12-20", "2027-01-03")
    assert pc.periodo_sri(sri) == (date(2026, 12, 1), date(2026, 12, 31),
                                   date(2026, 12, 1), date(2027, 1, 3))


def test_pipeline_con_error_agenda_limpieza(tmp_path, monkeypatch):
    pc, job_id = _job(tmp_path, monkeypatch)
    llamadas = []
    monkeypatch.setattr(pc, "schedule_cleanup", lambda *a, **k: llamadas.append((a, k)))
    monkeypatch.setattr(pc, "leer_sri_txts", lambda rutas: pd.DataFrame({"fecha": []}))
    pc.run_pipeline_compras(job_id)
    assert read_status(job_id)["estado"] == "error"
    assert llamadas == [((job_id,), {"delay_seconds": 7200})]


# ------------------------------------------------------------------ retenciones en Compras
# TXT sinteticos de punta a punta: lectura real de TXT, conciliacion y Excel reales; Odoo,
# descarga de XML y directorio de cache falsos (nunca red, nunca Odoo, nunca el cache real).

COLS_TXT = ["RUC_EMISOR", "RAZON_SOCIAL_EMISOR", "TIPO_COMPROBANTE", "SERIE_COMPROBANTE",
            "CLAVE_ACCESO", "FECHA_AUTORIZACION", "FECHA_EMISION", "IDENTIFICACION_RECEPTOR",
            "VALOR_SIN_IMPUESTOS", "IVA", "IMPORTE_TOTAL", "NUMERO_DOCUMENTO_MODIFICADO"]
RUC = "1790000000001"
CLAVE_FAC = "9" * 49
CLAVE_R1 = "1" * 49
CLAVE_R2 = "2" * 49
N_RET1 = "001-200-000004340"  # el n_ret que trae el XML sintetico de test_xml_cache
N_RET2 = "001-200-000004341"
ADV_FALLIDOS = ("1 XML de retenciones no se pudieron descargar o leer; sus retenciones se listan como "
                "'XML no disponible'")


def _escribir_txt(path, filas):
    lineas = ["\t".join(COLS_TXT)]
    for f in filas:
        d = dict(zip(COLS_TXT, [""] * len(COLS_TXT)))
        d.update(f)
        lineas.append("\t".join(d[c] for c in COLS_TXT))
    path.write_bytes("\n".join(lineas).encode("iso-8859-1"))


def _fac(clave=CLAVE_FAC, fecha="05/09/2026"):
    return {"RUC_EMISOR": RUC, "RAZON_SOCIAL_EMISOR": "PROVEEDOR SA", "TIPO_COMPROBANTE": "Factura",
            "SERIE_COMPROBANTE": "001-001-000000001", "CLAVE_ACCESO": clave, "FECHA_EMISION": fecha,
            "VALOR_SIN_IMPUESTOS": "100.00", "IVA": "15.00", "IMPORTE_TOTAL": "115.00"}


def _ret_txt(clave, serie=N_RET1, fecha="05/08/2026"):
    return {"RUC_EMISOR": RUC, "RAZON_SOCIAL_EMISOR": "PROVEEDOR SA",
            "TIPO_COMPROBANTE": "Comprobante de Retención", "SERIE_COMPROBANTE": serie,
            "CLAVE_ACCESO": clave, "FECHA_EMISION": fecha}


def _xml(cache, clave):
    cache.mkdir(parents=True, exist_ok=True)
    (cache / f"{clave}.xml").write_text(
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + _autorizacion(clave, "2.0.0"),
        encoding="utf-8")


def _linea_odoo(n_ret=N_RET1, id_ret=10, fecha="2026-08-05"):
    return {"id_ret": id_ret, "n_ret": n_ret, "fecha": pd.Timestamp(fecha), "estado": "posted",
            "socio": "PROVEEDOR SA", "ruc": RUC, "factura": "001-003-000088766", "pct": 1.75,
            "base": 100.0, "valor": 1.75, "cod": "312"}


class _Odoo:
    """Lector falso: registra las ventanas pedidas."""

    def __init__(self, ret_filas=None):
        self.compras = None
        self.retenciones = None
        self._ret = ret_filas if ret_filas is not None else [_linea_odoo()]

    def fetch_compras(self, desde, hasta, margen_dias=5):
        self.compras = (desde, hasta)
        return pd.DataFrame(columns=odoo_client.COLUMNAS)

    def fetch_retenciones(self, desde, hasta, margen_dias=5):
        self.retenciones = (desde, hasta)
        return pd.DataFrame(self._ret, columns=odoo_client.COLUMNAS_RET)


def _entorno(tmp_path, monkeypatch, txts, lector=None, descarga=None):
    monkeypatch.setattr("app.job_manager.JOBS_DIR", tmp_path / "jobs")
    import app.compras.pipeline_compras as pc
    job_id = create_job()
    sri_dir = get_job_dir(job_id) / "sri"
    sri_dir.mkdir()
    for nombre, filas in txts.items():
        _escribir_txt(sri_dir / nombre, filas)
    monkeypatch.setattr(pc, "schedule_cleanup", lambda *a, **k: None)
    monkeypatch.setattr("shutil.which", lambda *_: None)  # sin LibreOffice: Excel tal cual
    cache = tmp_path / "cache" / "xml"
    monkeypatch.setattr(pc, "CACHE_XML_DIR", cache)
    lector = lector or _Odoo()
    instancias = []

    def from_env(cls):
        instancias.append(1)
        return lector
    monkeypatch.setattr(pc.OdooReader, "from_env", classmethod(from_env))
    llamadas = []

    def descargar_falso(claves, cache_dir, on_progress=None, **kw):
        llamadas.append(list(claves))
        assert str(cache_dir) == str(cache)
        return descarga(claves, cache_dir, on_progress) if descarga else []
    monkeypatch.setattr(pc, "descargar_xmls", descargar_falso)
    return pc, job_id, lector, cache, llamadas, instancias


def _hojas(job_id):
    return openpyxl.load_workbook(get_job_dir(job_id) / "resultado.xlsx").sheetnames


def _espiar_status(pc, monkeypatch):
    vistos = []
    original = pc.write_status

    def espia(job_id, estado, mensaje, progreso=0.0, **kw):
        vistos.append((mensaje, round(progreso, 3)))
        return original(job_id, estado, mensaje, progreso, **kw)
    monkeypatch.setattr(pc, "write_status", espia)
    return vistos


def test_pipeline_con_compras_y_retenciones(tmp_path, monkeypatch):
    pc, job_id, lector, cache, llamadas, instancias = _entorno(
        tmp_path, monkeypatch, {"fac.txt": [_fac()], "ret.txt": [_ret_txt(CLAVE_R1)]})
    _xml(cache, CLAVE_R1)
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    hojas = _hojas(job_id)
    assert "Conciliacion" in hojas and "Ret_Conciliacion" in hojas and "Ret_Informativos" in hojas
    assert s["resumen"]["por_tipo"]["retenciones"]["coinciden"] == 1
    assert s["resumen"]["por_tipo"]["FAC"]["solo_sri"] == 1
    assert s["resumen"]["retenciones_fallidos"] == 0
    assert llamadas == [[CLAVE_R1]]
    assert len(instancias) == 1
    assert lector.compras is not None and lector.retenciones is not None
    assert not any("retenciones" in a for a in s["resumen"]["advertencias"])


_IVA_30 = ("<retencion><codigo>2</codigo><codigoRetencion>1</codigoRetencion><baseImponible>15.00</baseImponible>"
           "<porcentajeRetener>30</porcentajeRetener><valorRetenido>4.50</valorRetenido></retencion>")


def _xml_renta_iva(cache, clave, secuencial="000004340"):
    """XML v2 con DOS lineas (Renta 1.75% + IVA 30%) en un mismo comprobante."""
    cache.mkdir(parents=True, exist_ok=True)
    cuerpo = (_autorizacion(clave, "2.0.0").replace("</retencion></retenciones>", "</retencion>" + _IVA_30
                                                     + "</retenciones>")
              .replace("<secuencial>000004340</secuencial>", f"<secuencial>{secuencial}</secuencial>"))
    (cache / f"{clave}.xml").write_text('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + cuerpo,
                                        encoding="utf-8")


def test_pipeline_retencion_multilinea_cuenta_documentos(tmp_path, monkeypatch):
    # R1: Renta + IVA en SRI y en Odoo (1 documento, 2 lineas). R2: Renta + IVA solo en SRI.
    iva = {**_linea_odoo(), "pct": 30.0, "base": 15.0, "valor": 4.5, "cod": "1"}
    pc, job_id, _, cache, _, _ = _entorno(
        tmp_path, monkeypatch, {"ret.txt": [_ret_txt(CLAVE_R1), _ret_txt(CLAVE_R2, serie=N_RET2)]},
        lector=_Odoo([_linea_odoo(), iva]))
    _xml_renta_iva(cache, CLAVE_R1)
    _xml_renta_iva(cache, CLAVE_R2, secuencial="000004341")
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    r = s["resumen"]["por_tipo"]["retenciones"]
    assert (r["coinciden"], r["solo_sri"], r["diferencia"], r["solo_odoo"]) == (1, 1, 0, 0)
    assert (r["documentos"], r["lineas"]) == (2, 4)
    wb = openpyxl.load_workbook(get_job_dir(job_id) / "resultado.xlsx")
    assert wb["Ret_Documentos"].max_row - 1 == 2
    assert wb["Ret_Conciliacion"].max_row - 1 == 4
    assert wb["Ret_Solo_SRI"].max_row - 1 == 1


def test_pipeline_solo_retenciones(tmp_path, monkeypatch):
    pc, job_id, lector, cache, _, _ = _entorno(
        tmp_path, monkeypatch, {"ret.txt": [_ret_txt(CLAVE_R1)]})
    _xml(cache, CLAVE_R1)
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    adv = s["resumen"]["advertencias"]
    for t in ("facturas", "notas de crédito", "notas de débito"):
        assert f"Falta el TXT de {t}" in adv
    assert "Falta el TXT de retenciones" not in adv
    assert lector.compras is None
    assert s["resumen"]["por_tipo"]["retenciones"]["coinciden"] == 1
    hojas = _hojas(job_id)
    assert "Conciliacion" in hojas and "Ret_Conciliacion" in hojas


def test_pipeline_sin_retenciones_no_cambia(tmp_path, monkeypatch):
    pc, job_id, lector, _, llamadas, _ = _entorno(tmp_path, monkeypatch, {"fac.txt": [_fac()]})
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert "Falta el TXT de retenciones" in s["resumen"]["advertencias"]
    assert llamadas == []
    assert lector.retenciones is None
    assert lector.compras == (date(2026, 9, 1), date(2026, 9, 30))
    assert "retenciones" not in s["resumen"]["por_tipo"]
    assert "retenciones_fallidos" not in s["resumen"]
    assert not any(h.startswith("Ret_") for h in _hojas(job_id))


def test_pipeline_xml_fallidos_no_abortan(tmp_path, monkeypatch):
    lector = _Odoo([_linea_odoo(), _linea_odoo(n_ret=N_RET2, id_ret=11)])

    def descarga(claves, cache_dir, on_progress):
        _xml(cache_dir, CLAVE_R1)
        on_progress(2, 2, 1)
        return [{"clave": CLAVE_R2, "motivo": "NO AUTORIZADO"}]
    pc, job_id, _, _, _, _ = _entorno(
        tmp_path, monkeypatch,
        {"fac.txt": [_fac()], "ret.txt": [_ret_txt(CLAVE_R1), _ret_txt(CLAVE_R2, serie=N_RET2)]},
        lector=lector, descarga=descarga)
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert s["resumen"]["retenciones_fallidos"] == 1
    assert ADV_FALLIDOS in s["resumen"]["advertencias"]
    assert s["resumen"]["por_tipo"]["retenciones"]["informativos"] == 1
    wb = openpyxl.load_workbook(get_job_dir(job_id) / "resultado.xlsx")
    valores = [c for fila in wb["Ret_Informativos"].iter_rows(values_only=True) for c in fila]
    assert N_RET2 in valores


def test_pipeline_xml_ilegible_cuenta_como_fallido(tmp_path, monkeypatch):
    pc, job_id, _, cache, _, _ = _entorno(
        tmp_path, monkeypatch,
        {"ret.txt": [_ret_txt(CLAVE_R1), _ret_txt(CLAVE_R2, serie=N_RET2)]})
    _xml(cache, CLAVE_R1)
    (cache / f"{CLAVE_R2}.xml").write_text("no es xml <<<", encoding="utf-8")
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert s["resumen"]["retenciones_fallidos"] == 1
    assert ADV_FALLIDOS in s["resumen"]["advertencias"]


def test_pipeline_progreso_de_descarga(tmp_path, monkeypatch):
    def descarga(claves, cache_dir, on_progress):
        on_progress(1, 2, 0)
        on_progress(2, 2, 1)
        return [{"clave": CLAVE_R2, "motivo": "x"}]
    pc, job_id, _, _, _, _ = _entorno(
        tmp_path, monkeypatch,
        {"ret.txt": [_ret_txt(CLAVE_R1), _ret_txt(CLAVE_R2, serie=N_RET2)]}, descarga=descarga)
    vistos = _espiar_status(pc, monkeypatch)
    pc.run_pipeline_compras(job_id)
    assert read_status(job_id)["estado"] == "done"
    assert ("Descargando XML de retenciones (1/2, 0 fallidos)", 0.4) in vistos
    assert ("Descargando XML de retenciones (2/2, 1 fallidos)", 0.65) in vistos
    progresos = [p for m, p in vistos if m.startswith("Descargando")]
    assert progresos and all(0.15 <= p <= 0.65 for p in progresos)


def test_pipeline_progreso_completo_si_todo_en_cache(tmp_path, monkeypatch):
    pc, job_id, _, cache, _, _ = _entorno(tmp_path, monkeypatch, {"ret.txt": [_ret_txt(CLAVE_R1)]})
    _xml(cache, CLAVE_R1)
    vistos = _espiar_status(pc, monkeypatch)
    pc.run_pipeline_compras(job_id)
    assert read_status(job_id)["estado"] == "done"
    assert ("Descargando XML de retenciones (1/1, 0 fallidos)", 0.65) in vistos


def test_pipeline_ventana_de_retenciones_es_propia(tmp_path, monkeypatch):
    pc, job_id, lector, _, _, _ = _entorno(
        tmp_path, monkeypatch,
        {"fac.txt": [_fac(fecha="05/09/2026"), _fac(clave="8" * 49, fecha="20/09/2026")],
         "ret.txt": [_ret_txt(CLAVE_R1, fecha="05/08/2026"),
                     _ret_txt(CLAVE_R2, serie=N_RET2, fecha="06/08/2026"),
                     _ret_txt("3" * 49, serie="001-200-000004342", fecha="28/07/2026")]})
    pc.run_pipeline_compras(job_id)
    assert read_status(job_id)["estado"] == "done"
    assert lector.retenciones == (date(2026, 7, 28), date(2026, 8, 31))
    assert lector.compras == (date(2026, 9, 1), date(2026, 9, 30))


def test_pipeline_sin_ningun_tipo_da_error(tmp_path, monkeypatch):
    otro = dict(_fac(), TIPO_COMPROBANTE="Liquidación de compra de Bienes y Prestación de servicios")
    pc, job_id, lector, _, llamadas, _ = _entorno(tmp_path, monkeypatch, {"x.txt": [otro]})
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert ("Los TXT no contienen facturas, notas de crédito, notas de débito ni retenciones"
            in s["mensaje"])
    assert llamadas == [] and lector.compras is None and lector.retenciones is None


def test_pipeline_error_de_odoo_no_filtra_credenciales(tmp_path, monkeypatch):
    secretos = {"ODOO_URL": "https://usuario:clave-url-777@odoo.ejemplo.test",
                "ODOO_DB": "base-secreta-555", "ODOO_USER": "usuario-secreto-333",
                "ODOO_PASSWORD": "clave-secreta-999"}
    for k, v in secretos.items():
        monkeypatch.setenv(k, v)

    class ProxyQueFiltra:
        """Simula un error de bajo nivel cuyo texto incluye URL y credenciales."""

        def __init__(self, url, **kw):
            self.url = url

        def authenticate(self, db, user, password, ctx):
            raise ConnectionError(f"fallo {self.url} {db} {user} {password}")
    monkeypatch.setattr("xmlrpc.client.ServerProxy", ProxyQueFiltra)
    monkeypatch.setattr("app.job_manager.JOBS_DIR", tmp_path / "jobs")
    import app.compras.pipeline_compras as pc
    job_id = create_job()
    sri_dir = get_job_dir(job_id) / "sri"
    sri_dir.mkdir()
    _escribir_txt(sri_dir / "fac.txt", [_fac()])
    _escribir_txt(sri_dir / "ret.txt", [_ret_txt(CLAVE_R1)])
    monkeypatch.setattr(pc, "schedule_cleanup", lambda *a, **k: None)
    monkeypatch.setattr(pc, "CACHE_XML_DIR", tmp_path / "cache" / "xml")
    monkeypatch.setattr(pc, "descargar_xmls", lambda *a, **k: [])
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert s["mensaje"].startswith("OdooError")  # el fallo pasó por el lector real
    crudo = json.dumps(s, ensure_ascii=False)
    for v in secretos.values():
        assert v not in crudo
    assert "clave-url-777" not in crudo


def test_pipeline_retenciones_sin_fechas_validas_es_error_claro(tmp_path, monkeypatch):
    pc, job_id, _, _, llamadas, _ = _entorno(
        tmp_path, monkeypatch, {"ret.txt": [_ret_txt(CLAVE_R1, fecha="")]})
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert "fechas de emisión válidas" in s["mensaje"]
    assert llamadas == []  # no se intentó descargar nada


def test_pipeline_un_xml_no_disponible_por_clave_fallida(tmp_path, monkeypatch):
    # R1 con XML y gemelo; R2 sin XML con gemelo en Odoo; R3 sin XML y sin gemelo.
    clave_r3, n_ret3 = "3" * 49, "001-200-000004342"
    lector = _Odoo([_linea_odoo(), _linea_odoo(n_ret=N_RET2, id_ret=11)])

    def descarga(claves, cache_dir, on_progress):
        _xml(cache_dir, CLAVE_R1)
        return [{"clave": CLAVE_R2, "motivo": "x"}, {"clave": clave_r3, "motivo": "x"}]
    pc, job_id, _, _, _, _ = _entorno(
        tmp_path, monkeypatch,
        {"ret.txt": [_ret_txt(CLAVE_R1), _ret_txt(CLAVE_R2, serie=N_RET2),
                     _ret_txt(clave_r3, serie=n_ret3)]},
        lector=lector, descarga=descarga)
    pc.run_pipeline_compras(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert s["resumen"]["retenciones_fallidos"] == 2
    assert s["resumen"]["por_tipo"]["retenciones"]["informativos"] == 2
    wb = openpyxl.load_workbook(get_job_dir(job_id) / "resultado.xlsx")
    filas = list(wb["Ret_Conciliacion"].iter_rows(min_row=2, values_only=True))
    xml = [f for f in filas if f[0] == "XML no disponible"]
    assert len(xml) == s["resumen"]["retenciones_fallidos"]
    # Clave = columna P (indice 15); Documento y Estado_Documento van despues.
    assert sorted(f[15] for f in xml) == sorted([CLAVE_R2, clave_r3])
    por_clave = {f[15]: f for f in xml}
    assert por_clave[CLAVE_R2][14] == 11            # con gemelo: una sola fila, con datos de Odoo
    assert por_clave[clave_r3][1] == n_ret3         # sin gemelo: lado SRI del TXT
    assert por_clave[clave_r3][14] is None


def test_pipeline_lee_odoo_antes_de_descargar(tmp_path, monkeypatch):
    orden = []

    class Lector(_Odoo):
        def fetch_retenciones(self, desde, hasta, margen_dias=5):
            orden.append("fetch_retenciones")
            return super().fetch_retenciones(desde, hasta, margen_dias)

    def descarga(claves, cache_dir, on_progress):
        orden.append("descargar_xmls")
        return []
    pc, job_id, _, _, _, _ = _entorno(tmp_path, monkeypatch, {"ret.txt": [_ret_txt(CLAVE_R1)]},
                                      lector=Lector(), descarga=descarga)
    vistos = _espiar_status(pc, monkeypatch)
    pc.run_pipeline_compras(job_id)
    assert read_status(job_id)["estado"] == "done"
    assert orden == ["fetch_retenciones", "descargar_xmls"]
    mensajes = [m for m, _ in vistos]
    assert mensajes.index("Leyendo retenciones de Odoo...") < min(
        i for i, m in enumerate(mensajes) if m.startswith("Descargando"))
    progresos = [p for _, p in vistos]
    assert progresos == sorted(progresos)


def test_pipeline_odoo_caido_falla_antes_de_descargar(tmp_path, monkeypatch):
    class Caido(_Odoo):
        def fetch_retenciones(self, desde, hasta, margen_dias=5):
            raise OdooError("No se pudo conectar con Odoo")
    pc, job_id, _, _, llamadas, _ = _entorno(tmp_path, monkeypatch, {"ret.txt": [_ret_txt(CLAVE_R1)]},
                                             lector=Caido())
    pc.run_pipeline_compras(job_id)
    assert read_status(job_id)["estado"] == "error"
    assert llamadas == []


def test_pipeline_rotula_periodo_de_retenciones(tmp_path, monkeypatch):
    pc, job_id, _, cache, _, _ = _entorno(
        tmp_path, monkeypatch, {"fac.txt": [_fac()], "ret.txt": [_ret_txt(CLAVE_R1)]})
    _xml(cache, CLAVE_R1)
    pc.run_pipeline_compras(job_id)
    assert read_status(job_id)["estado"] == "done"
    wb = openpyxl.load_workbook(get_job_dir(job_id) / "resultado.xlsx")
    rs = wb["Resumen"]
    fila = [c.row for c in rs["A"] if c.value == "Retenciones"][0]
    assert rs[f"B{fila}"].value == "2026-08"
    textos = [c.value for row in rs.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("08/2026" in t for t in textos)
