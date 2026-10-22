# -*- coding: utf-8 -*-
from datetime import date

import openpyxl
import pandas as pd

from conciliar_compras import CAT_INFO, CAT_SOLO_ODOO
from excel_compras import construir_excel

DESDE, HASTA = date(2026, 9, 1), date(2026, 9, 30)


def _sri(filas):
    return pd.DataFrame(filas, columns=["clave", "serie", "ruc", "proveedor", "fecha", "tipo",
                                        "base", "iva", "total", "afectada"])


def _odoo(filas):
    return pd.DataFrame(filas, columns=["id", "num", "nfac", "socio", "ruc", "fecha", "tipo",
                                        "cod_comp", "clave", "base", "iva", "total", "afectada",
                                        "elegible"])


def _datos(proveedor="PROVEEDOR SA"):
    f = pd.Timestamp("2026-09-05")
    sri = _sri([
        ["1" * 49, "001-001-000000001", "1790000000001", proveedor, f, "FAC", 100.0, 15.0, 115.0, pd.NA],
        ["2" * 49, "001-001-000000002", "1790000000001", "OTRO SA", f, "FAC", 200.0, 30.0, 230.0, pd.NA],
        ["3" * 49, "001-001-000000003", "1790000000001", "OTRO SA", f, "NC", 10.0, 1.5, 11.5, pd.NA],
    ])
    odoo = _odoo([
        [10, "BILL/1", "001-001-000000001", proveedor, "1790000000001", f, "FAC", "01", "1" * 49,
         100.0, 15.0, 115.0, pd.NA, True],
        [11, "BILL/2", "001-001-000000002", "OTRO SA", "1790000000001", f, "FAC", "01", "2" * 49,
         200.0, 30.0, 230.0, pd.NA, True],
    ])
    pares = pd.DataFrame({
        "idx_sri": pd.array([0, 1, 2], dtype="Int64"),
        "id_odoo": pd.array([10, 11, None], dtype="Int64"),
        "tipo": ["FAC", "FAC", "NC"], "metodo": ["Clave de acceso", "Clave de acceso", ""],
        "cat": ["", "", ""]})
    return sri, odoo, pares


def _construir(tmp_path, **kw):
    sri, odoo, pares = _datos(**kw)
    salida = str(tmp_path / "out.xlsx")
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, salida)
    return res, openpyxl.load_workbook(salida)


def test_hojas_esperadas_y_colores(tmp_path):
    _, wb = _construir(tmp_path)
    assert wb.sheetnames == ["Resumen", "Conciliacion", "Solo_en_SRI", "Solo_en_Odoo", "Diferencia_Valor",
                             "Duplicados_Odoo", "Informativos", "SRI_Listado", "Odoo_Lineas"]
    colores = {"Resumen": "1F4E79", "Conciliacion": "1F4E79", "Solo_en_SRI": "ED7D31",
               "Solo_en_Odoo": "C00000", "Diferencia_Valor": "FFC000", "Duplicados_Odoo": "7030A0",
               "Informativos": "808080", "SRI_Listado": "808080", "Odoo_Lineas": "808080"}
    for nombre, color in colores.items():
        assert wb[nombre].sheet_properties.tabColor.rgb.endswith(color), nombre
    assert wb.calculation.fullCalcOnLoad is True


def test_resultado_es_formula_no_valor(tmp_path):
    _, wb = _construir(tmp_path)
    ws = wb["Conciliacion"]
    col = [c.value for c in ws[1]].index("Resultado") + 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        assert isinstance(v, str) and v.startswith("=")
    assert "0.02" in ws.cell(row=2, column=col).value


def test_resumen_cuenta_por_tipo(tmp_path):
    res, wb = _construir(tmp_path)
    assert res["FAC"] == {"coinciden": 2, "solo_sri": 0, "solo_odoo": 0, "diferencias": 0, "informativos": 0}
    assert res["NC"] == {"coinciden": 0, "solo_sri": 1, "solo_odoo": 0, "diferencias": 0, "informativos": 0}
    assert res["ND"] == {"coinciden": 0, "solo_sri": 0, "solo_odoo": 0, "diferencias": 0, "informativos": 0}
    formulas = [c.value for row in wb["Resumen"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=COUNTIFS")]
    assert any("Conciliacion!" in f for f in formulas)


def test_diferencia_solo_odoo_e_informativo(tmp_path):
    sri, odoo, pares = _datos()
    sri.at[1, "total"] = 240.0
    f = pd.Timestamp("2026-09-07")
    odoo = pd.concat([odoo, _odoo([
        [12, "BILL/3", "001-001-000000009", "X", "1790000000001", f, "FAC", "01", "9" * 49, 5.0, 0.0, 5.0, pd.NA, True],
        [13, "BILL/4", "", "Y", "", f, "FAC", "02", pd.NA, 5.0, 0.0, 5.0, pd.NA, False],
    ])], ignore_index=True)
    pares = pd.concat([pares, pd.DataFrame({
        "idx_sri": pd.array([None, None], dtype="Int64"), "id_odoo": pd.array([12, 13], dtype="Int64"),
        "tipo": ["FAC", "FAC"], "metodo": ["", ""], "cat": [CAT_SOLO_ODOO, CAT_INFO]})], ignore_index=True)
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, str(tmp_path / "o.xlsx"))
    assert res["FAC"] == {"coinciden": 1, "solo_sri": 0, "solo_odoo": 1, "diferencias": 1, "informativos": 1}


def test_texto_que_empieza_con_igual_no_se_guarda_como_formula(tmp_path):
    _, wb = _construir(tmp_path, proveedor='=HYPERLINK("http://x","click")')
    celdas = [c for ws in wb for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert celdas and all(c.data_type == "s" for c in celdas)


def test_fuente_arial(tmp_path):
    _, wb = _construir(tmp_path)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    assert c.font.name == "Arial", (ws.title, c.coordinate)
