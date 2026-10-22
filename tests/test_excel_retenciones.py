# -*- coding: utf-8 -*-
from datetime import date

import openpyxl
import pandas as pd
import pytest

import conciliar_retenciones as cr
from excel_compras import construir_excel

DESDE, HASTA = date(2026, 9, 1), date(2026, 9, 30)
HOJAS_RET = ["Ret_Conciliacion", "Ret_Solo_SRI", "Ret_Solo_Odoo", "Ret_Diferencia_Valor",
             "Ret_Factura_Incorrecta", "Ret_Informativos"]
COLORES_RET = {"Ret_Conciliacion": "1F4E79", "Ret_Solo_SRI": "ED7D31", "Ret_Solo_Odoo": "C00000",
               "Ret_Diferencia_Valor": "FFC000", "Ret_Factura_Incorrecta": "7030A0",
               "Ret_Informativos": "808080"}
F = pd.Timestamp("2026-09-05")


def _fila(estado, n_ret="001-200-000000001", proveedor="PROV SA", fac_s="001-001-000000010",
          fac_o="001-001-000000010", pct_s=1.0, pct_o=1.0, vs=10.0, vo=10.0, base_s=1000.0, base_o=1000.0,
          ruc="1790000000001", id_ret=5, clave="1" * 49):
    ambos = vs is not None and vo is not None
    return {"estado": estado, "n_ret": n_ret, "fecha": F, "ruc": ruc, "proveedor": proveedor,
            "factura_sri": fac_s, "factura_odoo": fac_o, "pct_sri": pct_s, "pct_odoo": pct_o,
            "base_sri": base_s if vs is not None else None, "base_odoo": base_o if vo is not None else None,
            "valor_sri": vs, "valor_odoo": vo, "dif_valor": round(vs - vo, 2) if ambos else None,
            "id_ret": id_ret if vo is not None else None, "clave": clave if vs is not None else None}


def _ret(**kw):
    filas = [
        _fila(cr.EST_OK, proveedor=kw.get("proveedor", "PROV SA")),
        _fila(cr.EST_DIF, n_ret="001-200-000000002", vs=12.0, vo=10.0),
        _fila(cr.EST_DIF, n_ret="001-200-000000003", pct_s=1.0, pct_o=2.0, vs=10.0, vo=10.0,
              base_s=1000.0, base_o=500.0),
        _fila(cr.EST_SOLO_SRI, n_ret="001-200-000000004", fac_o=None, vo=None, pct_o=None),
        _fila(cr.EST_SOLO_ODOO, n_ret="001-200-000000005", fac_s=None, vs=None, pct_s=None),
        _fila(cr.EST_FAC, n_ret="001-200-000000006", fac_o="002-001-000000010"),
        _fila(cr.EST_XML, n_ret="001-200-000000007", fac_s=None, vs=None, pct_s=None),
        _fila(cr.EST_INFO, n_ret="001-200-000000008", fac_s=None, vs=None, pct_s=None),
    ]
    df = pd.DataFrame(filas, columns=cr.COLUMNAS)
    df["id_ret"] = pd.to_numeric(df["id_ret"], errors="coerce").astype("Int64")
    return df


def _vacios():
    sri = pd.DataFrame(columns=["clave", "serie", "ruc", "proveedor", "fecha", "tipo", "base", "iva",
                                "total", "afectada"])
    import odoo_client
    odoo = pd.DataFrame(columns=odoo_client.COLUMNAS)
    pares = pd.DataFrame({"idx_sri": pd.array([], dtype="Int64"), "id_odoo": pd.array([], dtype="Int64"),
                          "tipo": [], "metodo": [], "cat": []})
    return sri, odoo, pares


def _construir(tmp_path, ret=None, **kw):
    sri, odoo, pares = _vacios()
    salida = str(tmp_path / "r.xlsx")
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, salida, ret=_ret(**kw) if ret is None else ret)
    return res, openpyxl.load_workbook(salida), salida


def test_hojas_ret_y_colores(tmp_path):
    _, wb, _ = _construir(tmp_path)
    for h in HOJAS_RET:
        assert h in wb.sheetnames
        assert wb[h].sheet_properties.tabColor.rgb.endswith(COLORES_RET[h]), h


def test_ret_estado_es_formula(tmp_path):
    _, wb, _ = _construir(tmp_path)
    ws = wb["Ret_Conciliacion"]
    col = [c.value for c in ws[1]].index("Estado") + 1
    estados = [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]
    assert estados[0].startswith("=") and "0.02" in estados[0]
    assert estados[1].startswith("=")
    assert estados[2] == cr.EST_DIF  # pareja por porcentaje distinto: texto fijo
    assert estados[3:] == [cr.EST_SOLO_SRI, cr.EST_SOLO_ODOO, cr.EST_FAC, cr.EST_XML, cr.EST_INFO]


def test_ret_resumen_en_hoja_resumen(tmp_path):
    res, wb, _ = _construir(tmp_path)
    assert res["retenciones"] == {"coinciden": 1, "diferencia": 2, "solo_sri": 1, "solo_odoo": 1,
                                  "factura_incorrecta": 1, "informativos": 2, "total_sri": 52.0,
                                  "total_odoo": 60.0}
    textos = [c.value for row in wb["Resumen"].iter_rows() for c in row if isinstance(c.value, str)]
    cf = [t for t in textos if t.startswith("=COUNTIFS") and "Ret_Conciliacion!" in t]
    for est in (cr.EST_OK, cr.EST_DIF, cr.EST_SOLO_SRI, cr.EST_SOLO_ODOO, cr.EST_INFO, cr.EST_XML):
        assert any(f'"{est}"' in t for t in cf), est
    assert any("Retenciones" == t for t in textos)


def test_ret_texto_peligroso_no_es_formula(tmp_path):
    _, wb, _ = _construir(tmp_path, proveedor='=HYPERLINK("http://x","click")')
    celdas = [c for h in HOJAS_RET for row in wb[h].iter_rows() for c in row
              if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert celdas and all(c.data_type == "s" for c in celdas)


def test_ret_fuente_arial(tmp_path):
    _, wb, _ = _construir(tmp_path)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    assert c.font.name == "Arial", (ws.title, c.coordinate)


def test_ret_diferencia_muestra_ambos_porcentajes(tmp_path):
    _, wb, _ = _construir(tmp_path)
    ws = wb["Ret_Diferencia_Valor"]
    enc = [c.value for c in ws[1]]
    for h in ("Base_SRI", "Base_Odoo", "Pct_SRI", "Pct_Odoo"):
        assert h in enc
    filas = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert len(filas) == 2
    f = [r for r in filas if r[enc.index("N_Retencion")] == "001-200-000000003"][0]
    assert (f[enc.index("Pct_SRI")], f[enc.index("Pct_Odoo")]) == (1.0, 2.0)
    assert (f[enc.index("Base_SRI")], f[enc.index("Base_Odoo")]) == (1000.0, 500.0)


def test_construir_excel_sin_ret_no_cambia(tmp_path):
    sri, odoo, pares = _vacios()
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, str(tmp_path / "n.xlsx"))
    assert "retenciones" not in res
    wb = openpyxl.load_workbook(str(tmp_path / "n.xlsx"))
    assert wb.sheetnames == ["Resumen", "Conciliacion", "Solo_en_SRI", "Solo_en_Odoo", "Diferencia_Valor",
                             "Duplicados_Odoo", "Informativos", "SRI_Listado", "Odoo_Lineas"]


def test_construir_excel_solo_retenciones(tmp_path):
    res, wb, _ = _construir(tmp_path)
    assert res["retenciones"]["coinciden"] == 1
    rs = wb["Resumen"]
    for row in rs.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=SUM(") and "!" not in c.value:
                a, b = c.value[5:-1].split(":")
                assert not (int("".join(ch for ch in a if ch.isdigit())) <= c.row
                            <= int("".join(ch for ch in b if ch.isdigit())))


def test_ret_vacio(tmp_path):
    res, wb, _ = _construir(tmp_path, ret=pd.DataFrame(columns=cr.COLUMNAS))
    assert all(v == 0 for v in res["retenciones"].values())
    assert set(HOJAS_RET) <= set(wb.sheetnames)


def test_formulas_evaluan_sin_errores(tmp_path):
    formulas = pytest.importorskip("formulas")
    _, _, salida = _construir(tmp_path)
    xl = formulas.ExcelModel().loads(salida).finish()
    sol = xl.calculate()
    errores = [(k, v.value[0, 0]) for k, v in sol.items()
               if hasattr(v, "value") and v.value.size and isinstance(v.value[0, 0], formulas.functions.XlError)]
    assert not errores, errores[:5]
