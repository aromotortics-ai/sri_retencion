# -*- coding: utf-8 -*-
"""Pipeline de Ventas (comprobantes EMITIDOS SRI vs Odoo por API).

TXT sintéticos de punta a punta: lectura real de TXT, clasificación y Excel reales; Odoo, descarga
de XML, caché y LibreOffice falsos (nunca red, nunca Odoo, nunca el caché real)."""
import json
from datetime import date

import openpyxl
import pandas as pd

import app.odoo.client as odoo_client
from _helpers import _autorizacion
from app.sri.emitidos_sri import RUC_EMPRESA
from app.job_manager import create_job, get_job_dir, read_status
from app.odoo.client import OdooError

COLS_TXT = ["COMPROBANTE", "SERIE_COMPROBANTE", "CLAVE_ACCESO", "FECHA_AUTORIZACION", "FECHA_EMISION",
            "VALOR_SIN_IMPUESTOS", "IVA", "IMPORTE_TOTAL"]
CLAVE_R1 = "1" * 49
CLAVE_R2 = "2" * 49
N_RET1 = "001-200-000004340"  # el N° que trae el XML sintético de _helpers
N_RET2 = "001-200-000004341"
ADV_LIBRE = "No se pudo recalcular con LibreOffice; Excel recalculará las fórmulas al abrir el archivo"


def _clave(n):
    return str(n).zfill(49)


def _fila(comp, serie, clave, fecha="05/09/2026", base="100.00", iva="15.00", total="115.00"):
    return [comp, serie, clave, fecha, fecha, base, iva, total]


def _fac(n=1, fecha="05/09/2026"):
    return _fila("Factura", f"001-001-{n:09d}", _clave(100 + n), fecha)


def _nc(n=1):
    return _fila("Nota de Crédito", f"001-002-{n:09d}", _clave(200 + n), base="10.00", iva="1.50", total="")


def _liq(n=1):
    return _fila("Liquidación de Compra", f"001-003-{n:09d}", _clave(300 + n))


def _ret(clave=CLAVE_R1, serie=N_RET1, fecha="05/09/2026"):
    return _fila("Comprobante de Retención", serie, clave, fecha, base="", iva="", total="")


def _escribir_txt(path, filas):
    lineas = ["\t".join(COLS_TXT)] + ["\t".join(f) for f in filas]
    path.write_bytes(("\n".join(lineas) + "\n").encode("iso-8859-1"))


def _xml(cache, clave):
    """XML v2 de una retención emitida por la empresa (N° 001-200-000004340, 1.75% de 100)."""
    cache.mkdir(parents=True, exist_ok=True)
    cuerpo = _autorizacion(clave, "2.0.0").replace("<ruc>1790000000001</ruc>", f"<ruc>{RUC_EMPRESA}</ruc>")
    (cache / f"{clave}.xml").write_text('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + cuerpo,
                                        encoding="utf-8")


def _doc_odoo(id_, tipo, nfac, clave, total=115.0, base=100.0):
    return {"id": id_, "num": f"{tipo}/{id_}", "nfac": nfac, "socio": "CLIENTE SA", "ruc": "1790000000001",
            "fecha": pd.Timestamp("2026-09-05"), "tipo": tipo, "clave": clave, "base": base,
            "iva": total - base, "total": total, "afectada": pd.NA, "estado": "posted", "electronico": True,
            "elegible": True}


def _linea_odoo(n_ret=N_RET1, id_ret=10):
    return {"id_ret": id_ret, "n_ret": n_ret, "fecha": pd.Timestamp("2026-09-05"), "estado": "posted",
            "socio": "PROVEEDOR SA", "ruc": RUC_EMPRESA, "factura": "001-003-000088766", "pct": 1.75,
            "base": 100.0, "valor": 1.75, "cod": "312", "clave": CLAVE_R1}


class _Odoo:
    """Lector falso: registra las ventanas pedidas y el orden de las llamadas."""

    def __init__(self, docs=None, ret=None, orden=None):
        self._docs = docs or []
        self._ret = ret if ret is not None else [_linea_odoo()]
        self.orden = orden if orden is not None else []
        self.emitidos = None
        self.retenciones = None

    def fetch_emitidos(self, desde, hasta, margen_dias=5):
        self.orden.append("fetch_emitidos")
        self.emitidos = (desde, hasta)
        return pd.DataFrame(self._docs, columns=odoo_client.COLUMNAS_EMITIDOS)

    def fetch_retenciones_emitidas(self, desde, hasta, margen_dias=5):
        self.orden.append("fetch_retenciones_emitidas")
        self.retenciones = (desde, hasta)
        return pd.DataFrame(self._ret, columns=odoo_client.COLUMNAS_RET + ["clave"])


def _entorno(tmp_path, monkeypatch, txts, lector=None, descarga=None):
    monkeypatch.setattr("app.job_manager.JOBS_DIR", tmp_path / "jobs")
    import app.ventas.pipeline_ventas as pv
    job_id = create_job()
    carpeta = get_job_dir(job_id) / "emitidos"
    carpeta.mkdir()
    for nombre, filas in txts.items():
        _escribir_txt(carpeta / nombre, filas)
    monkeypatch.setattr(pv, "schedule_cleanup", lambda *a, **k: None)
    monkeypatch.setattr("shutil.which", lambda *_: None)  # sin LibreOffice: Excel tal cual
    cache = tmp_path / "cache" / "xml"
    monkeypatch.setattr(pv, "CACHE_XML_DIR", cache)
    lector = lector or _Odoo()
    monkeypatch.setattr(pv.OdooReader, "from_env", classmethod(lambda cls: lector))
    llamadas = []

    def descargar_falso(claves, cache_dir, on_progress=None, **kw):
        lector.orden.append("descargar_xmls")
        llamadas.append(list(claves))
        assert str(cache_dir) == str(cache)
        return descarga(claves, cache_dir, on_progress) if descarga else []
    monkeypatch.setattr(pv, "descargar_xmls", descargar_falso)
    return pv, job_id, lector, cache, llamadas


def _espiar_status(pv, monkeypatch):
    vistos = []
    original = pv.write_status

    def espia(job_id, estado, mensaje, progreso=0.0, **kw):
        vistos.append((mensaje, round(progreso, 3)))
        return original(job_id, estado, mensaje, progreso, **kw)
    monkeypatch.setattr(pv, "write_status", espia)
    return vistos


def _todos():
    return {"2390000831001_EmitidosAutorizados (1).txt": [_fac(1), _fac(2), _nc(1), _liq(1)],
            "otro nombre.txt": [_ret(), _fila("Guía de Remisión", "001-004-000000001", _clave(900))]}


def test_pipeline_done_con_todos_los_tipos(tmp_path, monkeypatch):
    lector = _Odoo(docs=[_doc_odoo(1, "FAC", "001-001-000000001", _clave(101))])
    pv, job_id, lector, cache, llamadas = _entorno(tmp_path, monkeypatch, _todos(), lector=lector)
    _xml(cache, CLAVE_R1)
    vistos = _espiar_status(pv, monkeypatch)
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert (get_job_dir(job_id) / "resultado.xlsx").exists()
    r = s["resumen"]
    assert {"FAC", "NC", "LIQ", "RET"} <= set(r["por_tipo"])
    assert r["por_tipo"]["FAC"]["Coincide"] == 1
    assert r["por_tipo"]["FAC"]["Solo en SRI"] == 1
    assert r["por_tipo"]["RET"]["Coincide"] == 1
    assert r["periodo"] == ["2026-09-05", "2026-09-05"]
    assert r["comprobantes_sri"] == {"FAC": 2, "NC": 1, "ND": 0, "LIQ": 1, "RET": 1}
    assert r["retenciones_fallidos"] == 0
    assert r["recalculado"] is False
    assert any("guías de remisión" in a for a in r["advertencias"])
    assert ADV_LIBRE in r["advertencias"]
    json.dumps(r)  # serializable
    assert lector.emitidos == (date(2026, 9, 5), date(2026, 9, 5))
    assert lector.retenciones == (date(2026, 9, 5), date(2026, 9, 5))
    assert llamadas == [[CLAVE_R1]]
    hojas = openpyxl.load_workbook(get_job_dir(job_id) / "resultado.xlsx").sheetnames
    assert "Conciliacion" in hojas and "Ret_Documentos" in hojas
    progresos = [p for _, p in vistos]
    assert progresos == sorted(progresos)
    assert vistos[-1] == ("Completado", 1.0)


def test_pipeline_sin_txt_validos(tmp_path, monkeypatch):
    guia = _fila("Guía de Remisión", "001-004-000000001", _clave(900))
    pv, job_id, lector, _, llamadas = _entorno(tmp_path, monkeypatch, {"x.txt": [guia]})
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert ("Los TXT no contienen facturas, notas de crédito, notas de débito, liquidaciones de compra "
            "ni retenciones" in s["mensaje"])
    assert lector.orden == [] and llamadas == []


def test_pipeline_sin_ningun_txt_es_error(tmp_path, monkeypatch):
    pv, job_id, lector, _, _ = _entorno(tmp_path, monkeypatch, {})
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert "Los TXT no contienen" in s["mensaje"]


def test_pipeline_odoo_falla_antes_de_descargar_xml(tmp_path, monkeypatch):
    class Caido(_Odoo):
        def fetch_retenciones_emitidas(self, desde, hasta, margen_dias=5):
            raise OdooError("No se pudo conectar con Odoo")
    pv, job_id, lector, _, llamadas = _entorno(tmp_path, monkeypatch, _todos(), lector=Caido())
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert s["mensaje"].startswith("OdooError")
    assert llamadas == []


def test_pipeline_sin_credenciales_no_descarga(tmp_path, monkeypatch):
    pv, job_id, _, _, llamadas = _entorno(tmp_path, monkeypatch, _todos())

    def sin_env(cls):
        raise OdooError("Falta configurar ODOO_URL")
    monkeypatch.setattr(pv.OdooReader, "from_env", classmethod(sin_env))
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error" and "ODOO_URL" in s["mensaje"]
    assert llamadas == []


def test_pipeline_lee_odoo_antes_de_descargar(tmp_path, monkeypatch):
    pv, job_id, lector, _, _ = _entorno(tmp_path, monkeypatch, _todos())
    pv.run_pipeline_ventas(job_id)
    assert read_status(job_id)["estado"] == "done"
    assert lector.orden == ["fetch_emitidos", "fetch_retenciones_emitidas", "descargar_xmls"]


def test_pipeline_xml_todos_fallidos(tmp_path, monkeypatch):
    lector = _Odoo(ret=[_linea_odoo(), _linea_odoo(n_ret=N_RET2, id_ret=11)])

    def descarga(claves, cache_dir, on_progress):
        on_progress(len(claves), len(claves), len(claves))
        return [{"clave": c, "motivo": "NO AUTORIZADO"} for c in claves]
    pv, job_id, _, _, _ = _entorno(
        tmp_path, monkeypatch, {"r.txt": [_ret(), _ret(CLAVE_R2, N_RET2)]}, lector=lector, descarga=descarga)
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    r = s["resumen"]
    assert r["retenciones_fallidos"] == 2
    assert ("2 XML de retenciones no se pudieron descargar o leer; sus retenciones se listan como "
            "'XML no disponible'") in r["advertencias"]
    assert "Motivo de las descargas fallidas: NO AUTORIZADO (×2)" in r["advertencias"]
    assert r["por_tipo"]["RET"]["XML no disponible"] == 2
    json.dumps(r)


def test_pipeline_solo_documentos_no_descarga_xml(tmp_path, monkeypatch):
    pv, job_id, lector, _, llamadas = _entorno(tmp_path, monkeypatch, {"f.txt": [_fac(1), _nc(1)]})
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert llamadas == []
    assert lector.retenciones is None
    assert "RET" not in s["resumen"]["por_tipo"]
    assert s["resumen"]["retenciones_fallidos"] == 0
    assert s["resumen"]["comprobantes_sri"]["RET"] == 0
    hojas = openpyxl.load_workbook(get_job_dir(job_id) / "resultado.xlsx").sheetnames
    assert not any(h.startswith("Ret_") for h in hojas)


def test_pipeline_solo_retenciones_no_lee_documentos_de_odoo(tmp_path, monkeypatch):
    pv, job_id, lector, cache, _ = _entorno(tmp_path, monkeypatch, {"r.txt": [_ret()]})
    _xml(cache, CLAVE_R1)
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert lector.emitidos is None
    assert s["resumen"]["por_tipo"]["RET"]["Coincide"] == 1
    assert s["resumen"]["por_tipo"]["FAC"]["total"] == 0


ADV_FALTA_FAC = ("Falta el TXT de facturas: los 1 registros de Odoo de ese tipo aparecerán como "
                 "'Solo en Odoo'")


def test_pipeline_advierte_tipo_sin_txt_con_registros_en_odoo(tmp_path, monkeypatch):
    fuera = dict(_doc_odoo(3, "LIQ", "001-003-000000009", _clave(309)), fecha=pd.Timestamp("2026-09-30"))
    borrador = dict(_doc_odoo(4, "ND", "001-005-000000009", _clave(409)), estado="draft", elegible=False)
    lector = _Odoo(docs=[_doc_odoo(1, "FAC", "001-001-000000001", _clave(101)),
                         _doc_odoo(2, "NC", "001-002-000000001", _clave(201), total=11.5, base=10.0),
                         fuera, borrador])
    pv, job_id, _, _, _ = _entorno(tmp_path, monkeypatch, {"nc.txt": [_nc(1)]}, lector=lector)
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    adv = s["resumen"]["advertencias"]
    assert ADV_FALTA_FAC in adv
    assert not any(a.startswith("Falta el TXT de notas de crédito") for a in adv)
    # fuera del período (LIQ) o borrador (ND): no se advierte
    assert not any("liquidaciones" in a or "notas de débito" in a for a in adv)
    assert s["resumen"]["por_tipo"]["FAC"]["Solo en Odoo"] == 1


def test_pipeline_sin_advertencia_de_tipo_si_hay_txt(tmp_path, monkeypatch):
    lector = _Odoo(docs=[_doc_odoo(1, "FAC", "001-001-000000001", _clave(101))])
    pv, job_id, _, _, _ = _entorno(tmp_path, monkeypatch, {"f.txt": [_fac(1)]}, lector=lector)
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert not any(a.startswith("Falta el TXT") for a in s["resumen"]["advertencias"])


def test_pipeline_rango_amplio_advierte(tmp_path, monkeypatch):
    pv, job_id, lector, _, _ = _entorno(
        tmp_path, monkeypatch, {"f.txt": [_fac(1, "01/07/2026"), _fac(2, "20/09/2026")]})
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert s["resumen"]["periodo"] == ["2026-07-01", "2026-09-20"]
    assert lector.emitidos == (date(2026, 7, 1), date(2026, 9, 20))
    assert any("abarcan 81 días" in a for a in s["resumen"]["advertencias"])


def test_pipeline_retenciones_sin_fechas_validas_es_error(tmp_path, monkeypatch):
    pv, job_id, lector, _, llamadas = _entorno(
        tmp_path, monkeypatch, {"f.txt": [_fac(1)], "r.txt": [_ret(fecha="")]})
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert "fechas de emisión válidas" in s["mensaje"]
    assert llamadas == [] and lector.orden == []


def test_pipeline_progreso_de_descarga(tmp_path, monkeypatch):
    def descarga(claves, cache_dir, on_progress):
        on_progress(1, 2, 0)
        on_progress(2, 2, 1)
        return [{"clave": CLAVE_R2, "motivo": "x"}]
    pv, job_id, _, _, _ = _entorno(
        tmp_path, monkeypatch, {"r.txt": [_ret(), _ret(CLAVE_R2, N_RET2)]}, descarga=descarga)
    vistos = _espiar_status(pv, monkeypatch)
    pv.run_pipeline_ventas(job_id)
    assert read_status(job_id)["estado"] == "done"
    assert ("Descargando XML de retenciones (1/2, 0 fallidos)", 0.4) in vistos
    assert ("Descargando XML de retenciones (2/2, 1 fallidos)", 0.65) in vistos


def test_pipeline_con_libreoffice_sin_advertencia(tmp_path, monkeypatch):
    pv, job_id, _, _, _ = _entorno(tmp_path, monkeypatch, {"f.txt": [_fac(1)]})
    import app.ventas.excel_emitidos as excel_emitidos
    monkeypatch.setattr(excel_emitidos, "_recalcular", lambda ruta: True)
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "done", s
    assert s["resumen"]["recalculado"] is True
    assert ADV_LIBRE not in s["resumen"]["advertencias"]


def test_pipeline_con_error_agenda_limpieza(tmp_path, monkeypatch):
    pv, job_id, _, _, _ = _entorno(tmp_path, monkeypatch, {})
    llamadas = []
    monkeypatch.setattr(pv, "schedule_cleanup", lambda *a, **k: llamadas.append((a, k)))
    pv.run_pipeline_ventas(job_id)
    s = read_status(job_id)
    assert s["estado"] == "error"
    assert s["error"]  # traceback
    assert llamadas == [((job_id,), {"delay_seconds": 7200})]
