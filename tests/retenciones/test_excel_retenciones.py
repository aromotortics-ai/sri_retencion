# -*- coding: utf-8 -*-
from datetime import date

import openpyxl
import pandas as pd
import pytest

import app.retenciones.conciliar_retenciones as cr
from openpyxl.utils import get_column_letter

from app.compras.excel_compras import construir_excel

DESDE, HASTA = date(2026, 9, 1), date(2026, 9, 30)
HOJAS_RET = ["Ret_Documentos", "Ret_Conciliacion", "Ret_Solo_SRI", "Ret_Solo_Odoo", "Ret_Diferencia_Valor",
             "Ret_Factura_Incorrecta", "Ret_Informativos"]
COLORES_RET = {"Ret_Documentos": "1F4E79", "Ret_Conciliacion": "1F4E79", "Ret_Solo_SRI": "ED7D31", "Ret_Solo_Odoo": "C00000",
               "Ret_Diferencia_Valor": "FFC000", "Ret_Factura_Incorrecta": "7030A0",
               "Ret_Informativos": "808080"}
F = pd.Timestamp("2026-09-05")


def _fila(estado, n_ret="001-200-000000001", proveedor="PROV SA", fac_s="001-001-000000010",
          fac_o="001-001-000000010", pct_s=1.0, pct_o=1.0, vs=10.0, vo=10.0, base_s=1000.0, base_o=1000.0,
          ruc="1790000000001", id_ret=5, clave="1" * 49):
    ambos = vs is not None and vo is not None
    return {"estado": estado, "n_ret": n_ret, "fecha": F, "ruc": ruc, "proveedor": proveedor,
            "factura_sri": fac_s, "factura_odoo": fac_o, "pct_sri": pct_s, "pct_odoo": pct_o,
            "base_sri": base_s if vs is not None else None, "base_odoo": base_o if vo is not None else None,
            "valor_sri": vs, "valor_odoo": vo, "dif_valor": round(vs - vo, 2) if ambos else None,
            "id_ret": id_ret if vo is not None else None, "clave": clave if vs is not None else None}


def _ret(**kw):
    filas = [
        _fila(cr.EST_OK, proveedor=kw.get("proveedor", "PROV SA")),
        _fila(cr.EST_DIF, n_ret="001-200-000000002", vs=12.0, vo=10.0),
        _fila(cr.EST_DIF, n_ret="001-200-000000003", pct_s=1.0, pct_o=2.0, vs=10.0, vo=10.0,
              base_s=1000.0, base_o=500.0),
        _fila(cr.EST_SOLO_SRI, n_ret="001-200-000000004", fac_o=None, vo=None, pct_o=None),
        _fila(cr.EST_SOLO_ODOO, n_ret="001-200-000000005", fac_s=None, vs=None, pct_s=None),
        _fila(cr.EST_FAC, n_ret="001-200-000000006", fac_o="002-001-000000010"),
        _fila(cr.EST_XML, n_ret="001-200-000000007", fac_s=None, vs=None, pct_s=None),
        _fila(cr.EST_INFO, n_ret="001-200-000000008", fac_s=None, vs=None, pct_s=None),
    ]
    df = pd.DataFrame(filas, columns=cr.COLUMNAS)
    df["id_ret"] = pd.to_numeric(df["id_ret"], errors="coerce").astype("Int64")
    return df


def _vacios():
    sri = pd.DataFrame(columns=["clave", "serie", "ruc", "proveedor", "fecha", "tipo", "base", "iva",
                                "total", "afectada"])
    import app.odoo.client as odoo_client
    odoo = pd.DataFrame(columns=odoo_client.COLUMNAS)
    pares = pd.DataFrame({"idx_sri": pd.array([], dtype="Int64"), "id_odoo": pd.array([], dtype="Int64"),
                          "tipo": [], "metodo": [], "cat": []})
    return sri, odoo, pares


def _construir(tmp_path, ret=None, **kw):
    sri, odoo, pares = _vacios()
    salida = str(tmp_path / "r.xlsx")
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, salida, ret=_ret(**kw) if ret is None else ret)
    return res, openpyxl.load_workbook(salida), salida


def test_hojas_ret_y_colores(tmp_path):
    _, wb, _ = _construir(tmp_path)
    for h in HOJAS_RET:
        assert h in wb.sheetnames
        assert wb[h].sheet_properties.tabColor.rgb.endswith(COLORES_RET[h]), h


def test_ret_estado_es_formula(tmp_path):
    _, wb, _ = _construir(tmp_path)
    ws = wb["Ret_Conciliacion"]
    col = [c.value for c in ws[1]].index("Estado") + 1
    estados = [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]
    assert estados[0].startswith("=") and "0.02" in estados[0]
    assert estados[1].startswith("=")
    assert estados[2] == cr.EST_DIF  # pareja por porcentaje distinto: texto fijo
    assert estados[3:] == [cr.EST_SOLO_SRI, cr.EST_SOLO_ODOO, cr.EST_FAC, cr.EST_XML, cr.EST_INFO]


def test_ret_resumen_en_hoja_resumen(tmp_path):
    res, wb, _ = _construir(tmp_path)
    assert res["retenciones"] == {"coinciden": 1, "diferencia": 2, "solo_sri": 1, "solo_odoo": 1,
                                  "factura_incorrecta": 1, "informativos": 2, "total_sri": 52.0,
                                  "total_odoo": 60.0, "documentos": 8, "lineas": 8}
    textos = [c.value for row in wb["Resumen"].iter_rows() for c in row if isinstance(c.value, str)]
    # Los conteos del Resumen son por DOCUMENTO: COUNTIFS sobre Ret_Documentos.
    cf = [t for t in textos if t.startswith("=COUNTIFS") and "Ret_Documentos!" in t]
    for est in (cr.EST_OK, cr.EST_DIF, cr.EST_SOLO_SRI, cr.EST_SOLO_ODOO, cr.EST_INFO, cr.EST_XML):
        assert any(f'"{est}"' in t for t in cf), est
    assert any("Retenciones" == t for t in textos)


def test_ret_texto_peligroso_no_es_formula(tmp_path):
    _, wb, _ = _construir(tmp_path, proveedor='=HYPERLINK("http://x","click")')
    celdas = [c for h in HOJAS_RET for row in wb[h].iter_rows() for c in row
              if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert celdas and all(c.data_type == "s" for c in celdas)


def test_ret_fuente_arial(tmp_path):
    _, wb, _ = _construir(tmp_path)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    assert c.font.name == "Arial", (ws.title, c.coordinate)


def test_ret_diferencia_muestra_ambos_porcentajes(tmp_path):
    _, wb, _ = _construir(tmp_path)
    ws = wb["Ret_Diferencia_Valor"]
    enc = [c.value for c in ws[1]]
    for h in ("Base_SRI", "Base_Odoo", "Pct_SRI", "Pct_Odoo"):
        assert h in enc
    filas = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert len(filas) == 2
    f = [r for r in filas if r[enc.index("N_Retencion")] == "001-200-000000003"][0]
    assert (f[enc.index("Pct_SRI")], f[enc.index("Pct_Odoo")]) == (1.0, 2.0)
    assert (f[enc.index("Base_SRI")], f[enc.index("Base_Odoo")]) == (1000.0, 500.0)


def test_ret_diferencia_es_formula(tmp_path):
    _, wb, _ = _construir(tmp_path)
    ws = wb["Ret_Diferencia_Valor"]
    enc = [c.value for c in ws[1]]
    ci, cj, ck = (enc.index(h) + 1 for h in ("Valor_SRI", "Valor_Odoo", "Diferencia"))
    assert enc[0] == "Estado_Linea"
    assert (ci, cj, ck) == (10, 11, 12)
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=ck)
        assert c.value == f"=ROUND(J{r}-K{r},2)" and c.data_type == "f", c.value
        assert c.number_format == "#,##0.00"


def test_ret_diferencia_formula_evalua(tmp_path):
    formulas = pytest.importorskip("formulas")
    _, wb, salida = _construir(tmp_path)
    ws = wb["Ret_Diferencia_Valor"]
    esperados = {r: round(ws[f"J{r}"].value - ws[f"K{r}"].value, 2) for r in range(2, ws.max_row + 1)}
    sol = formulas.ExcelModel().loads(salida).finish().calculate()
    for r, v in esperados.items():
        k = [k for k in sol if k.upper().endswith(f"RET_DIFERENCIA_VALOR'!L{r}")]
        assert k, r
        assert float(sol[k[0]].value[0, 0]) == pytest.approx(v)


def test_ret_xml_sin_gemelo_contadores_consistentes(tmp_path):
    ret = _ret()
    solo_xml = _fila(cr.EST_XML, n_ret="001-200-000000009", fac_s=None, fac_o=None, vs=None, vo=None,
                     pct_s=None, pct_o=None)
    solo_xml["clave"] = "1" * 49  # fila del lado SRI (TXT): clave sin valores
    ret = pd.concat([ret, pd.DataFrame([solo_xml], columns=cr.COLUMNAS)], ignore_index=True)
    ret["id_ret"] = pd.to_numeric(ret["id_ret"], errors="coerce").astype("Int64")
    res, wb, _ = _construir(tmp_path, ret=ret)
    n_info = wb["Ret_Informativos"].max_row - 1
    assert res["retenciones"]["informativos"] == n_info == 3
    assert res["retenciones"]["total_odoo"] == 60.0
    enc = [c.value for c in wb["Ret_Informativos"][1]]
    filas = list(wb["Ret_Informativos"].iter_rows(min_row=2, values_only=True))
    assert any(f[enc.index("N_Retencion")] == "001-200-000000009" and f[enc.index("Clave")] == "1" * 49
               for f in filas)


def _textos_resumen(wb):
    return [c.value for row in wb["Resumen"].iter_rows() for c in row if isinstance(c.value, str)]


def test_ret_periodo_distinto_se_rotula(tmp_path):
    sri, odoo, pares = _vacios()
    salida = str(tmp_path / "p.xlsx")
    construir_excel(sri, odoo, pares, DESDE, HASTA, salida, ret=_ret(),
                    periodo_ret=(date(2026, 8, 1), date(2026, 8, 31)))
    wb = openpyxl.load_workbook(salida)
    rs = wb["Resumen"]
    fila = [c.row for c in rs["A"] if c.value == "Retenciones"][0]
    assert rs[f"B{fila}"].value == "2026-08"
    assert any("08/2026" in t for t in _textos_resumen(wb))


def test_ret_periodo_igual_no_agrega_nota(tmp_path):
    sri, odoo, pares = _vacios()
    salida = str(tmp_path / "p.xlsx")
    construir_excel(sri, odoo, pares, DESDE, HASTA, salida, ret=_ret(), periodo_ret=(DESDE, HASTA))
    wb = openpyxl.load_workbook(salida)
    assert not any("Período de retenciones" in t for t in _textos_resumen(wb))


def test_construir_excel_sin_ret_no_cambia(tmp_path):
    sri, odoo, pares = _vacios()
    res = construir_excel(sri, odoo, pares, DESDE, HASTA, str(tmp_path / "n.xlsx"))
    assert "retenciones" not in res
    wb = openpyxl.load_workbook(str(tmp_path / "n.xlsx"))
    assert wb.sheetnames == ["Resumen", "Conciliacion", "Solo_en_SRI", "Solo_en_Odoo", "Diferencia_Valor",
                             "Duplicados_Odoo", "Informativos", "SRI_Listado", "Odoo_Lineas"]


def test_construir_excel_solo_retenciones(tmp_path):
    res, wb, _ = _construir(tmp_path)
    assert res["retenciones"]["coinciden"] == 1
    rs = wb["Resumen"]
    for row in rs.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=SUM(") and "!" not in c.value:
                a, b = c.value[5:-1].split(":")
                assert not (int("".join(ch for ch in a if ch.isdigit())) <= c.row
                            <= int("".join(ch for ch in b if ch.isdigit())))


def test_ret_vacio(tmp_path):
    res, wb, _ = _construir(tmp_path, ret=pd.DataFrame(columns=cr.COLUMNAS))
    assert all(v == 0 for v in res["retenciones"].values())
    assert set(HOJAS_RET) <= set(wb.sheetnames)


def test_formulas_evaluan_sin_errores(tmp_path):
    formulas = pytest.importorskip("formulas")
    _, _, salida = _construir(tmp_path)
    xl = formulas.ExcelModel().loads(salida).finish()
    sol = xl.calculate()
    errores = [(k, v.value[0, 0]) for k, v in sol.items()
               if hasattr(v, "value") and v.value.size and isinstance(v.value[0, 0], formulas.functions.XlError)]
    assert not errores, errores[:5]


# ------------------------------------------------------------- vista por documento (multi-linea)
# Un comprobante (N° de retencion) con varias lineas (Renta + IVA) es UN documento.

def _ret_multi():
    A, B, C, D, E, I = (f"001-200-00000010{k}" for k in range(1, 7))
    filas = [
        _fila(cr.EST_OK, n_ret=A, pct_s=1.0, pct_o=1.0),
        _fila(cr.EST_OK, n_ret=B, pct_s=30.0, pct_o=30.0),
        _fila(cr.EST_SOLO_SRI, n_ret=C, fac_o=None, vo=None, pct_o=None, vs=7.0),
        _fila(cr.EST_OK, n_ret=A, pct_s=2.0, pct_o=2.0),
        _fila(cr.EST_DIF, n_ret=B, pct_s=1.75, pct_o=1.75, vs=12.0, vo=10.0),
        _fila(cr.EST_SOLO_ODOO, n_ret=D, fac_s=None, vs=None, pct_s=None, id_ret=9),
        _fila(cr.EST_SOLO_SRI, n_ret=C, pct_s=30.0, fac_o=None, vo=None, pct_o=None, vs=3.0),
        _fila(cr.EST_OK, n_ret=E, pct_s=1.75, pct_o=1.75),
        _fila(cr.EST_OK, n_ret=A, pct_s=30.0, pct_o=30.0),
        _fila(cr.EST_SOLO_SRI, n_ret=E, pct_s=30.0, fac_o=None, vo=None, pct_o=None, vs=4.5),
        _fila(cr.EST_SOLO_ODOO, n_ret=D, pct_o=30.0, fac_s=None, vs=None, pct_s=None, id_ret=9),
        _fila(cr.EST_INFO, n_ret=I, fac_s=None, vs=None, pct_s=None, id_ret=12),
    ]
    df = pd.DataFrame(filas, columns=cr.COLUMNAS)
    df["id_ret"] = pd.to_numeric(df["id_ret"], errors="coerce").astype("Int64")
    return df


ENC_DOC = ["Estado", "N_Retencion", "Fecha", "RUC", "Proveedor", "Factura_SRI", "Factura_Odoo", "Lineas_SRI",
           "Lineas_Odoo", "Valor_SRI", "Valor_Odoo", "Diferencia", "Detalle", "Documento"]


def _calc(salida):
    formulas = pytest.importorskip("formulas")
    sol = formulas.ExcelModel().loads(salida).finish().calculate()

    def val(hoja, celda):
        k = [k for k in sol if k.upper().endswith(f"{hoja.upper()}'!{celda}")]
        assert k, (hoja, celda)
        return sol[k[0]].value[0, 0]
    return val


def test_doc_hoja_documentos_una_fila_por_documento(tmp_path):
    _, wb, _ = _construir(tmp_path, ret=_ret_multi())
    ret_hojas = [h for h in wb.sheetnames if h.startswith("Ret_")]
    assert ret_hojas[0] == "Ret_Documentos"
    ws = wb["Ret_Documentos"]
    assert [c.value for c in ws[1]] == ENC_DOC
    assert ws.max_row - 1 == 6
    docs = [ws.cell(row=r, column=14).value for r in range(2, ws.max_row + 1)]
    assert len(set(docs)) == 6
    est = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert est["001-200-000000101"] == cr.EST_OK
    assert est["001-200-000000102"] == cr.EST_DIF
    assert est["001-200-000000103"] == cr.EST_SOLO_SRI
    assert est["001-200-000000104"] == cr.EST_SOLO_ODOO
    assert est["001-200-000000105"] == cr.EST_DIF
    assert est["001-200-000000106"] == cr.EST_INFO
    for r in range(2, ws.max_row + 1):
        for col in "HIJKL":
            assert ws[f"{col}{r}"].data_type == "f", (col, r)
        assert ws[f"C{r}"].number_format == "dd/mm/yyyy"
        assert ws[f"J{r}"].number_format == "#,##0.00"
    assert ws.freeze_panes and ws.auto_filter.ref


def test_doc_contadores_cuentan_documentos(tmp_path):
    res, _, _ = _construir(tmp_path, ret=_ret_multi())
    c = res["retenciones"]
    assert {k: c[k] for k in ("coinciden", "diferencia", "solo_sri", "solo_odoo", "factura_incorrecta",
                              "informativos", "documentos", "lineas")} == {
        "coinciden": 1, "diferencia": 2, "solo_sri": 1, "solo_odoo": 1, "factura_incorrecta": 0,
        "informativos": 1, "documentos": 6, "lineas": 12}
    assert c["total_sri"] == pytest.approx(10 * 3 + 10 + 12 + 7 + 3 + 10 + 4.5)
    assert c["total_odoo"] == pytest.approx(10 * 3 + 10 + 10 + 10 * 2 + 10)


def test_doc_lineas_contiguas_y_estado_documento_formula(tmp_path):
    _, wb, _ = _construir(tmp_path, ret=_ret_multi())
    ws = wb["Ret_Conciliacion"]
    enc = [c.value for c in ws[1]]
    assert enc[:16] == ["Estado", "N_Retencion", "Fecha", "RUC", "Proveedor", "Factura_SRI", "Factura_Odoo",
                        "Pct_SRI", "Pct_Odoo", "Base_SRI", "Base_Odoo", "Valor_SRI", "Valor_Odoo", "Diferencia",
                        "ID_Odoo", "Clave"]
    assert enc[16:] == ["Documento", "Estado_Documento"]
    docs = [ws[f"Q{r}"].value for r in range(2, ws.max_row + 1)]
    assert len(docs) == 12
    vistos = []
    for d in docs:
        if not vistos or vistos[-1] != d:
            assert d not in vistos, docs  # cada documento en un bloque contiguo
            vistos.append(d)
    for r in range(2, ws.max_row + 1):
        assert ws[f"R{r}"].data_type == "f" and "Ret_Documentos!" in ws[f"R{r}"].value


def test_doc_formulas_evaluan(tmp_path):
    _, wb, salida = _construir(tmp_path, ret=_ret_multi())
    val = _calc(salida)
    ws = wb["Ret_Documentos"]
    fila = {ws[f"B{r}"].value: r for r in range(2, ws.max_row + 1)}
    r = fila["001-200-000000101"]
    assert (val("Ret_Documentos", f"H{r}"), val("Ret_Documentos", f"I{r}")) == (3, 3)
    assert float(val("Ret_Documentos", f"J{r}")) == pytest.approx(30.0)
    r = fila["001-200-000000102"]
    assert float(val("Ret_Documentos", f"L{r}")) == pytest.approx(2.0)
    r = fila["001-200-000000103"]
    assert (val("Ret_Documentos", f"H{r}"), val("Ret_Documentos", f"I{r}")) == (2, 0)
    assert float(val("Ret_Documentos", f"J{r}")) == pytest.approx(10.0)
    r = fila["001-200-000000106"]  # documento solo informativo: muestra su valor de Odoo
    assert float(val("Ret_Documentos", f"K{r}")) == pytest.approx(10.0)
    wc = wb["Ret_Conciliacion"]
    for rr in range(2, wc.max_row + 1):
        assert val("Ret_Conciliacion", f"R{rr}") == ws[f"A{fila[wc[f'B{rr}'].value]}"].value


def test_doc_resumen_cuenta_documentos(tmp_path):
    _, wb, salida = _construir(tmp_path, ret=_ret_multi())
    rs = wb["Resumen"]
    r = [c.row for c in rs["A"] if c.value == "Retenciones"][0]
    assert "Ret_Documentos!" in rs[f"D{r}"].value
    assert "COUNTA(Ret_Conciliacion!" in rs[f"A{r + 1}"].value
    val = _calc(salida)
    got = {col: val("Resumen", f"{col}{r}") for col in "CDEFGI"}
    assert got == {"C": 4, "D": 1, "E": 2, "F": 1, "G": 1, "I": 1}
    assert float(val("Resumen", f"K{r}")) == pytest.approx(76.5)
    assert float(val("Resumen", f"L{r}")) == pytest.approx(80.0)
    assert "12" in str(val("Resumen", f"A{r + 1}"))


def test_doc_hojas_por_estado_una_fila_por_documento(tmp_path):
    _, wb, salida = _construir(tmp_path, ret=_ret_multi())
    assert wb["Ret_Solo_SRI"].max_row - 1 == 1
    assert wb["Ret_Solo_Odoo"].max_row - 1 == 1
    assert wb["Ret_Informativos"].max_row - 1 == 1
    assert wb["Ret_Factura_Incorrecta"].max_row - 1 == 0
    ws = wb["Ret_Solo_SRI"]
    enc = [c.value for c in ws[1]]
    assert "Lineas" in enc and "Valor" in enc and "Documento" in enc
    cl, cv = (get_column_letter(enc.index(h) + 1) for h in ("Lineas", "Valor"))
    assert ws[f"{cl}2"].data_type == "f" and ws[f"{cv}2"].data_type == "f"
    val = _calc(salida)
    assert val("Ret_Solo_SRI", f"{cl}2") == 2
    assert float(val("Ret_Solo_SRI", f"{cv}2")) == pytest.approx(10.0)


def test_doc_diferencia_lista_todas_las_lineas(tmp_path):
    _, wb, _ = _construir(tmp_path, ret=_ret_multi())
    ws = wb["Ret_Diferencia_Valor"]
    enc = [c.value for c in ws[1]]
    filas = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    nret = [f[enc.index("N_Retencion")] for f in filas]
    assert sorted(nret) == ["001-200-000000102"] * 2 + ["001-200-000000105"] * 2
    assert nret[0] == nret[1] and nret[2] == nret[3]  # agrupadas por documento
    estados = sorted(f[0] for f in filas)
    assert estados == sorted([cr.EST_OK, cr.EST_DIF, cr.EST_OK, cr.EST_SOLO_SRI])
    k = enc.index("Diferencia") + 1
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=k)
        if ws.cell(row=r, column=enc.index("Valor_Odoo") + 1).value in (None, ""):
            assert c.value in (None, "")  # linea de un solo lado: sin formula (evita #VALUE!)
        else:
            assert c.data_type == "f"


def test_doc_texto_peligroso_en_documentos_y_clave(tmp_path):
    ret = _ret_multi()
    ret["proveedor"] = '=HYPERLINK("http://x","click")'
    ret.loc[ret["n_ret"] == "001-200-000000104", "n_ret"] = "=1+1"
    _, wb, _ = _construir(tmp_path, ret=ret)
    celdas = [c for h in HOJAS_RET for row in wb[h].iter_rows() for c in row
              if isinstance(c.value, str) and ("HYPERLINK" in c.value or c.value == "=1+1")]
    assert any(c.parent.title == "Ret_Documentos" for c in celdas)
    assert any(c.parent.title == "Ret_Conciliacion" and c.column_letter == "Q" for c in celdas)
    assert celdas and all(c.data_type == "s" for c in celdas)


def test_doc_formulas_sin_errores_multilinea(tmp_path):
    formulas = pytest.importorskip("formulas")
    _, _, salida = _construir(tmp_path, ret=_ret_multi())
    sol = formulas.ExcelModel().loads(salida).finish().calculate()
    errores = [(k, v.value[0, 0]) for k, v in sol.items()
               if hasattr(v, "value") and v.value.size and isinstance(v.value[0, 0], formulas.functions.XlError)]
    assert not errores, errores[:5]


def test_doc_mismo_id_odoo_con_ruc_distinto_cuenta_un_documento(tmp_path):
    ret = pd.DataFrame([_fila(cr.EST_OK, ruc="1790000000001", id_ret=5),
                        _fila(cr.EST_SOLO_ODOO, ruc="1790000000", fac_s=None, vs=None, pct_s=None, pct_o=30.0,
                              id_ret=5)], columns=cr.COLUMNAS)
    ret["id_ret"] = pd.to_numeric(ret["id_ret"], errors="coerce").astype("Int64")
    res, wb, _ = _construir(tmp_path, ret=ret)
    c = res["retenciones"]
    assert (c["documentos"], c["diferencia"], c["coinciden"], c["solo_odoo"]) == (1, 1, 0, 0)
    assert wb["Ret_Documentos"].max_row - 1 == 1
    assert len({wb["Ret_Conciliacion"][f"Q{r}"].value for r in (2, 3)}) == 1
