import pytest
from pathlib import Path
from build_workbook_final import build_workbook

FIXTURES = Path(__file__).parent.parent

def test_build_workbook_returns_stats(tmp_path):
    """build_workbook debe retornar dict con n_listado y n_xml."""
    result = build_workbook(
        txt_path=str(FIXTURES / "2390000831001_Recibidos.txt"),
        xml_glob=str(FIXTURES / "xml_samples/[0-9]*.xml"),
        output_path=str(tmp_path / "out.xlsx"),
    )
    assert "n_listado" in result
    assert "n_xml" in result
    assert "n_pendientes" in result
    assert result["n_listado"] > 0
    assert Path(tmp_path / "out.xlsx").exists()

def test_build_workbook_file_created(tmp_path):
    """El archivo Excel de salida debe crearse en la ruta indicada."""
    out = tmp_path / "subdir" / "resultado.xlsx"
    out.parent.mkdir()
    build_workbook(
        txt_path=str(FIXTURES / "2390000831001_Recibidos.txt"),
        xml_glob=str(FIXTURES / "xml_samples/[0-9]*.xml"),
        output_path=str(out),
    )
    assert out.exists()
    assert out.stat().st_size > 0


def test_build_workbook_ignora_xml_fuera_del_txt(tmp_path):
    """Los XML en caché de otros períodos (claves ausentes del TXT cargado) no
    deben entrar a SRI_Detalle_Retenciones, SRI_Por_Doc_Sustento ni n_xml."""
    import shutil
    from openpyxl import load_workbook

    # TXT del "período nuevo": cabecera + 1 sola fila del TXT real
    lineas = (FIXTURES / "2390000831001_Recibidos.txt").read_bytes().splitlines(keepends=True)
    txt = tmp_path / "Recibidos.txt"
    txt.write_bytes(b"".join(lineas[:2]))
    from parse_recibidos_txt import parse_recibidos_txt
    claves_txt = set(parse_recibidos_txt(str(txt))["CLAVE_ACCESO"].dropna())
    assert len(claves_txt) == 1

    # Caché compartido: el XML del TXT + otros XML de "otro período"
    xmls = sorted((FIXTURES / "xml_samples").glob("[0-9]*.xml"))
    dentro = [x for x in xmls if x.stem in claves_txt]
    fuera = [x for x in xmls if x.stem not in claves_txt][:3]
    assert dentro and fuera
    cache = tmp_path / "cache"
    cache.mkdir()
    for x in dentro + fuera:
        shutil.copy(x, cache / x.name)

    out = tmp_path / "out.xlsx"
    result = build_workbook(
        txt_path=str(txt),
        xml_glob=str(cache / "[0-9]*.xml"),
        output_path=str(out),
    )

    ws = load_workbook(out)["SRI_Detalle_Retenciones"]
    headers = [c.value for c in ws[1]]
    idx = headers.index("Clave Acceso")
    claves_hoja = {r[idx] for r in ws.iter_rows(min_row=2, values_only=True) if r[idx]}
    assert claves_hoja <= claves_txt
    assert result["n_xml"] == 1
