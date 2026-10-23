# -*- coding: utf-8 -*-
from datetime import date

import openpyxl
import pandas as pd
import pytest

import conciliar_retenciones as cr
import excel_emitidos
from conciliar_emitidos import (EST_ANULADO, EST_DIF_DATOS, EST_DIF_VALOR, EST_INFO, EST_OK,
                                EST_SOLO_ODOO, EST_SOLO_SRI, clasificar_documentos)
from excel_emitidos import construir_excel_emitidos

DESDE, HASTA = date(2026, 9, 1), date(2026, 9, 30)
F = pd.Timestamp("2026-09-05")
SRI_COLS = ["clave", "serie", "fecha", "tipo", "base", "iva", "total", "ruc", "proveedor", "afectada"]
ODOO_COLS = ["id", "num", "nfac", "socio", "ruc", "fecha", "tipo", "clave", "base", "iva", "total",
             "afectada", "estado", "electronico", "elegible"]
HOJAS = ["Resumen", "Conciliacion", "Solo_en_SRI", "Solo_en_Odoo", "Diferencias", "Anulados_Odoo"]


@pytest.fixture(autouse=True)
def _sin_libreoffice(monkeypatch):
    monkeypatch.setattr(excel_emitidos, "_recalcular", lambda ruta: False)


def _sri(*filas):
    base = {"clave": "1" * 49, "serie": "001-001-000000100", "fecha": F, "tipo": "FAC", "base": 100.0,
            "iva": 15.0, "total": 115.0, "ruc": pd.NA, "proveedor": pd.NA, "afectada": pd.NA}
    return pd.DataFrame([{**base, **f} for f in filas], columns=SRI_COLS)


def _odoo(*filas):
    base = {"id": 1, "num": "FAC/2026/0100", "nfac": "001-001-000000100", "socio": "CLIENTE SA",
            "ruc": "1790000000001", "fecha": F, "tipo": "FAC", "clave": "1" * 49, "base": 100.0,
            "iva": 15.0, "total": 115.0, "afectada": pd.NA, "estado": "posted", "electronico": True,
            "elegible": True}
    return pd.DataFrame([{**base, **f} for f in filas], columns=ODOO_COLS)


def _datos(cliente="CLIENTE SA"):
    sri = _sri(
        {},  # coincide
        {"clave": "2" * 49, "serie": "001-001-000000101", "total": 120.0},  # diferencia de valor
        {"clave": "3" * 49, "serie": "001-001-000000102"},  # solo SRI
        {"clave": "4" * 49, "serie": "001-001-000000103"},  # anulado en Odoo
        {"clave": "5" * 49, "serie": "001-002-000000001", "tipo": "NC", "base": 10.0, "iva": 1.5,
         "total": 11.5},
    )
    odoo = _odoo(
        {"socio": cliente},
        {"id": 2, "num": "FAC/2026/0101", "nfac": "001-001-000000101", "clave": "2" * 49},
        {"id": 4, "num": "FAC/2026/0103", "nfac": "001-001-000000103", "clave": "4" * 49,
         "estado": "cancel"},
        {"id": 5, "num": "FAC/2026/0104", "nfac": "001-001-000000104", "clave": "6" * 49},  # solo Odoo
        {"id": 6, "num": "NC/2026/0001", "nfac": "001-002-000000001", "clave": "5" * 49, "tipo": "NC",
         "base": 10.0, "iva": 1.5, "total": 11.5},
    )
    return sri, odoo


def _fila_ret(estado, n_ret, vs=10.0, vo=10.0):
    return {"estado": estado, "n_ret": n_ret, "fecha": F, "ruc": "1790000000001", "proveedor": "CLIENTE SA",
            "factura_sri": "001-001-000000010", "factura_odoo": "001-001-000000010", "pct_sri": 1.0,
            "pct_odoo": 1.0, "base_sri": 1000.0, "base_odoo": 1000.0, "valor_sri": vs, "valor_odoo": vo,
            "dif_valor": None, "id_ret": 7, "clave": "9" * 49}


def _ret():
    df = pd.DataFrame([_fila_ret(cr.EST_OK, "001-001-000000001"),
                       _fila_ret(cr.EST_SOLO_SRI, "001-001-000000002", vo=None),
                       _fila_ret(cr.EST_INFO, "001-001-000000003", vs=None),
                       _fila_ret(cr.EST_ANULADO, "001-001-000000004")], columns=cr.COLUMNAS)
    df["id_ret"] = df["id_ret"].astype("Int64")
    return df


def _construir(tmp_path, sri=None, odoo=None, ret=None, **kw):
    if sri is None:
        sri, odoo = _datos(**kw)
    docs = clasificar_documentos(sri, odoo, DESDE, HASTA)
    salida = str(tmp_path / "out.xlsx")
    res = construir_excel_emitidos(sri, odoo, docs, ret, DESDE, HASTA, salida)
    return res, openpyxl.load_workbook(salida, data_only=False)


def _col(ws, nombre):
    return [c.value for c in ws[1]].index(nombre) + 1


def test_hojas_y_resumen(tmp_path):
    res, wb = _construir(tmp_path, ret=_ret())
    assert wb.sheetnames[:len(HOJAS)] == HOJAS
    assert {"Ret_Documentos", "Ret_Conciliacion"} <= set(wb.sheetnames)
    colores = {"Resumen": "1F4E79", "Conciliacion": "1F4E79", "Solo_en_SRI": "ED7D31",
               "Solo_en_Odoo": "C00000", "Diferencias": "FFC000", "Anulados_Odoo": "7030A0"}
    for nombre, color in colores.items():
        assert wb[nombre].sheet_properties.tabColor.rgb.endswith(color), nombre
    formulas = [c.value for row in wb["Resumen"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=COUNTIFS")]
    assert any("Conciliacion!" in f for f in formulas)
    assert any("Ret_Documentos!" in f for f in formulas)
    # el informativo de documentos y el de retenciones nunca se mezclan
    assert any(f'"{EST_INFO}"' in f and "Conciliacion!" in f for f in formulas)
    assert all(f'"{EST_INFO}"' not in f for f in formulas if "Ret_Documentos!" in f)
    fac = res["documentos"]["FAC"]
    assert fac["total"] == 5
    assert (fac[EST_OK], fac[EST_DIF_VALOR], fac[EST_SOLO_SRI], fac[EST_SOLO_ODOO], fac[EST_ANULADO]) == \
        (1, 1, 1, 1, 1)
    assert res["documentos"]["NC"]["total"] == 1 and res["documentos"]["NC"][EST_OK] == 1
    r = res["documentos"]["RET"]
    assert r["total"] == 4
    assert (r[cr.EST_OK], r[cr.EST_SOLO_SRI], r[cr.EST_INFO], r[cr.EST_ANULADO]) == (1, 1, 1, 1)
    # hojas filtradas: copias por estado
    assert wb["Solo_en_SRI"].max_row == 2 and wb["Solo_en_Odoo"].max_row == 2
    assert wb["Diferencias"].max_row == 2 and wb["Anulados_Odoo"].max_row == 2
    ws = wb["Conciliacion"]
    assert ws.max_row == 1 + 6
    assert [c.value for c in ws[1]][:10] == ["Tipo", "Estado", "Número", "Fecha_SRI", "Fecha_Odoo", "Cliente",
                                             "Total_SRI", "Total_Odoo", "Diferencia", "Detalle"]


def _fila_resumen(ws, etiqueta):
    return next(r for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=1).value == etiqueta)


def test_resumen_formulas_fila_fac(tmp_path):
    _, wb = _construir(tmp_path)
    rs = wb["Resumen"]
    r = _fila_resumen(rs, "FAC")
    assert r == 5
    conc = "Conciliacion!${c}$2:${c}$7"
    a, b, g = (conc.format(c=c) for c in "ABG")
    assert rs["D5"].value == f'=COUNTIFS({a},$A5,{b},"{EST_OK}")'
    assert rs["K5"].value == f'=SUMIFS({g},{a},$A5,{b},"<>Duplicado en Odoo",{b},"<>{EST_INFO}")'


def test_total_odoo_retenciones_excluye_anuladas(tmp_path):
    import re
    ret = _ret()
    _, wb = _construir(tmp_path, ret=ret)
    rs = wb["Resumen"]
    r = _fila_resumen(rs, "Retenciones")
    f = rs[f"L{r}"].value
    assert f'"<>{cr.EST_INFO}"' in f and f'"<>{cr.EST_ANULADO}"' in f
    assert f.startswith("=SUMIFS(Ret_Documentos!$K$")
    # evaluar el SUMIFS: estados (columna A, texto) de Ret_Documentos y valor Odoo por documento
    excluidos = set(re.findall(r'"<>([^"]*)"', f))
    odoo_por_doc = ret.groupby("n_ret")["valor_odoo"].sum(min_count=1).fillna(0).to_dict()
    ws = wb["Ret_Documentos"]
    total = sum(odoo_por_doc[ws.cell(row=i, column=2).value] for i in range(2, ws.max_row + 1)
                if ws.cell(row=i, column=1).value not in excluidos)
    # solo cuenta la retención que coincide (10); la anulada (10) y la informativa (10) no
    assert total == 10.0
    assert rs[f"M{r}"].value == f"=K{r}-L{r}"


def test_diferencia_es_formula(tmp_path):
    _, wb = _construir(tmp_path)
    for hoja in ("Conciliacion", "Diferencias"):
        ws = wb[hoja]
        col = _col(ws, "Diferencia")
        con_formula = 0
        for r in range(2, ws.max_row + 1):
            c = ws.cell(row=r, column=col)
            if c.value is not None:
                assert isinstance(c.value, str) and c.value.startswith("=ROUND("), (hoja, r, c.value)
                assert c.data_type == "f"
                con_formula += 1
        assert con_formula >= 1, hoja


def test_nombre_cliente_con_signo_igual_no_es_formula(tmp_path):
    _, wb = _construir(tmp_path, cliente="=1+1")
    for hoja in ("Conciliacion",):
        ws = wb[hoja]
        col = _col(ws, "Cliente")
        celdas = [ws.cell(row=r, column=col) for r in range(2, ws.max_row + 1)]
        c = next(c for c in celdas if c.value == "=1+1")
        assert c.data_type != "f"


def test_sin_retenciones(tmp_path):
    res, wb = _construir(tmp_path, ret=None)
    assert not [s for s in wb.sheetnames if s.startswith("Ret_")]
    tipos = [wb["Resumen"].cell(row=r, column=1).value for r in range(1, wb["Resumen"].max_row + 1)]
    assert "Retenciones" not in tipos
    assert "RET" not in res["documentos"]


def test_envio_con_un_solo_tipo(tmp_path):
    sri = _sri({"tipo": "NC", "total": 11.5, "base": 10.0, "iva": 1.5})
    odoo = _odoo({"tipo": "NC", "total": 11.5, "base": 10.0, "iva": 1.5})
    res, wb = _construir(tmp_path, sri=sri, odoo=odoo)
    docs = res["documentos"]
    assert docs["NC"]["total"] == 1 and docs["NC"][EST_OK] == 1
    for t in ("FAC", "ND", "LIQ"):
        assert docs[t]["total"] == 0
        assert all(v == 0 for v in docs[t].values())
    assert EST_DIF_DATOS in docs["FAC"]


def test_fuente_arial(tmp_path):
    _, wb = _construir(tmp_path)
    assert wb["Resumen"]["A1"].font.name == "Arial"
    assert wb["Conciliacion"]["A2"].font.name == "Arial"
